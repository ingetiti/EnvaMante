"""
SECTOR DE MANTENIMIENTO - Carga de actividades (v3, multiusuario)
Réplica de UserForm1 de CargaActividades.xlsm sobre SQLite.

- La primera vez importa la hoja 'bbdd' y las listas de 'Datos' del Excel (solo lectura).
- La ubicación de la base se define en config.ini (junto al programa). Puede estar en una carpeta de red
  compartida para que varias PC trabajen sobre la misma base.
- Agregar / Editar / Eliminar guardan directo en la base.
- Respaldo automático diario, refresco automático de la lista y aviso si otro usuario modificó
  la actividad que se está editando.

Requiere: pip install openpyxl   (sqlite3 y tkinter vienen con Python)
"""
import os
import re
import sys
import socket
import hashlib
import secrets
import sqlite3
import configparser
from datetime import datetime, date, time

from openpyxl import load_workbook, Workbook

# ======================= Configuración =======================
CARPETA = os.path.dirname(sys.executable if getattr(sys, "frozen", False) else os.path.abspath(__file__))
RUTA_CONFIG = os.path.join(CARPETA, "config.ini")
EQUIPO = os.environ.get("COMPUTERNAME") or socket.gethostname()
LIMITE_REFRESCO_AUTO = 3000   # listas más grandes se refrescan solo con el botón REFRESCAR
CONFIG_DEFECTO = {
    "base": {"ruta": "mantenimiento.db"},
    "respaldo": {"carpeta": "", "dias": "30"},
    "app": {"refresco_seg": "30"},
}
HOJA_BBDD, HOJA_DATOS = "bbdd", "Datos"
INICIO_CICLOS = date(2026, 1, 3)   # igual que en el VBA
DIAS_CICLO = 8

# Listas que se leen de la hoja Datos: nombre -> (columna, fila_desde, fila_hasta)
LISTAS_DATOS = {
    "operarios": ("A", 2, 30),
    "sectores": ("D", 2, 30),
    "estado": ("B", 27, 36),
    "prioridad": ("B", 38, 41),
    "planificado": ("B", 45, 46),
    "extras": ("B", 57, 59),
    "maq_parada": ("B", 62, 67),
    "tipo_pcpg": ("B", 69, 74),
    "supervisores": ("B", 82, 85),
}

# Encabezados de bbdd usados por el formulario
OT, OT_SEC, OPERARIO, FECHA, SECTOR, MAQUINA = "OT MAN", "OT SEC", "OPERARIO", "FECHA", "SECTOR", "MAQUINA"
PRIORIDAD, ACTIVIDAD, HI, HF, OT_REPLAN, QREPLAN = "PRIORIDAD", "ACTIVIDAD", "HI", "HF", "OT REPLAN", "QREPLAN"
TIPO_END, SEM, ANIO, CONJ, SUBCONJ, CODACT = "TIPO (E/N/D)", "SEM/CIC", "AÑO", "CONJUNTO", "SUB-CONJ", "COD/ACT"
ESTADO, PLANIF, MAQ_PAR, TIPO_PCPG, CAT = "ESTADO", "PLANIF (S/N)", "MAQ_PAR", "TIPO (PV-CO-PD-GE)", "CAT"
REPUESTOS, MEJORAS, COMENT, CONTROL = "REPUESTOS UTILIZADOS", "MEJORAS PROPUESTAS", "COMENTARIOS", "CONTROL_ACTIVIDAD"

AYUDA = {
    OT_SEC: "OT del sector que pidió la actividad.",
    OT_REPLAN: "OT de la actividad replanificada (prototipo aún).",
    QREPLAN: "Cantidad de veces que fue replanificada la actividad.",
    FECHA: "Fecha de realización de la actividad (dd-mm-aaaa). Si figura 11/11/1111, es una actividad semanal.",
    SECTOR: "Sector donde se realizará (o realizó) la actividad.",
    MAQUINA: "Máquina o sector específico dónde se realizará (o realizó) la actividad.",
    CONJ: "Parte MACRO de una máquina o sector (prototipo).",
    SUBCONJ: "Sub parte de una máquina o sector (prototipo), debe estar contenida en un CONJUNTO.",
    CODACT: "Código que tendrá una pieza/parte/sector/etc en específico. Es un código único, sigue en prototipo.",
    PRIORIDAD: "ALTA, MEDIA, BAJA.",
    ACTIVIDAD: "Descripción de la actividad a realizar o realizada. Ser lo más específico que se pueda por favor.",
    MEJORAS: "Posibles mejoras que pueden hacerse en la pieza, máquina, sector y/o proceso que se intervino.",
    REPUESTOS: "Escribir con detalle los repuestos que se utilizaron para facilitar su reposición. "
               "En caso de hacer un pedido de insumos, detallar el número de SRH.",
    COMENT: "Escriba cualquier cosa que enriquezca la descripción de la actividad. "
            "O alguna observación que notó durante el desarrollo de la actividad.",
    PLANIF: "Por defecto, SI. A cualquier actividad que no fue planificada con antelación, se le pone NO "
            "(por ejemplo: cuando nos buscan para reparar algo urgente en máquina).",
    ESTADO: "TER: Terminado.\nEP_FT: En proceso, faltó tiempo (especificar cuánto más falta).\n"
            "EP_FM: En proceso, faltó material (Hacer SRH).\n"
            "EP_PR: En proceso, es proyecto (especificar cuánto tiempo más falta).\n"
            "NOTER: No hecho (especificar por qué).",
    CAT: "ID de proyecto (se guarda en la columna CAT, igual que en el Excel).",
    TIPO_PCPG: "PV: Preventivo.\nCO: Correctivo.\nPD: Predictivo.\nGE: General.",
    MAQ_PAR: "SI_Planif: con parada de máquina planificada.\nSI_PorFallo: con parada de máquina por falla.\n"
             "NO_EnMov: con la máquina en producción.\nNO_SinProd: con la máquina parada sin producción.\n"
             "NO_APLICA: la actividad no se realiza sobre una máquina.",
    TIPO_END: "NOEXTRA, EXTRA, DEV. DEV: devolución de horas por un cambio de jornada.",
    HI: "Hora a la que se inició la actividad (hh:mm).",
    HF: "Hora a la que se terminó la actividad (hh:mm).",
}


# ======================= Utilidades =======================
def ciclo_actual(hoy=None):
    hoy = hoy or date.today()
    return f"{(hoy - INICIO_CICLOS).days // DIAS_CICLO + 1:02d}"


def armar_codigo(anio2, ciclo, operario):
    """AÑO(2) + CICLO + 2 primeras letras del operario (igual que ActualizarCodigo del VBA)."""
    anio2, ciclo, operario = anio2.strip(), ciclo.strip(), operario.strip()
    if len(anio2) >= 2 and ciclo:
        return anio2[-2:] + ciclo + (operario[:2] if len(operario) >= 2 else "")
    return ""


def valor_a_db(v):
    """Excel -> SQLite. Fechas ISO (ordenables), horas hh:mm."""
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d") if v.time() == time(0) else v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, time):
        return v.strftime("%H:%M")
    if isinstance(v, str):
        return v.strip() or None
    return v


def db_a_texto(v):
    """SQLite -> texto en pantalla (fechas dd-mm-aaaa)."""
    if v is None:
        return ""
    if isinstance(v, str):
        m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})( \d{2}:\d{2})?", v)
        if m:
            return f"{m[3]}-{m[2]}-{m[1]}"
    return str(v)


def fecha_a_db(t):
    """'dd-mm-aaaa' o 'dd/mm/aaaa' -> ISO. Devuelve None si no es válida."""
    m = re.fullmatch(r"(\d{1,2})[-/.](\d{1,2})[-/.](\d{2}|\d{4})", t.strip())
    if not m:
        return None
    a = int(m[3])
    if a < 100:
        a += 2000 if a < 30 else 1900
    try:
        return date(a, int(m[2]), int(m[1])).strftime("%Y-%m-%d")
    except ValueError:
        return None


def hora_a_db(t):
    """'8', '8:5', '08:05', '08:05:00' -> 'hh:mm'. None si no es válida."""
    m = re.fullmatch(r"(\d{1,2})(?::(\d{1,2}))?(?::\d{1,2})?", t.strip())
    if not m:
        return None
    h, mi = int(m[1]), int(m[2] or 0)
    return f"{h:02d}:{mi:02d}" if h < 24 and mi < 60 else None


def db_a_excel(v):
    if isinstance(v, str):
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", v):
            return datetime.strptime(v, "%Y-%m-%d")
        if re.fullmatch(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2}", v):
            return datetime.strptime(v, "%Y-%m-%d %H:%M")
        if re.fullmatch(r"\d{2}:\d{2}", v):
            return datetime.strptime(v, "%H:%M").time()
    return v


def validar(d, supervisor):
    """Mismas reglas que btnAgregar del VBA. Devuelve el mensaje de error o None."""
    vacio = lambda k: not str(d.get(k) or "").strip()
    if vacio(OPERARIO):
        return "Nombre de operario en blanco."
    if vacio(FECHA):
        return "Debe ingresar una fecha."
    if vacio(SECTOR) or vacio(MAQUINA):
        return "Debe ingresar una máquina/sector."
    if supervisor:
        return None
    if vacio(HI) or vacio(HF):
        return "Debe ingresar una hora de inicio y fin."
    if vacio(ACTIVIDAD):
        return "Debe ingresar una actividad."
    if vacio(TIPO_END):
        return "Debe especificar si fueron extras, no extras o devolución."
    if vacio(COMENT):
        return "Debe ingresar un comentario o propuesta de mejora."
    if vacio(ESTADO):
        return "Agregar un estado de la actividad."
    if d[ESTADO] == "EP_FM" and vacio(REPUESTOS):
        return "Debe especificar el número de SRH del pedido de insumos que realizó para esta actividad."
    return None


# ======================= config.ini y respaldos =======================
def leer_config():
    cp = configparser.ConfigParser()
    cp.read_dict(CONFIG_DEFECTO)
    if os.path.exists(RUTA_CONFIG):
        cp.read(RUTA_CONFIG, encoding="utf-8")
    else:
        guardar_config(cp)
    return cp


def guardar_config(cp):
    try:
        with open(RUTA_CONFIG, "w", encoding="utf-8") as f:
            f.write("; [base] ruta: archivo de la base. Relativa a la carpeta del programa o absoluta,\n"
                    ";   por ejemplo: \\\\PC-TALLER\\Mantenimiento\\mantenimiento.db\n"
                    "; [respaldo] carpeta: vacío = subcarpeta 'respaldos' junto a la base. dias: copias a conservar.\n"
                    "; [app] refresco_seg: cada cuántos segundos se buscan cambios de otros usuarios (0 = nunca).\n\n")
            cp.write(f)
    except OSError:
        pass


def ruta_abs(p):
    """Rutas relativas se toman desde la carpeta del programa. Acepta rutas de red (\\\\servidor\\carpeta)."""
    p = os.path.expandvars(p.strip())
    return p if os.path.isabs(p) else os.path.join(CARPETA, p)


def respaldar(ruta_db, carpeta="", dias=30, manual=False):
    """Copia consistente de la base (API de respaldo de SQLite, segura aunque otros estén usando la base).
    Automático: una copia por día, la hace la primera PC que abre el programa. Devuelve la ruta o None."""
    if not os.path.exists(ruta_db):
        return None
    carpeta = carpeta.strip() or os.path.join(os.path.dirname(ruta_db), "respaldos")
    os.makedirs(carpeta, exist_ok=True)
    base = os.path.splitext(os.path.basename(ruta_db))[0]
    sufijo = f"{datetime.now():%Y%m%d_%H%M%S}_manual" if manual else f"{date.today():%Y%m%d}"
    destino = os.path.join(carpeta, f"{base}_{sufijo}.db")
    if os.path.exists(destino):
        return None
    tmp = f"{destino}.{EQUIPO}.tmp"
    src, dst = sqlite3.connect(ruta_db, timeout=15), sqlite3.connect(tmp)
    try:
        src.backup(dst)
    finally:
        dst.close()
        src.close()
    os.replace(tmp, destino)
    if not manual and dias > 0:  # rotación: se conservan los últimos 'dias' respaldos automáticos
        patron = re.compile(re.escape(base) + r"_\d{8}\.db")
        viejos = sorted(f for f in os.listdir(carpeta) if patron.fullmatch(f))
        for f in viejos[:-dias]:
            try:
                os.remove(os.path.join(carpeta, f))
            except OSError:
                pass
    return destino


def marca_modificacion():
    return f"{EQUIPO} {datetime.now():%d-%m-%Y %H:%M}"


# ======================= Base de datos =======================
class BaseDatos:
    def __init__(self, ruta):
        self.ruta = ruta
        self.con = sqlite3.connect(ruta, timeout=15)  # espera hasta 15 s si otra PC está guardando
        self.con.executescript("""
            CREATE TABLE IF NOT EXISTS columnas (pos INTEGER PRIMARY KEY, nombre TEXT);
            CREATE TABLE IF NOT EXISTS listas (lista TEXT, orden INTEGER, valor TEXT, padre TEXT);
            CREATE TABLE IF NOT EXISTS config (clave TEXT PRIMARY KEY, valor TEXT);
        """)
        self.encabezados = [r[0] for r in self.con.execute("SELECT nombre FROM columnas ORDER BY pos")]
        self._migrar()

    def _migrar(self):
        """Agrega las columnas de control de cambios a bases creadas con versiones anteriores."""
        if self.vacia:
            return
        existentes = {r[1] for r in self.con.execute("PRAGMA table_info(actividades)")}
        if "rev" not in existentes:
            self.con.execute("ALTER TABLE actividades ADD COLUMN rev INTEGER DEFAULT 0")
        if "mod_info" not in existentes:
            self.con.execute("ALTER TABLE actividades ADD COLUMN mod_info TEXT")
        self.con.commit()

    def data_version(self):
        """Cambia cuando OTRA conexión (otra PC) guarda algo en la base."""
        return self.con.execute("PRAGMA data_version").fetchone()[0]

    @property
    def _select(self):
        return f"SELECT id, {', '.join(self.cols)} FROM actividades"

    @property
    def vacia(self):
        return not self.encabezados

    def col(self, nombre):
        return f"c{self.encabezados.index(nombre) + 1:02d}"

    @property
    def cols(self):
        return [f"c{i + 1:02d}" for i in range(len(self.encabezados))]

    # ---------- Importación ----------
    def importar_excel(self, ruta_excel):
        wb = load_workbook(ruta_excel, read_only=True, data_only=True)  # el Excel nunca se guarda
        try:
            n = self._importar_bbdd(wb[HOJA_BBDD])
            self._importar_listas(wb[HOJA_DATOS])
        finally:
            wb.close()
        self.con.commit()
        return n

    def _importar_bbdd(self, ws):
        filas = ws.iter_rows(values_only=True)
        enc = [str(h).strip() if h is not None else f"COL{i + 1}" for i, h in enumerate(next(filas))]
        n = len(enc)
        cols = [f"c{i + 1:02d}" for i in range(n)]
        cur = self.con.cursor()
        cur.execute("DROP TABLE IF EXISTS actividades")
        cur.execute("DELETE FROM columnas")
        cur.execute(f"CREATE TABLE actividades (id INTEGER PRIMARY KEY AUTOINCREMENT, {', '.join(cols)}, "
                    f"rev INTEGER DEFAULT 0, mod_info TEXT)")
        cur.executemany("INSERT INTO columnas VALUES (?, ?)", list(enumerate(enc, 1)))
        self.encabezados = enc
        cur.execute(f"CREATE INDEX idx_ot ON actividades ({self.col(OT)})")
        datos = []
        for f in filas:
            f = list(f[:n]) + [None] * (n - len(f))
            if any(v is not None for v in f):
                datos.append([valor_a_db(v) for v in f])
        cur.executemany(f"INSERT INTO actividades ({', '.join(cols)}) VALUES ({', '.join('?' * n)})", datos)
        return len(datos)

    def _importar_listas(self, ws):
        celdas = {}
        for r, fila in enumerate(ws.iter_rows(min_row=1, max_row=100, max_col=26, values_only=True), 1):
            for ci, v in enumerate(fila, 1):
                if v is not None and str(v).strip():
                    celdas[(chr(64 + ci), r)] = str(v).strip()
        reg = []
        for nombre, (col, desde, hasta) in LISTAS_DATOS.items():
            vals = [celdas.get((col, r)) for r in range(desde, hasta + 1)]
            reg += [(nombre, i, v, None) for i, v in enumerate(x for x in vals if x)]
        # Máquinas por sector: encabezados en E1:Z1
        for ci in range(5, 27):
            letra = chr(64 + ci)
            sector = celdas.get((letra, 1))
            if sector:
                for r in range(2, 101):
                    v = celdas.get((letra, r))
                    if v:
                        reg.append(("maquinas", r, v, sector))
        self.con.execute("DELETE FROM listas")
        self.con.executemany("INSERT INTO listas VALUES (?, ?, ?, ?)", reg)

    def lista(self, nombre, padre=None):
        sql, p = "SELECT valor FROM listas WHERE lista = ?", [nombre]
        if padre is not None:
            sql += " AND padre = ?"
            p.append(padre)
        vistos, res = set(), []
        for (v,) in self.con.execute(sql + " ORDER BY orden", p):
            if v not in vistos:
                vistos.add(v)
                res.append(v)
        return res

    # ---------- Consultas ----------
    def consultar(self, where="", params=(), orden=None):
        orden = orden or f"{self.col(FECHA)}, id"
        sql = self._select + (f" WHERE {where}" if where else "") + f" ORDER BY {orden}"
        return self.con.execute(sql, params).fetchall()

    def por_codigo(self, prefijo):
        return self.consultar(f"CAST({self.col(OT)} AS TEXT) LIKE ?", (prefijo + "%",))

    def filtrar(self, codigo="", sector="", maquina="", proyecto=""):
        w, p = [], []
        if codigo and codigo.upper() != "TT":
            w.append(f"CAST({self.col(OT)} AS TEXT) LIKE ?"); p.append(f"%{codigo}%")
        if sector:
            w.append(f"{self.col(SECTOR)} LIKE ?"); p.append(f"%{sector}%")
        if maquina:
            w.append(f"{self.col(MAQUINA)} LIKE ?"); p.append(f"%{maquina}%")
        if proyecto:
            w.append(f"CAST({self.col(CAT)} AS TEXT) LIKE ?"); p.append(f"%{proyecto}%")
        return self.consultar(" AND ".join(w), p)

    def obtener(self, id_):
        return self.con.execute(self._select + " WHERE id = ?", (id_,)).fetchone()

    def obtener_rev(self, id_):
        """(rev, mod_info) del registro, o None si fue eliminado."""
        return self.con.execute("SELECT rev, mod_info FROM actividades WHERE id = ?", (id_,)).fetchone()

    def nueva_ot(self, codigo):
        """codigo + (mayor correlativo existente + 1), como GenerarNuevaOT del VBA."""
        mayor = 0
        for (ot,) in self.con.execute(f"SELECT {self.col(OT)} FROM actividades WHERE CAST({self.col(OT)} AS TEXT) LIKE ?",
                                      (codigo + "%",)):
            resto = str(ot)[len(codigo):]
            if resto.isdigit():
                mayor = max(mayor, int(resto))
        return f"{codigo}{mayor + 1}"

    # ---------- Escritura ----------
    def insertar(self, valores_dict):
        cols = [self.col(k) for k in valores_dict] + ["mod_info"]
        self.con.execute(f"INSERT INTO actividades ({', '.join(cols)}) VALUES ({', '.join('?' * len(cols))})",
                         list(valores_dict.values()) + [marca_modificacion()])
        self.con.commit()

    def insertar_nueva(self, codigo, valores_dict):
        """Genera la OT y la inserta en una sola transacción bloqueante, así dos PC
        que cargan con el mismo código al mismo tiempo no obtienen la misma OT."""
        self.con.execute("BEGIN IMMEDIATE")
        try:
            ot = self.nueva_ot(codigo)
            datos = dict(valores_dict)
            datos[OT] = ot
            cols = [self.col(k) for k in datos] + ["mod_info"]
            self.con.execute(f"INSERT INTO actividades ({', '.join(cols)}) VALUES ({', '.join('?' * len(cols))})",
                             list(datos.values()) + [marca_modificacion()])
            self.con.commit()
            return ot
        except Exception:
            self.con.rollback()
            raise

    def actualizar(self, id_, valores_dict, rev=None):
        """Guarda solo si nadie modificó el registro desde que se cargó (rev).
        Devuelve 'ok', 'conflicto' (otro lo modificó) o 'eliminado'. Con rev=None guarda siempre."""
        sets = ", ".join(f"{self.col(k)} = ?" for k in valores_dict)
        sql = f"UPDATE actividades SET {sets}, rev = COALESCE(rev, 0) + 1, mod_info = ? WHERE id = ?"
        params = list(valores_dict.values()) + [marca_modificacion(), id_]
        if rev is not None:
            sql += " AND COALESCE(rev, 0) = ?"
            params.append(rev)
        cur = self.con.execute(sql, params)
        self.con.commit()
        if cur.rowcount:
            return "ok"
        return "eliminado" if self.obtener_rev(id_) is None else "conflicto"

    def actualizar_campo(self, ids, nombre, valor):
        marca = marca_modificacion()
        self.con.executemany(f"UPDATE actividades SET {self.col(nombre)} = ?, rev = COALESCE(rev, 0) + 1, "
                             f"mod_info = ? WHERE id = ?", [(valor, marca, i) for i in ids])
        self.con.commit()

    def eliminar(self, id_):
        self.con.execute("DELETE FROM actividades WHERE id = ?", (id_,))
        self.con.commit()

    def exportar(self, ruta, filas=None):
        wb = Workbook(write_only=True)
        ws = wb.create_sheet(HOJA_BBDD)
        ws.append(self.encabezados)
        for f in (filas if filas is not None else self.consultar(orden="id")):
            ws.append([db_a_excel(v) for v in f[1:]])
        wb.save(ruta)

    # ---------- Contraseña supervisor ----------
    def hay_clave(self):
        return self.con.execute("SELECT 1 FROM config WHERE clave = 'sup_hash'").fetchone() is not None

    def fijar_clave(self, clave):
        sal = secrets.token_hex(16)
        h = hashlib.sha256((sal + clave).encode()).hexdigest()
        self.con.execute("INSERT OR REPLACE INTO config VALUES ('sup_hash', ?)", (f"{sal}${h}",))
        self.con.commit()

    def verificar_clave(self, clave):
        r = self.con.execute("SELECT valor FROM config WHERE clave = 'sup_hash'").fetchone()
        if not r:
            return False
        sal, h = r[0].split("$")
        return secrets.compare_digest(hashlib.sha256((sal + clave).encode()).hexdigest(), h)


# ======================= Interfaz =======================
def iniciar_interfaz():
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox, simpledialog

    class App(tk.Tk):
        def __init__(self):
            super().__init__()
            self.title("SECTOR DE MANTENIMIENTO")
            self.geometry("1320x720")
            self.minsize(1100, 650)
            self.cfg = leer_config()
            ruta = self._resolver_ruta_db(ruta_abs(self.cfg.get("base", "ruta")))
            if not ruta:
                self.destroy()
                raise SystemExit
            self.msg_respaldo = self._respaldo_automatico(ruta)
            try:
                self.db = BaseDatos(ruta)
            except sqlite3.Error as ex:
                messagebox.showerror("Base de datos", f"No se pudo abrir la base:\n{ruta}\n\n{ex}")
                self.destroy()
                raise SystemExit
            self.supervisor = False
            self.id_cargado = None       # registro cargado en el formulario
            self.rev_cargado = None      # versión del registro al cargarlo (para detectar cambios de otros)
            self.filas_lista = []        # filas mostradas en la lista
            self.consulta = None         # función que vuelve a generar la lista (BUSCAR o Filtrar)
            self.cambios_pendientes = False
            self.ultimo_refresco = None
            self.campos = {}             # encabezado -> widget
            self.solo_sup = []           # widgets habilitados solo en modo supervisor
            self._armar()
            self._aplicar_permisos()
            self.dv = self.db.data_version()
            self._actualizar_estado()
            if self.db.vacia:
                self.after(200, self.importar)
            else:
                self._cargar_listas()
            seg = self.cfg.getint("app", "refresco_seg", fallback=30)
            if seg > 0:
                self.after(seg * 1000, lambda: self._vigilar(seg))

        # ---------------- Base compartida ----------------
        def _resolver_ruta_db(self, ruta):
            """Verifica que la carpeta de la base sea accesible (por ejemplo, la carpeta de red)."""
            while not os.path.isdir(os.path.dirname(ruta) or "."):
                if not messagebox.askyesno(
                        "Base de datos no accesible",
                        f"No se encuentra la carpeta de la base:\n{os.path.dirname(ruta)}\n\n"
                        "Puede ser un problema de red o una ruta mal configurada en config.ini.\n\n"
                        "¿Desea elegir otra ubicación? (No = cerrar el programa)", icon="warning"):
                    return None
                nueva = self._elegir_ruta_db()
                if nueva:
                    ruta = nueva
            return ruta

        def _elegir_ruta_db(self):
            ruta = filedialog.asksaveasfilename(
                title="Elegir la base existente o una nueva", defaultextension=".db",
                filetypes=[("Base de datos", "*.db")], initialfile="mantenimiento.db", confirmoverwrite=False)
            if ruta:
                self.cfg.set("base", "ruta", os.path.normpath(ruta))
                guardar_config(self.cfg)
            return ruta

        def cambiar_ubicacion(self):
            if self._elegir_ruta_db():
                messagebox.showinfo("Ubicación de la base",
                                    "Ubicación guardada en config.ini.\nCierre y vuelva a abrir el programa.")

        def _respaldo_automatico(self, ruta):
            try:
                hecho = respaldar(ruta, self.cfg.get("respaldo", "carpeta"),
                                  self.cfg.getint("respaldo", "dias", fallback=30))
                return f"Respaldo del día creado: {os.path.basename(hecho)}" if hecho else ""
            except Exception as ex:
                return f"⚠ No se pudo crear el respaldo diario: {ex}"

        def respaldar_ahora(self):
            try:
                hecho = respaldar(self.db.ruta, self.cfg.get("respaldo", "carpeta"), manual=True)
                messagebox.showinfo("Respaldo", f"Respaldo creado:\n{hecho}")
            except Exception as ex:
                messagebox.showerror("Respaldo", f"No se pudo crear el respaldo:\n{ex}")

        def _vigilar(self, seg):
            """Cada 'seg' segundos revisa si otra PC guardó cambios y refresca la lista."""
            try:
                dv = self.db.data_version()
                if dv != self.dv:
                    self.dv = dv
                    if self.consulta and len(self.filas_lista) <= LIMITE_REFRESCO_AUTO:
                        self.refrescar()
                    elif self.consulta:
                        self.cambios_pendientes = True
                        self._actualizar_estado()
                    self._revisar_cargado()
                self.conexion_ok = True
            except sqlite3.Error:
                self.conexion_ok = False
                self._actualizar_estado()
            self.after(seg * 1000, lambda: self._vigilar(seg))

        def _revisar_cargado(self):
            """Avisa si la actividad que se está editando fue modificada o eliminada por otro usuario."""
            if self.id_cargado is None:
                return
            r = self.db.obtener_rev(self.id_cargado)
            if r is None:
                self._ayuda("⚠ La actividad que está editando fue ELIMINADA por otro usuario.")
            elif (r[0] or 0) != self.rev_cargado:
                self._ayuda(f"⚠ Otro usuario modificó esta actividad ({r[1]}). "
                            "Si guarda, se le preguntará qué versión conservar.")

        def _actualizar_estado(self):
            partes = [f"{len(self.filas_lista)} actividades en la lista"]
            if self.cambios_pendientes:
                partes.append("⚠ Hay cambios de otros usuarios: presione REFRESCAR")
            if getattr(self, "conexion_ok", True) is False:
                partes.append("⚠ SIN CONEXIÓN CON LA BASE")
            if self.ultimo_refresco:
                partes.append(f"Actualizado {self.ultimo_refresco:%H:%M:%S}")
            partes.append(f"Base: {self.db.ruta}")
            if self.msg_respaldo:
                partes.append(self.msg_respaldo)
            self.lbl_cuenta.config(text="   |   ".join(partes))

        # ---------------- Armado ----------------
        def _armar(self):
            menu = tk.Menu(self)
            archivo = tk.Menu(menu, tearoff=0)
            archivo.add_command(label="Importar desde Excel...", command=self.importar)
            archivo.add_command(label="Exportar lista actual a Excel...", command=lambda: self.exportar(False))
            archivo.add_command(label="Exportar base completa a Excel...", command=lambda: self.exportar(True))
            archivo.add_separator()
            archivo.add_command(label="Refrescar lista (F5)", command=self.refrescar)
            archivo.add_command(label="Respaldar ahora", command=self.respaldar_ahora)
            archivo.add_command(label="Ubicación de la base...", command=self.cambiar_ubicacion)
            archivo.add_separator()
            archivo.add_command(label="Salir", command=self.destroy)
            menu.add_cascade(label="Archivo", menu=archivo)
            self.config(menu=menu)

            self._armar_cabecera()
            self._armar_lista()
            self._armar_formulario()

        def _armar_cabecera(self):
            cab = ttk.Frame(self, padding=(10, 8, 10, 0))
            cab.pack(fill="x")

            ttk.Label(cab, text="OPERARIO").grid(row=0, column=0, sticky="w")
            self.cmb_oper = ttk.Combobox(cab, width=22, state="readonly")
            self.cmb_oper.grid(row=1, column=0, padx=(0, 10))
            ttk.Label(cab, text="AÑO").grid(row=0, column=1)
            self.var_anio = tk.StringVar(value=str(date.today().year)[-2:])
            ttk.Entry(cab, width=5, textvariable=self.var_anio, justify="center").grid(row=1, column=1, padx=3)
            ttk.Label(cab, text="CICLO").grid(row=0, column=2)
            self.var_ciclo = tk.StringVar(value=ciclo_actual())
            ttk.Entry(cab, width=5, textvariable=self.var_ciclo, justify="center").grid(row=1, column=2, padx=3)
            ttk.Label(cab, text="CÓDIGO").grid(row=0, column=3)
            self.var_codigo = tk.StringVar()
            ttk.Entry(cab, width=12, textvariable=self.var_codigo, state="readonly",
                      justify="center").grid(row=1, column=3, padx=(3, 15))
            for v in (self.var_anio, self.var_ciclo):
                v.trace_add("write", lambda *a: self._actualizar_codigo())
            self.cmb_oper.bind("<<ComboboxSelected>>", lambda e: self._actualizar_codigo())

            botones = ttk.Frame(cab)
            botones.grid(row=0, column=4, rowspan=2, padx=5)
            for i, (txt, cmd) in enumerate([("BUSCAR", self.buscar), ("AGREGAR", self.agregar),
                                            ("EDITAR", self.editar), ("ELIMINAR", self.eliminar),
                                            ("LIMPIAR", self.limpiar), ("REFRESCAR", self.refrescar)]):
                ttk.Button(botones, text=txt, width=11, command=cmd).grid(row=i // 3, column=i % 3, padx=2, pady=2)
            self.bind("<F5>", lambda e: self.refrescar())

            # Filtros de supervisor (reemplazan L. CÓD / L. SEC / L. MÁQ / L. PROY)
            filt = ttk.LabelFrame(cab, text="Filtros (supervisor)", padding=5)
            filt.grid(row=0, column=5, rowspan=2, padx=10, sticky="n")
            ttk.Label(filt, text="Código").grid(row=0, column=0, sticky="w")
            self.f_codigo = ttk.Entry(filt, width=10)
            self.f_codigo.grid(row=1, column=0, padx=2)
            ttk.Label(filt, text="Sector").grid(row=0, column=1, sticky="w")
            self.f_sector = ttk.Combobox(filt, width=9, state="readonly")
            self.f_sector.grid(row=1, column=1, padx=2)
            self.f_sector.bind("<<ComboboxSelected>>", lambda e: self._maquinas_de(self.f_sector, self.f_maquina))
            ttk.Label(filt, text="Máquina").grid(row=0, column=2, sticky="w")
            self.f_maquina = ttk.Combobox(filt, width=10)
            self.f_maquina.grid(row=1, column=2, padx=2)
            ttk.Label(filt, text="Proyecto").grid(row=0, column=3, sticky="w")
            self.f_proyecto = ttk.Entry(filt, width=9)
            self.f_proyecto.grid(row=1, column=3, padx=2)
            b_filtrar = ttk.Button(filt, text="Filtrar", command=self.filtrar)
            b_filtrar.grid(row=1, column=4, padx=4)
            for w in (self.f_codigo, self.f_sector, self.f_maquina, self.f_proyecto, b_filtrar):
                self.solo_sup.append(w)
            self.f_codigo.bind("<FocusIn>", lambda e: self._ayuda("Filtra por OT que contenga el texto. 'TT' trae todo."))


        def _armar_lista(self):
            ttk.Label(self, text="VISTA DE ACTIVIDADES", padding=(10, 6, 0, 0)).pack(anchor="w")
            frm = ttk.Frame(self, padding=(10, 0, 10, 0))
            frm.pack(fill="both", expand=True)
            self.tree = ttk.Treeview(frm, show="headings", selectmode="extended", height=10)
            sy = ttk.Scrollbar(frm, orient="vertical", command=self.tree.yview)
            sx = ttk.Scrollbar(frm, orient="horizontal", command=self.tree.xview)
            self.tree.configure(yscrollcommand=sy.set, xscrollcommand=sx.set)
            self.tree.grid(row=0, column=0, sticky="nsew")
            sy.grid(row=0, column=1, sticky="ns")
            sx.grid(row=1, column=0, sticky="ew")
            frm.rowconfigure(0, weight=1)
            frm.columnconfigure(0, weight=1)
            self.tree.tag_configure("extra", foreground="red")
            self.tree.tag_configure("dev", foreground="blue")
            self.tree.bind("<Double-1>", lambda e: self.cargar_seleccion())
            self.lbl_cuenta = ttk.Label(self, text="", padding=(10, 0))
            self.lbl_cuenta.pack(anchor="w")

        def _armar_formulario(self):
            frm = ttk.Frame(self, padding=(10, 4, 10, 8))
            frm.pack(fill="x")

            def entrada(padre, clave, fila, col, ancho=14, sup=False):
                ttk.Label(padre, text=etiquetas.get(clave, clave)).grid(row=fila, column=col, sticky="w", padx=4)
                w = ttk.Entry(padre, width=ancho)
                w.grid(row=fila + 1, column=col, sticky="we", padx=4, pady=(0, 4))
                self._registrar(clave, w, sup)
                return w

            def combo(padre, clave, fila, col, ancho=14, sup=False, editable=False):
                ttk.Label(padre, text=etiquetas.get(clave, clave)).grid(row=fila, column=col, sticky="w", padx=4)
                w = ttk.Combobox(padre, width=ancho, state="normal" if editable else "readonly")
                w.grid(row=fila + 1, column=col, sticky="we", padx=4, pady=(0, 4))
                self._registrar(clave, w, sup)
                return w

            def texto(padre, clave, fila, col, alto, ancho=48, colspan=1):
                ttk.Label(padre, text=etiquetas.get(clave, clave)).grid(row=fila, column=col, columnspan=colspan,
                                                                     sticky="w", padx=4)
                w = tk.Text(padre, height=alto, width=ancho, wrap="word", font=("Segoe UI", 9))
                w.grid(row=fila + 1, column=col, columnspan=colspan, sticky="we", padx=4, pady=(0, 4))
                self._registrar(clave, w, False)
                return w

            etiquetas = {ACTIVIDAD: "ACTIVIDAD A REALIZAR O REALIZADA (Detallar lo máximo posible)",
                         REPUESTOS: "REPUESTOS UTILIZADOS (ESCRIBIR SRH EN SU DEFECTO)",
                         PLANIF: "¿FUE PLANIFICADO?", MAQ_PAR: "¿MAQ PARADA?", TIPO_END: "EXTRAS/DEV",
                         TIPO_PCPG: "TIPO (PCPG)", CAT: "ID PROYECTO", HI: "HORA INICIO", HF: "HORA FIN",
                         SUBCONJ: "SUBCONJUNTO", MAQUINA: "MÁQUINA"}

            # Columna 1 y 2
            c1 = ttk.Frame(frm)
            c1.grid(row=0, column=0, sticky="n")
            entrada(c1, OT_SEC, 0, 0, sup=True)
            entrada(c1, OT_REPLAN, 2, 0, sup=True)
            entrada(c1, QREPLAN, 4, 0, sup=True)
            entrada(c1, FECHA, 6, 0)
            self.cmb_sector = combo(c1, SECTOR, 8, 0)
            self.cmb_maquina = combo(c1, MAQUINA, 0, 1, editable=True)
            combo(c1, CONJ, 2, 1, sup=True, editable=True)
            combo(c1, SUBCONJ, 4, 1, sup=True, editable=True)
            combo(c1, CODACT, 6, 1, sup=True, editable=True)
            combo(c1, PRIORIDAD, 8, 1)
            self.cmb_sector.bind("<<ComboboxSelected>>",
                                 lambda e: self._maquinas_de(self.cmb_sector, self.cmb_maquina, limpiar=True))

            # Columna central: textos
            c2 = ttk.Frame(frm)
            c2.grid(row=0, column=1, sticky="n", padx=6)
            texto(c2, ACTIVIDAD, 0, 0, 4, colspan=2)
            entrada(c2, MEJORAS, 2, 0, ancho=48).grid(columnspan=2)
            entrada(c2, REPUESTOS, 4, 0, ancho=34)
            combo(c2, MAQ_PAR, 4, 1, ancho=12)
            texto(c2, COMENT, 6, 0, 3, ancho=34)
            c2b = ttk.Frame(c2)
            c2b.grid(row=7, column=1, sticky="n")
            combo(c2b, TIPO_END, 0, 0, ancho=12)

            # Columna derecha
            c3 = ttk.Frame(frm)
            c3.grid(row=0, column=2, sticky="n")
            combo(c3, PLANIF, 0, 0, ancho=18)
            combo(c3, ESTADO, 2, 0, ancho=18)
            entrada(c3, CAT, 4, 0, ancho=18, sup=True)
            combo(c3, TIPO_PCPG, 6, 0, ancho=18, sup=True)
            entrada(c3, HI, 8, 0, ancho=18)
            entrada(c3, HF, 10, 0, ancho=18)

            # Ayuda + control de actividad
            c4 = ttk.Frame(frm)
            c4.grid(row=0, column=3, sticky="nw", padx=(15, 0))
            self.lbl_ayuda = ttk.Label(c4, text="Espacio para mensajes de ayuda", wraplength=240,
                                       justify="left", foreground="#333")
            self.lbl_ayuda.grid(row=0, column=0, columnspan=3, sticky="nw", pady=(0, 10))
            c4.rowconfigure(0, minsize=100)
            ttk.Label(c4, text="CONTROL DE ACTIVIDAD").grid(row=1, column=0, columnspan=3, sticky="w")
            self.cmb_control = ttk.Combobox(c4, width=12, state="readonly")
            self.cmb_control.grid(row=2, column=0, columnspan=2, sticky="we")
            b_ind = ttk.Button(c4, text="OK IND", command=lambda: self.control(False))
            b_ind.grid(row=3, column=0, sticky="w", pady=3)
            b_todo = ttk.Button(c4, text="OK TODO", command=lambda: self.control(True))
            b_todo.grid(row=3, column=1, sticky="w", padx=3, pady=3)
            self.solo_sup += [self.cmb_control, b_ind, b_todo]
            ttk.Label(c4, text="Todas las actividades deben ser evaluadas por un supervisor.",
                      wraplength=240, foreground="gray").grid(row=4, column=0, columnspan=3, sticky="w", pady=4)
            self.btn_sup = ttk.Button(c4, text="PASAR A SUP", command=self.cambiar_modo)
            self.btn_sup.grid(row=5, column=0, sticky="w", pady=(8, 0), ipady=6)
            self.lbl_modo = ttk.Label(c4, text="Modo: OPERARIO", foreground="gray")
            self.lbl_modo.grid(row=5, column=1, columnspan=2, sticky="w", pady=(8, 0))

            self.cmb_oper.bind("<FocusIn>", lambda e: self._ayuda("Seleccione su nombre."))

        def _registrar(self, clave, w, sup):
            self.campos[clave] = w
            if sup:
                self.solo_sup.append(w)
            if clave in AYUDA:
                w.bind("<FocusIn>", lambda e, c=clave: self._ayuda(AYUDA[c]), add="+")
            if clave == FECHA:
                w.bind("<FocusOut>", lambda e: self._normalizar(FECHA, fecha_a_db, "Ingrese una fecha válida (ej: 10-06-2025)."))
            if clave in (HI, HF):
                w.bind("<FocusOut>", lambda e, c=clave: self._normalizar(c, hora_a_db, "Ingrese una hora válida (ej: 08:00 o 14:30)."))

        # ---------------- Utilidades de interfaz ----------------
        def _ayuda(self, texto):
            self.lbl_ayuda.config(text=texto)

        def _get(self, clave):
            w = self.campos[clave]
            return w.get("1.0", "end").strip() if isinstance(w, tk.Text) else w.get().strip()

        def _set(self, clave, valor):
            w = self.campos[clave]
            valor = "" if valor is None else str(valor)
            if isinstance(w, tk.Text):
                w.delete("1.0", "end")
                w.insert("1.0", valor)
            elif isinstance(w, ttk.Combobox):
                estado = w.cget("state")
                w.config(state="normal")
                w.set(valor)
                w.config(state=estado)
            else:
                estado = w.cget("state")
                w.config(state="normal")
                w.delete(0, "end")
                w.insert(0, valor)
                w.config(state=estado)

        def _normalizar(self, clave, func, msg):
            t = self._get(clave)
            if not t:
                return
            v = func(t)
            if v is None:
                messagebox.showwarning("Dato inválido", msg)
                self.after(10, self.campos[clave].focus_set)
            else:
                self._set(clave, db_a_texto(v))

        def _actualizar_codigo(self):
            self.var_codigo.set(armar_codigo(self.var_anio.get(), self.var_ciclo.get(), self.cmb_oper.get()))

        def _maquinas_de(self, cmb_sec, cmb_maq, limpiar=False):
            cmb_maq["values"] = self.db.lista("maquinas", cmb_sec.get())
            if limpiar:
                cmb_maq.set("")

        def _cargar_listas(self):
            self.cmb_oper["values"] = self.db.lista("operarios")
            sectores = self.db.lista("sectores")
            self.cmb_sector["values"] = sectores
            self.f_sector["values"] = [""] + sectores
            for clave, lista in [(ESTADO, "estado"), (PRIORIDAD, "prioridad"), (PLANIF, "planificado"),
                                 (TIPO_END, "extras"), (MAQ_PAR, "maq_parada"), (TIPO_PCPG, "tipo_pcpg")]:
                self.campos[clave]["values"] = [""] + self.db.lista(lista)
            self.cmb_control["values"] = [""] + self.db.lista("supervisores")
            # Columnas de la lista: las 31 de bbdd
            cols = self.db.encabezados
            self.tree["columns"] = cols
            anchos = {ACTIVIDAD: 300, COMENT: 200, REPUESTOS: 160, MEJORAS: 160, MAQUINA: 90, OPERARIO: 85,
                      FECHA: 90, HI: 50, HF: 50, TIPO_PCPG: 110, CONTROL: 110}
            for c in cols:
                self.tree.heading(c, text=c)
                self.tree.column(c, width=anchos.get(c, 70), stretch=False, anchor="w")

        def _aplicar_permisos(self):
            for w in self.solo_sup:
                if isinstance(w, ttk.Combobox):
                    editable = w in (self.f_maquina,) or w in (self.campos.get(CONJ), self.campos.get(SUBCONJ),
                                                              self.campos.get(CODACT))
                    w.config(state=("normal" if editable else "readonly") if self.supervisor else "disabled")
                else:
                    w.config(state="normal" if self.supervisor else "disabled")
            self.btn_sup.config(text="VOLVER A OP" if self.supervisor else "PASAR A SUP")
            self.lbl_modo.config(text="Modo: SUPERVISOR" if self.supervisor else "Modo: OPERARIO",
                                 foreground="darkred" if self.supervisor else "gray")

        def _mostrar(self, filas, conservar=False):
            sel = self.tree.selection() if conservar else ()
            vista = self.tree.yview()[0] if conservar else 0
            self.filas_lista = filas
            self.tree.delete(*self.tree.get_children())
            i_tipo = self.db.encabezados.index(TIPO_END) + 1
            for f in filas:
                tipo = str(f[i_tipo] or "").strip().lower()
                tag = ("extra",) if tipo == "extra" else ("dev",) if tipo == "dev" else ()
                self.tree.insert("", "end", iid=str(f[0]), values=[db_a_texto(v) for v in f[1:]], tags=tag)
            if conservar:
                existentes = [i for i in sel if self.tree.exists(i)]
                if existentes:
                    self.tree.selection_set(existentes)
                self.tree.yview_moveto(vista)
            self._actualizar_estado()

        def refrescar(self):
            """Vuelve a ejecutar la última búsqueda/filtro: muestra altas, cambios y bajas de todas las PC."""
            if self.consulta is None:
                return
            try:
                filas = self.consulta()
            except sqlite3.Error as ex:
                self.conexion_ok = False
                self._actualizar_estado()
                messagebox.showerror("Base de datos", f"No se pudo leer la base:\n{ex}")
                return
            self.conexion_ok = True
            self.cambios_pendientes = False
            self.ultimo_refresco = datetime.now()
            self.dv = self.db.data_version()
            self._mostrar(filas, conservar=True)

        def _limpiar_campos(self):
            for clave in self.campos:
                self._set(clave, "")
            self.cmb_maquina["values"] = []
            self.id_cargado = None
            self.rev_cargado = None
            self.cmb_oper.focus_set()

        def _datos_formulario(self):
            d = {k: self._get(k) for k in self.campos}
            d[OPERARIO] = self.cmb_oper.get().strip()
            return d

        def _a_guardar(self, d):
            """Convierte los datos del formulario al formato de la base."""
            out = {}
            for k, v in d.items():
                if k == FECHA:
                    iso = fecha_a_db(v)
                    if iso and iso >= "1900":
                        out[k] = iso
                        out[ANIO] = int(iso[:4])
                    elif iso:  # fechas especiales (11/11/1111 = actividad semanal): se guardan como texto
                        out[k] = f"{iso[8:10]}/{iso[5:7]}/{iso[:4]}"
                    else:
                        out[k] = v or None
                elif k in (HI, HF):
                    out[k] = hora_a_db(v) or (v or None)
                else:
                    out[k] = v or None
            return out

        # ---------------- Acciones ----------------
        def buscar(self):
            if not self.cmb_oper.get():
                messagebox.showwarning("Atención", "Nombre de operario en blanco.")
                return
            codigo = self.var_codigo.get()
            if len(codigo) < 6:
                messagebox.showwarning("Atención", "Faltan datos del usuario (año, ciclo u operario).")
                return
            prefijo = codigo[:6]
            self.consulta = lambda: self.db.por_codigo(prefijo)
            self.refrescar()
            if not self.filas_lista:
                messagebox.showinfo("Buscar", "No se encontraron coincidencias.")

        def agregar(self):
            if self.db.vacia:
                return
            d = self._datos_formulario()
            err = validar(d, self.supervisor)
            if err:
                messagebox.showwarning("Falta dato", err)
                return
            codigo = self.var_codigo.get()
            if len(codigo) < 6:
                messagebox.showwarning("Atención", "No se pudo armar el código (año, ciclo y operario).")
                return
            datos = self._a_guardar(d)
            datos[SEM] = self.var_ciclo.get().strip()
            datos.setdefault(ANIO, 2000 + int(self.var_anio.get()[-2:]))
            try:
                ot = self.db.insertar_nueva(codigo, datos)
            except sqlite3.OperationalError as ex:
                messagebox.showerror("No se pudo guardar",
                                     f"La base está ocupada o no accesible. Intente de nuevo.\n\n{ex}")
                return
            if self.consulta is None:  # sin lista activa: mostrar las actividades del operario en el ciclo
                prefijo = codigo[:6]
                self.consulta = lambda: self.db.por_codigo(prefijo)
            self.refrescar()
            nuevo = self.db.consultar(f"{self.db.col(OT)} = ?", (ot,))
            if nuevo and not self.tree.exists(str(nuevo[0][0])):  # no entra en el filtro actual: se agrega igual
                self._mostrar(self.filas_lista + nuevo, conservar=True)
            if nuevo:
                self.tree.see(str(nuevo[0][0]))
            self._limpiar_campos()
            self._ayuda(f"Actividad {ot} guardada.")

        def cargar_seleccion(self):
            sel = self.tree.selection()
            if sel:
                self._cargar_id(int(sel[0]))

        def _cargar_id(self, id_):
            fila = self.db.obtener(id_)
            if not fila:
                messagebox.showwarning("Atención", "La actividad ya no existe (fue eliminada por otro usuario).")
                self.refrescar()
                return
            enc = self.db.encabezados
            self._limpiar_campos()
            self.id_cargado = fila[0]
            self.rev_cargado = (self.db.obtener_rev(id_) or (0,))[0] or 0
            sector = fila[enc.index(SECTOR) + 1]
            self._set(SECTOR, sector)
            self.cmb_maquina["values"] = self.db.lista("maquinas", sector or "")
            for clave in self.campos:
                if clave in enc and clave != SECTOR:
                    self._set(clave, db_a_texto(fila[enc.index(clave) + 1]))
            self.cmb_control.set(db_a_texto(fila[enc.index(CONTROL) + 1]))
            self._ayuda(f"Editando OT {fila[enc.index(OT) + 1]}. Modifique y presione EDITAR.")

        def editar(self):
            if self.id_cargado is None:
                messagebox.showwarning("Atención", "Haga doble clic en una fila de la lista para cargarla.")
                return
            d = self._datos_formulario()
            err = validar(d, self.supervisor)
            if err:   # se valida ANTES de modificar (el VBA lo hacía después)
                messagebox.showwarning("Falta dato", err)
                return
            datos = self._a_guardar(d)
            if not self.supervisor:  # el operario no modifica campos de supervisor
                for k in (OT_SEC, OT_REPLAN, QREPLAN, CONJ, SUBCONJ, CODACT, CAT, TIPO_PCPG):
                    datos.pop(k, None)
            try:
                res = self.db.actualizar(self.id_cargado, datos, self.rev_cargado)
            except sqlite3.OperationalError as ex:
                messagebox.showerror("No se pudo guardar",
                                     f"La base está ocupada o no accesible. Intente de nuevo.\n\n{ex}")
                return
            if res == "eliminado":
                messagebox.showwarning("Atención", "No se guardó: la actividad fue eliminada por otro usuario.")
                self._limpiar_campos()
                self.refrescar()
                return
            if res == "conflicto":
                info = (self.db.obtener_rev(self.id_cargado) or (None, "?"))[1]
                r = messagebox.askyesnocancel(
                    "Actividad modificada por otro usuario",
                    f"Otro usuario modificó esta actividad mientras usted la editaba ({info}).\n\n"
                    "Sí: guardar SUS cambios igual (reemplaza los del otro usuario).\n"
                    "No: descartar sus cambios y cargar la versión actual.\n"
                    "Cancelar: volver al formulario sin guardar.", icon="warning")
                if r is None:
                    return
                if r is False:
                    self._cargar_id(self.id_cargado)
                    self.refrescar()
                    return
                self.db.actualizar(self.id_cargado, datos)
            self.refrescar()
            self._limpiar_campos()
            self._ayuda("Cambios guardados.")

        def eliminar(self):
            sel = self.tree.selection()
            if not sel:
                messagebox.showwarning("Atención", "Seleccione una fila de la lista para eliminar.")
                return
            if not self.supervisor:
                messagebox.showwarning("Atención", "Debe ser supervisor para eliminar de la base de datos.")
                return
            id_ = int(sel[0])
            ot = self.tree.set(sel[0], OT)
            if not messagebox.askyesno("Confirmar eliminación",
                                       f"¿Está seguro que desea eliminar la OT '{ot}'?\nEsta acción no se puede deshacer.",
                                       icon="warning", default="no"):
                return
            self.db.eliminar(id_)
            self._mostrar([f for f in self.filas_lista if f[0] != id_], conservar=True)
            self._limpiar_campos()

        def limpiar(self):
            self.consulta = None
            self.cambios_pendientes = False
            self._mostrar([])
            self._limpiar_campos()
            self._ayuda("Espacio para mensajes de ayuda")

        def filtrar(self):
            codigo = self.f_codigo.get().strip()
            if not any([codigo, self.f_sector.get(), self.f_maquina.get().strip(), self.f_proyecto.get().strip()]):
                messagebox.showinfo("Filtros", "Complete al menos un filtro ('TT' en Código trae todo).")
                return
            self.config(cursor="watch")
            self.update()
            args = (codigo, self.f_sector.get(), self.f_maquina.get().strip(), self.f_proyecto.get().strip())
            try:
                self.consulta = lambda: self.db.filtrar(*args)
                self.refrescar()
                self.tree.yview_moveto(0)
            finally:
                self.config(cursor="")

        def control(self, todo):
            sup = self.cmb_control.get()
            if not sup:
                messagebox.showwarning("Atención", "Debe seleccionar un supervisor.")
                return
            if todo:
                ids = [f[0] for f in self.filas_lista]
                if not ids:
                    messagebox.showwarning("Atención", "No hay filas en la lista.")
                    return
                if not messagebox.askyesno("Confirmar", f"¿Marcar las {len(ids)} actividades de la lista como "
                                                        f"controladas por {sup}?"):
                    return
            else:
                ids = [int(i) for i in self.tree.selection()]
                if not ids:
                    messagebox.showwarning("Atención", "Debe seleccionar una fila.")
                    return
            self.db.actualizar_campo(ids, CONTROL, sup)
            self.refrescar()
            self._ayuda(f"Control de actividad actualizado en {len(ids)} fila(s).")

        def cambiar_modo(self):
            if self.supervisor:
                self.supervisor = False
            else:
                if not self.db.hay_clave():
                    c1 = simpledialog.askstring("Contraseña de supervisor",
                                                "No hay contraseña definida.\nIngrese una nueva:", show="*", parent=self)
                    if not c1:
                        return
                    c2 = simpledialog.askstring("Contraseña de supervisor", "Repítala:", show="*", parent=self)
                    if c1 != c2:
                        messagebox.showerror("Error", "Las contraseñas no coinciden.")
                        return
                    self.db.fijar_clave(c1)
                else:
                    c = simpledialog.askstring("Acceso restringido", "Ingrese la contraseña de supervisor:",
                                               show="*", parent=self)
                    if c is None:
                        return
                    if not self.db.verificar_clave(c):
                        messagebox.showerror("Error", "Contraseña incorrecta. Permanece en modo OPERARIO.")
                        return
                self.supervisor = True
            self._aplicar_permisos()

        def importar(self):
            if not self.db.vacia:
                if not self.supervisor:
                    messagebox.showwarning("Importar", "Solo un supervisor puede reimportar el Excel.")
                    return
                if not messagebox.askyesno(
                        "Importar", "Esto REEMPLAZA todos los datos y listas de la base por los del Excel.\n\n"
                                    "Asegúrese de que nadie más esté usando el programa en otra PC.\n"
                                    "Antes de importar se crea un respaldo.\n\n¿Continuar?", icon="warning"):
                    return
            ruta = filedialog.askopenfilename(title="Seleccionar CargaActividades",
                                              filetypes=[("Excel", "*.xlsm *.xlsx")])
            if not ruta:
                return
            self.config(cursor="watch")
            self.update()
            try:
                if not self.db.vacia:
                    respaldar(self.db.ruta, self.cfg.get("respaldo", "carpeta"), manual=True)
                n = self.db.importar_excel(ruta)
            except KeyError as ex:
                messagebox.showerror("Error", f"El archivo no tiene la hoja {ex}.")
                return
            except Exception as ex:
                messagebox.showerror("Error al importar", str(ex))
                return
            finally:
                self.config(cursor="")
            self._cargar_listas()
            self.limpiar()
            messagebox.showinfo("Importación", f"Se importaron {n} registros y las listas de '{HOJA_DATOS}'.")

        def exportar(self, completa):
            if self.db.vacia:
                return
            if not completa and not self.filas_lista:
                messagebox.showinfo("Exportar", "La lista está vacía.")
                return
            ruta = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")],
                                                initialfile=f"actividades_{datetime.now():%Y%m%d_%H%M}.xlsx")
            if not ruta:
                return
            try:
                self.db.exportar(ruta, None if completa else self.filas_lista)
            except PermissionError:
                messagebox.showerror("Error", "No se pudo guardar: el archivo está abierto en Excel.")
                return
            messagebox.showinfo("Exportar", f"Exportado a:\n{ruta}")

    App().mainloop()


if __name__ == "__main__":
    iniciar_interfaz()
