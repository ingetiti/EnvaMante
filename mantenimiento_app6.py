"""
SECTOR DE MANTENIMIENTO - Carga de actividades (v3, multiusuario)
Réplica de UserForm1 de CargaActividades.xlsm sobre SQLite.

- La primera vez importa la hoja 'bbdd' y las listas de 'Datos' del Excel (solo lectura).
- La ubicación de la base se define en config.ini (junto al programa). Puede estar en una carpeta de red
  compartida para que varias PC trabajen sobre la misma base.
- Agregar / Editar / Eliminar guardan directo en la base.
- Respaldo automático diario, refresco automático de la lista y aviso si otro usuario modificó
  la actividad que se está editando.

Requiere: pip install openpyxl   (sqlite3 y tkinter vienen con Python)
"""
import os
import re
import sys
import socket
import hashlib
import secrets
import sqlite3
import configparser
import calendar
import bisect
from datetime import datetime, date, time, timedelta

from openpyxl import load_workbook, Workbook

# ======================= Configuración =======================
CARPETA = os.path.dirname(sys.executable if getattr(sys, "frozen", False) else os.path.abspath(__file__))
RUTA_CONFIG = os.path.join(CARPETA, "config.ini")
EQUIPO = os.environ.get("COMPUTERNAME") or socket.gethostname()
LIMITE_REFRESCO_AUTO = 3000   # listas más grandes se refrescan solo con el botón REFRESCAR
CONFIG_DEFECTO = {
    "base": {"ruta": "mantenimiento.db"},
    "respaldo": {"carpeta": "", "dias": "30"},
    "app": {"refresco_seg": "30"},
    "plan": {"tolerancia_dias": "3", "anticipacion_orden_dias": "15"},
}
HOJA_BBDD, HOJA_DATOS = "bbdd", "Datos"
INICIO_CICLOS = date(2026, 1, 3)   # igual que en el VBA
DIAS_CICLO = 8

# Listas que se leen de la hoja Datos: nombre -> (columna, fila_desde, fila_hasta)
LISTAS_DATOS = {
    "operarios": ("A", 2, 30),
    "sectores": ("D", 2, 30),
    "estado": ("B", 27, 36),
    "prioridad": ("B", 38, 41),
    "planificado": ("B", 45, 46),
    "extras": ("B", 57, 59),
    "maq_parada": ("B", 62, 67),
    "tipo_pcpg": ("B", 69, 74),
    "supervisores": ("B", 82, 85),
}

# Encabezados de bbdd usados por el formulario
OT, OT_SEC, OPERARIO, FECHA, SECTOR, MAQUINA = "OT MAN", "OT SEC", "OPERARIO", "FECHA", "SECTOR", "MAQUINA"
PRIORIDAD, ACTIVIDAD, HI, HF, OT_REPLAN, QREPLAN = "PRIORIDAD", "ACTIVIDAD", "HI", "HF", "OT REPLAN", "QREPLAN"
TIPO_END, SEM, ANIO, CONJ, SUBCONJ, CODACT = "TIPO (E/N/D)", "SEM/CIC", "AÑO", "CONJUNTO", "SUB-CONJ", "COD/ACT"
ESTADO, PLANIF, MAQ_PAR, TIPO_PCPG, CAT = "ESTADO", "PLANIF (S/N)", "MAQ_PAR", "TIPO (PV-CO-PD-GE)", "CAT"
REPUESTOS, MEJORAS, COMENT, CONTROL = "REPUESTOS UTILIZADOS", "MEJORAS PROPUESTAS", "COMENTARIOS", "CONTROL_ACTIVIDAD"

AYUDA = {
    OT_SEC: "OT del sector que pidió la actividad.",
    OT_REPLAN: "OT de la actividad replanificada (prototipo aún).",
    QREPLAN: "Cantidad de veces que fue replanificada la actividad.",
    FECHA: "Fecha de realización de la actividad (dd-mm-aaaa). Si figura 11/11/1111, es una actividad semanal.",
    SECTOR: "Sector donde se realizará (o realizó) la actividad.",
    MAQUINA: "Máquina o sector específico dónde se realizará (o realizó) la actividad.",
    CONJ: "Parte MACRO de una máquina o sector (prototipo).",
    SUBCONJ: "Sub parte de una máquina o sector (prototipo), debe estar contenida en un CONJUNTO.",
    CODACT: "Código que tendrá una pieza/parte/sector/etc en específico. Es un código único, sigue en prototipo.",
    PRIORIDAD: "ALTA, MEDIA, BAJA.",
    ACTIVIDAD: "Descripción de la actividad a realizar o realizada. Ser lo más específico que se pueda por favor.",
    MEJORAS: "Posibles mejoras que pueden hacerse en la pieza, máquina, sector y/o proceso que se intervino.",
    REPUESTOS: "Escribir con detalle los repuestos que se utilizaron para facilitar su reposición. "
               "En caso de hacer un pedido de insumos, detallar el número de SRH.",
    COMENT: "Escriba cualquier cosa que enriquezca la descripción de la actividad. "
            "O alguna observación que notó durante el desarrollo de la actividad.",
    PLANIF: "Por defecto, SI. A cualquier actividad que no fue planificada con antelación, se le pone NO "
            "(por ejemplo: cuando nos buscan para reparar algo urgente en máquina).",
    ESTADO: "TER: Terminado.\nEP_FT: En proceso, faltó tiempo (especificar cuánto más falta).\n"
            "EP_FM: En proceso, faltó material (Hacer SRH).\n"
            "EP_PR: En proceso, es proyecto (especificar cuánto tiempo más falta).\n"
            "NOTER: No hecho (especificar por qué).",
    CAT: "ID de proyecto (se guarda en la columna CAT, igual que en el Excel).",
    TIPO_PCPG: "PV: Preventivo.\nCO: Correctivo.\nPD: Predictivo.\nGE: General.",
    MAQ_PAR: "SI_Planif: con parada de máquina planificada.\nSI_PorFallo: con parada de máquina por falla.\n"
             "NO_EnMov: con la máquina en producción.\nNO_SinProd: con la máquina parada sin producción.\n"
             "NO_APLICA: la actividad no se realiza sobre una máquina.",
    TIPO_END: "NOEXTRA, EXTRA, DEV. DEV: devolución de horas por un cambio de jornada.",
    HI: "Hora a la que se inició la actividad (hh:mm).",
    HF: "Hora a la que se terminó la actividad (hh:mm).",
}


# ======================= Utilidades =======================
def ciclo_actual(hoy=None):
    hoy = hoy or date.today()
    return f"{(hoy - INICIO_CICLOS).days // DIAS_CICLO + 1:02d}"


def armar_codigo(anio2, ciclo, operario):
    """AÑO(2) + CICLO + 2 primeras letras del operario (igual que ActualizarCodigo del VBA)."""
    anio2, ciclo, operario = anio2.strip(), ciclo.strip(), operario.strip()
    if len(anio2) >= 2 and ciclo:
        return anio2[-2:] + ciclo + (operario[:2] if len(operario) >= 2 else "")
    return ""


def valor_a_db(v):
    """Excel -> SQLite. Fechas ISO (ordenables), horas hh:mm."""
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d") if v.time() == time(0) else v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, time):
        return v.strftime("%H:%M")
    if isinstance(v, str):
        return v.strip() or None
    return v


def db_a_texto(v):
    """SQLite -> texto en pantalla (fechas dd-mm-aaaa)."""
    if v is None:
        return ""
    if isinstance(v, str):
        m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})( \d{2}:\d{2})?", v)
        if m:
            return f"{m[3]}-{m[2]}-{m[1]}"
    return str(v)


def fecha_a_db(t):
    """'dd-mm-aaaa' o 'dd/mm/aaaa' -> ISO. Devuelve None si no es válida."""
    m = re.fullmatch(r"(\d{1,2})[-/.](\d{1,2})[-/.](\d{2}|\d{4})", t.strip())
    if not m:
        return None
    a = int(m[3])
    if a < 100:
        a += 2000 if a < 30 else 1900
    try:
        return date(a, int(m[2]), int(m[1])).strftime("%Y-%m-%d")
    except ValueError:
        return None


def hora_a_db(t):
    """'8', '8:5', '08:05', '08:05:00' -> 'hh:mm'. None si no es válida."""
    m = re.fullmatch(r"(\d{1,2})(?::(\d{1,2}))?(?::\d{1,2})?", t.strip())
    if not m:
        return None
    h, mi = int(m[1]), int(m[2] or 0)
    return f"{h:02d}:{mi:02d}" if h < 24 and mi < 60 else None


def db_a_excel(v):
    if isinstance(v, str):
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", v):
            return datetime.strptime(v, "%Y-%m-%d")
        if re.fullmatch(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2}", v):
            return datetime.strptime(v, "%Y-%m-%d %H:%M")
        if re.fullmatch(r"\d{2}:\d{2}", v):
            return datetime.strptime(v, "%H:%M").time()
    return v


def validar(d, supervisor):
    """Mismas reglas que btnAgregar del VBA. Devuelve el mensaje de error o None."""
    vacio = lambda k: not str(d.get(k) or "").strip()
    if vacio(OPERARIO):
        return "Nombre de operario en blanco."
    if vacio(FECHA):
        return "Debe ingresar una fecha."
    if vacio(SECTOR) or vacio(MAQUINA):
        return "Debe ingresar una máquina/sector."
    if supervisor:
        return None
    if vacio(HI) or vacio(HF):
        return "Debe ingresar una hora de inicio y fin."
    if vacio(ACTIVIDAD):
        return "Debe ingresar una actividad."
    if vacio(TIPO_END):
        return "Debe especificar si fueron extras, no extras o devolución."
    if vacio(COMENT):
        return "Debe ingresar un comentario o propuesta de mejora."
    if vacio(ESTADO):
        return "Agregar un estado de la actividad."
    if d[ESTADO] == "EP_FM" and vacio(REPUESTOS):
        return "Debe especificar el número de SRH del pedido de insumos que realizó para esta actividad."
    return None


# ======================= config.ini y respaldos =======================
def leer_config():
    cp = configparser.ConfigParser()
    cp.read_dict(CONFIG_DEFECTO)
    if os.path.exists(RUTA_CONFIG):
        cp.read(RUTA_CONFIG, encoding="utf-8")
    else:
        guardar_config(cp)
    return cp


def guardar_config(cp):
    try:
        with open(RUTA_CONFIG, "w", encoding="utf-8") as f:
            f.write("; [base] ruta: archivo de la base. Relativa a la carpeta del programa o absoluta,\n"
                    ";   por ejemplo: \\\\PC-TALLER\\Mantenimiento\\mantenimiento.db\n"
                    "; [respaldo] carpeta: vacío = subcarpeta 'respaldos' junto a la base. dias: copias a conservar.\n"
                    "; [app] refresco_seg: cada cuántos segundos se buscan cambios de otros usuarios (0 = nunca).\n\n")
            cp.write(f)
    except OSError:
        pass


def ruta_abs(p):
    """Rutas relativas se toman desde la carpeta del programa. Acepta rutas de red (\\\\servidor\\carpeta)."""
    p = os.path.expandvars(p.strip())
    return p if os.path.isabs(p) else os.path.join(CARPETA, p)


def respaldar(ruta_db, carpeta="", dias=30, manual=False):
    """Copia consistente de la base (API de respaldo de SQLite, segura aunque otros estén usando la base).
    Automático: una copia por día, la hace la primera PC que abre el programa. Devuelve la ruta o None."""
    if not os.path.exists(ruta_db):
        return None
    carpeta = carpeta.strip() or os.path.join(os.path.dirname(ruta_db), "respaldos")
    os.makedirs(carpeta, exist_ok=True)
    base = os.path.splitext(os.path.basename(ruta_db))[0]
    sufijo = f"{datetime.now():%Y%m%d_%H%M%S}_manual" if manual else f"{date.today():%Y%m%d}"
    destino = os.path.join(carpeta, f"{base}_{sufijo}.db")
    if os.path.exists(destino):
        return None
    tmp = f"{destino}.{EQUIPO}.tmp"
    src, dst = sqlite3.connect(ruta_db, timeout=15), sqlite3.connect(tmp)
    try:
        src.backup(dst)
    finally:
        dst.close()
        src.close()
    os.replace(tmp, destino)
    if not manual and dias > 0:  # rotación: se conservan los últimos 'dias' respaldos automáticos
        patron = re.compile(re.escape(base) + r"_\d{8}\.db")
        viejos = sorted(f for f in os.listdir(carpeta) if patron.fullmatch(f))
        for f in viejos[:-dias]:
            try:
                os.remove(os.path.join(carpeta, f))
            except OSError:
                pass
    return destino


# Listas editables (nombre interno, nombre en pantalla)
NOMBRES_LISTAS = [("operarios", "Operarios"), ("sectores", "Sectores"), ("maquinas", "Máquinas por sector"),
                  ("estado", "Estados"), ("prioridad", "Prioridades"), ("planificado", "¿Fue planificado?"),
                  ("extras", "Extras / Dev"), ("maq_parada", "¿Máq. parada?"), ("tipo_pcpg", "Tipo (PCPG)"),
                  ("supervisores", "Supervisores"), ("tipos_tarea", "Tipos de tarea (planificación)")]
TIPOS_TAREA_DEFECTO = ["MEC", "ELE"]
# Valores que el programa usa en reglas: se avisa antes de renombrarlos o eliminarlos
VALORES_CON_USO = {
    ("estado", "EP_FM"): "la validación que exige el n° de SRH en REPUESTOS",
    ("extras", "EXTRA"): "el color rojo en la lista",
    ("extras", "DEV"): "el color azul en la lista",
}


COLS_BUSQUEDA = [OT, OT_SEC, OPERARIO, SECTOR, MAQUINA, CONJ, SUBCONJ, ACTIVIDAD, ESTADO, REPUESTOS, MEJORAS, COMENT,
                 CONTROL]
_SIN_ACENTOS = str.maketrans("áéíóúüàèìòùâêîôû", "aeiouuaeiouaeiou")


def normalizar(t):
    """Minúsculas y sin acentos (la ñ se conserva), para buscar 'valvula' y encontrar 'VÁLVULA'."""
    return str(t or "").lower().translate(_SIN_ACENTOS)


def marca_modificacion():
    return f"{EQUIPO} {datetime.now():%d-%m-%Y %H:%M}"


# ======================= Planificación: fechas y estados =======================
TAREA_CAMPOS = ["ot_base", "tipo", "sector", "maquina", "conjunto", "subconjunto", "tarea", "asignado", "prioridad",
                "comentarios", "fecha", "periodo", "unidad", "fecha_fin", "creada", "creada_por", "rev", "mod_info"]
UNIDADES = {"dias": "días", "semanas": "semanas", "meses": "meses"}   # interno -> pantalla
ESTADOS_PLAN = ["Pendiente", "Vencida", "En orden", "Cargada", "En proceso", "No hecha", "Cumplida", "Cancelada",
                "Sin programar"]
MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre", "Octubre",
         "Noviembre", "Diciembre"]
PALABRAS_MESES = re.compile(r"mensual|bimestral|trimestral|cuatrimestral|semestral|anual", re.I)


def iso_a_date(s):
    if isinstance(s, date):
        return s
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", str(s or ""))
    if not m or m[1] < "1900":
        return None
    try:
        return date(int(m[1]), int(m[2]), int(m[3]))
    except ValueError:
        return None


def semana_excel(d):
    """Igual que =NUM.DE.SEMANA(fecha;2) de Excel (semanas de lunes a domingo, el 1/1 es semana 1)."""
    return (d.timetuple().tm_yday + date(d.year, 1, 1).weekday() - 1) // 7 + 1


def sumar_periodo(inicio, k, n, unidad):
    """k-ésima ocurrencia desde 'inicio' cada n días/semanas/meses (meses: respeta fin de mes)."""
    if unidad == "meses":
        m = inicio.month - 1 + k * n
        y, m = inicio.year + m // 12, m % 12 + 1
        return date(y, m, min(inicio.day, calendar.monthrange(y, m)[1]))
    return inicio + timedelta(days=k * n * (7 if unidad == "semanas" else 1))


def ocurrencias(inicio, n, unidad, hasta, fin=None):
    """Fechas desde 'inicio' hasta 'hasta' (o 'fin'), más una posterior para asignar bien las actividades."""
    res, k = [], 0
    while k < 50000:
        d = sumar_periodo(inicio, k, n, unidad)
        if fin and d > fin:
            break
        res.append(d)
        if d > hasta:
            break
        k += 1
    return res


def asignar_actividades(acts, occs):
    """Asigna cada actividad a una ocurrencia: la vinculada explícitamente o, si no, la de fecha más cercana."""
    res = {o: [] for o in occs}
    if not occs:
        return res
    for a in acts:
        o = iso_a_date(a["fecha_occ"]) if a["fecha_occ"] else None
        if a["fecha_occ"] is not None and o in res:
            res[o].append(a)
            continue
        if len(occs) == 1:
            res[occs[0]].append(a)
            continue
        f = iso_a_date(a["fecha"])
        if f is None:
            continue
        i = bisect.bisect_left(occs, f)
        cands = [occs[j] for j in (i - 1, i) if 0 <= j < len(occs)]
        res[min(cands, key=lambda o: abs((o - f).days))].append(a)
    return res


def estado_y_realizada(acts, marca, item, fecha_occ, hoy):
    """Estado de una ocurrencia y la fecha en que se realizó (para el retraso).
    Prioridad: marca manual del supervisor > ítem de una orden de preventivo > actividades cargadas > fecha."""
    if marca:
        return {"Hecha": "Cumplida"}.get(marca, marca), None
    if item:
        if item["resultado"] == "Hecho":
            return "Cumplida", iso_a_date(item["fecha_real"])
        if item["resultado"] == "No hecho":
            return "No hecha", None
        return "En orden", None
    if acts:
        ter = sorted(a["fecha"] or "" for a in acts if str(a["estado"] or "").upper() == "TER")
        if ter:   # realizada = primera actividad TER
            return "Cumplida", iso_a_date(ter[0])
        e = str(max(acts, key=lambda a: (a["fecha"] or "", a["id"]))["estado"] or "").upper()
        if e.startswith("EP"):
            return "En proceso", None
        if e == "NOTER":
            return "No hecha", None
        return "Cargada", None
    if fecha_occ and fecha_occ < hoy:
        return "Vencida", None
    return "Pendiente", None


def ocurrencias_tarea(t, acts, marcas, items, hasta, hoy):
    """Todas las ocurrencias de una tarea hasta 'hasta' con su estado. None si es periódica sin programar."""
    if es_periodica(t):
        ini = iso_a_date(t["fecha"])
        if not ini or not t["periodo"]:
            return None
        limite = max([hasta] + [d for d in (iso_a_date(a["fecha"]) for a in acts) if d])
        occs = ocurrencias(ini, int(t["periodo"]), t["unidad"], limite, iso_a_date(t["fecha_fin"]))
    else:
        occs = [iso_a_date(t["fecha"])]
    asig = asignar_actividades(acts, occs)
    res = []
    for o in occs:
        clave = o.isoformat() if o else ""
        marca = marcas.get((t["id"], clave))
        item = items.get((t["id"], clave))
        est, realizada = estado_y_realizada(asig[o], marca[0] if marca else None, item, o, hoy)
        res.append(dict(fecha=o, clave=clave, estado=est, acts=asig[o], marca=marca, item=item, realizada=realizada))
    return res


def es_periodica(t):
    return bool(t.get("unidad"))


def texto_periodicidad(t):
    if not es_periodica(t):
        return "Única"
    if not t.get("periodo"):
        return "Periódica (sin frecuencia)"
    if int(t["periodo"]) == 1:
        return {"dias": "Diaria", "semanas": "Semanal", "meses": "Mensual"}.get(t["unidad"], "Cada 1")
    return f"Cada {t['periodo']} {UNIDADES.get(t['unidad'], t['unidad'])}"


# ======================= Base de datos =======================
class BaseDatos:
    def __init__(self, ruta):
        self.ruta = ruta
        self.con = sqlite3.connect(ruta, timeout=15)  # espera hasta 15 s si otra PC está guardando
        self.con.create_function("norm", 1, normalizar, deterministic=True)
        self.con.executescript("""
            CREATE TABLE IF NOT EXISTS columnas (pos INTEGER PRIMARY KEY, nombre TEXT);
            CREATE TABLE IF NOT EXISTS listas (lista TEXT, orden INTEGER, valor TEXT, padre TEXT);
            CREATE TABLE IF NOT EXISTS config (clave TEXT PRIMARY KEY, valor TEXT);
            CREATE TABLE IF NOT EXISTS tareas (id INTEGER PRIMARY KEY AUTOINCREMENT, ot_base TEXT UNIQUE,
                tipo TEXT, sector TEXT, maquina TEXT, conjunto TEXT, subconjunto TEXT, tarea TEXT, asignado TEXT,
                prioridad TEXT, comentarios TEXT, fecha TEXT, periodo INTEGER, unidad TEXT, fecha_fin TEXT,
                creada TEXT, creada_por TEXT, rev INTEGER DEFAULT 0, mod_info TEXT);
            CREATE TABLE IF NOT EXISTS vinculos (actividad_id INTEGER PRIMARY KEY, tarea_id INTEGER, fecha_occ TEXT);
            CREATE TABLE IF NOT EXISTS marcas (tarea_id INTEGER, fecha_occ TEXT, estado TEXT, info TEXT,
                PRIMARY KEY (tarea_id, fecha_occ));
            CREATE TABLE IF NOT EXISTS ordenes (id INTEGER PRIMARY KEY AUTOINCREMENT, tarea_id INTEGER UNIQUE,
                ot TEXT UNIQUE, sector TEXT, maquina TEXT, fecha TEXT, creada TEXT, creada_por TEXT);
            CREATE TABLE IF NOT EXISTS orden_items (orden_id INTEGER, tarea_id INTEGER, fecha_occ TEXT,
                resultado TEXT, fecha_real TEXT, comentario TEXT, info TEXT,
                PRIMARY KEY (orden_id, tarea_id, fecha_occ));
        """)
        self.encabezados = [r[0] for r in self.con.execute("SELECT nombre FROM columnas ORDER BY pos")]
        self.ultimo_id = None
        self._migrar()

    def _migrar(self):
        """Agrega las columnas de control de cambios a bases creadas con versiones anteriores."""
        if self.vacia:
            return
        existentes = {r[1] for r in self.con.execute("PRAGMA table_info(actividades)")}
        if "rev" not in existentes:
            self.con.execute("ALTER TABLE actividades ADD COLUMN rev INTEGER DEFAULT 0")
        if "mod_info" not in existentes:
            self.con.execute("ALTER TABLE actividades ADD COLUMN mod_info TEXT")
        self.con.execute(f"CREATE INDEX IF NOT EXISTS idx_otsec ON actividades ({self.col(OT_SEC)})")
        if not self.con.execute("SELECT 1 FROM listas WHERE lista = 'tipos_tarea'").fetchone():
            self.con.executemany("INSERT INTO listas VALUES ('tipos_tarea', ?, ?, NULL)",
                                 list(enumerate(TIPOS_TAREA_DEFECTO)))
        self.con.commit()

    def data_version(self):
        """Cambia cuando OTRA conexión (otra PC) guarda algo en la base."""
        return self.con.execute("PRAGMA data_version").fetchone()[0]

    @property
    def _select(self):
        return f"SELECT id, {', '.join(self.cols)} FROM actividades"

    @property
    def vacia(self):
        return not self.encabezados

    def col(self, nombre):
        return f"c{self.encabezados.index(nombre) + 1:02d}"

    @property
    def cols(self):
        return [f"c{i + 1:02d}" for i in range(len(self.encabezados))]

    # ---------- Importación ----------
    def importar_excel(self, ruta_excel, listas=True):
        wb = load_workbook(ruta_excel, read_only=True, data_only=True)  # el Excel nunca se guarda
        try:
            n = self._importar_bbdd(wb[HOJA_BBDD])
            if listas:
                self._importar_listas(wb[HOJA_DATOS])
        finally:
            wb.close()
        self.con.commit()
        return n

    # ---------- Listas (hoja Datos) ----------
    def listas_todas(self):
        """{(lista, padre): [valores en orden]}. Las máquinas usan padre = sector."""
        res = {}
        for lista, valor, padre in self.con.execute("SELECT lista, valor, padre FROM listas ORDER BY lista, padre, orden"):
            vals = res.setdefault((lista, padre), [])
            if valor not in vals:
                vals.append(valor)
        return res

    def guardar_listas(self, listas):
        self.con.execute("BEGIN IMMEDIATE")
        try:
            self.con.execute("DELETE FROM listas")
            self.con.executemany("INSERT INTO listas VALUES (?, ?, ?, ?)",
                                 [(lista, i, v, padre) for (lista, padre), vals in listas.items()
                                  for i, v in enumerate(vals)])
            self.con.commit()
        except Exception:
            self.con.rollback()
            raise

    def _importar_bbdd(self, ws):
        filas = ws.iter_rows(values_only=True)
        enc = [str(h).strip() if h is not None else f"COL{i + 1}" for i, h in enumerate(next(filas))]
        n = len(enc)
        cols = [f"c{i + 1:02d}" for i in range(n)]
        cur = self.con.cursor()
        cur.execute("DROP TABLE IF EXISTS actividades")
        cur.execute("DELETE FROM columnas")
        cur.execute(f"CREATE TABLE actividades (id INTEGER PRIMARY KEY AUTOINCREMENT, {', '.join(cols)}, "
                    f"rev INTEGER DEFAULT 0, mod_info TEXT)")
        cur.executemany("INSERT INTO columnas VALUES (?, ?)", list(enumerate(enc, 1)))
        self.encabezados = enc
        cur.execute(f"CREATE INDEX idx_ot ON actividades ({self.col(OT)})")
        cur.execute(f"CREATE INDEX idx_otsec ON actividades ({self.col(OT_SEC)})")
        datos = []
        for f in filas:
            f = list(f[:n]) + [None] * (n - len(f))
            if any(v is not None for v in f):
                datos.append([valor_a_db(v) for v in f])
        cur.executemany(f"INSERT INTO actividades ({', '.join(cols)}) VALUES ({', '.join('?' * n)})", datos)
        return len(datos)

    def _importar_listas(self, ws):
        celdas = {}
        for r, fila in enumerate(ws.iter_rows(min_row=1, max_row=100, max_col=26, values_only=True), 1):
            for ci, v in enumerate(fila, 1):
                if v is not None and str(v).strip():
                    celdas[(chr(64 + ci), r)] = str(v).strip()
        reg = []
        for nombre, (col, desde, hasta) in LISTAS_DATOS.items():
            vals = [celdas.get((col, r)) for r in range(desde, hasta + 1)]
            reg += [(nombre, i, v, None) for i, v in enumerate(x for x in vals if x)]
        # Máquinas por sector: encabezados en E1:Z1
        for ci in range(5, 27):
            letra = chr(64 + ci)
            sector = celdas.get((letra, 1))
            if sector:
                for r in range(2, 101):
                    v = celdas.get((letra, r))
                    if v:
                        reg.append(("maquinas", r, v, sector))
        reg += [("tipos_tarea", i, v, None) for i, v in enumerate(TIPOS_TAREA_DEFECTO)]
        self.con.execute("DELETE FROM listas")
        self.con.executemany("INSERT INTO listas VALUES (?, ?, ?, ?)", reg)

    def lista(self, nombre, padre=None):
        sql, p = "SELECT valor FROM listas WHERE lista = ?", [nombre]
        if padre is not None:
            sql += " AND padre = ?"
            p.append(padre)
        vistos, res = set(), []
        for (v,) in self.con.execute(sql + " ORDER BY orden", p):
            if v not in vistos:
                vistos.add(v)
                res.append(v)
        return res

    # ---------- Consultas ----------
    def consultar(self, where="", params=(), orden=None):
        orden = orden or f"{self.col(FECHA)}, id"
        sql = self._select + (f" WHERE {where}" if where else "") + f" ORDER BY {orden}"
        return self.con.execute(sql, params).fetchall()

    def por_codigo(self, prefijo):
        return self.consultar(f"CAST({self.col(OT)} AS TEXT) LIKE ?", (prefijo + "%",))

    def filtrar(self, codigo="", sector="", maquina="", proyecto=""):
        w, p = [], []
        if codigo and codigo.upper() != "TT":
            w.append(f"CAST({self.col(OT)} AS TEXT) LIKE ?"); p.append(f"%{codigo}%")
        if sector:
            w.append(f"{self.col(SECTOR)} LIKE ?"); p.append(f"%{sector}%")
        if maquina:
            w.append(f"{self.col(MAQUINA)} LIKE ?"); p.append(f"%{maquina}%")
        if proyecto:
            w.append(f"CAST({self.col(CAT)} AS TEXT) LIKE ?"); p.append(f"%{proyecto}%")
        return self.consultar(" AND ".join(w), p)

    def buscar_texto(self, texto, limite=500):
        """Búsqueda por palabras (todas deben aparecer, en cualquier orden y columna; sin distinguir mayúsculas
        ni acentos). Devuelve (total, filas) con las más recientes primero."""
        palabras = [normalizar(p) for p in texto.split() if p.strip()]
        if not palabras or self.vacia:
            return 0, []
        cols = [self.col(c) for c in COLS_BUSQUEDA if c in self.encabezados]
        junto = "norm(" + " || ' ' || ".join(f"COALESCE(CAST({c} AS TEXT), '')" for c in cols) + ")"
        where = " AND ".join(f"instr({junto}, ?) > 0" for _ in palabras)
        filas = self.con.execute(f"{self._select} WHERE {where} ORDER BY {self.col(FECHA)} DESC, id DESC",
                                 palabras).fetchall()   # una sola pasada: total y filas juntos
        return len(filas), filas[:limite]

    def obtener(self, id_):
        return self.con.execute(self._select + " WHERE id = ?", (id_,)).fetchone()

    def obtener_rev(self, id_):
        """(rev, mod_info) del registro, o None si fue eliminado."""
        return self.con.execute("SELECT rev, mod_info FROM actividades WHERE id = ?", (id_,)).fetchone()

    def nueva_ot(self, codigo):
        """codigo + (mayor correlativo existente + 1), como GenerarNuevaOT del VBA."""
        mayor = 0
        for (ot,) in self.con.execute(f"SELECT {self.col(OT)} FROM actividades WHERE CAST({self.col(OT)} AS TEXT) LIKE ?",
                                      (codigo + "%",)):
            resto = str(ot)[len(codigo):]
            if resto.isdigit():
                mayor = max(mayor, int(resto))
        return f"{codigo}{mayor + 1}"

    # ---------- Escritura ----------
    def insertar(self, valores_dict):
        cols = [self.col(k) for k in valores_dict] + ["mod_info"]
        self.con.execute(f"INSERT INTO actividades ({', '.join(cols)}) VALUES ({', '.join('?' * len(cols))})",
                         list(valores_dict.values()) + [marca_modificacion()])
        self.con.commit()

    def insertar_nueva(self, codigo, valores_dict):
        """Genera la OT y la inserta en una sola transacción bloqueante, así dos PC
        que cargan con el mismo código al mismo tiempo no obtienen la misma OT."""
        self.con.execute("BEGIN IMMEDIATE")
        try:
            ot = self.nueva_ot(codigo)
            datos = dict(valores_dict)
            datos[OT] = ot
            cols = [self.col(k) for k in datos] + ["mod_info"]
            cur = self.con.execute(f"INSERT INTO actividades ({', '.join(cols)}) VALUES ({', '.join('?' * len(cols))})",
                                   list(datos.values()) + [marca_modificacion()])
            self.ultimo_id = cur.lastrowid
            self.con.commit()
            return ot
        except Exception:
            self.con.rollback()
            raise

    def actualizar(self, id_, valores_dict, rev=None):
        """Guarda solo si nadie modificó el registro desde que se cargó (rev).
        Devuelve 'ok', 'conflicto' (otro lo modificó) o 'eliminado'. Con rev=None guarda siempre."""
        sets = ", ".join(f"{self.col(k)} = ?" for k in valores_dict)
        sql = f"UPDATE actividades SET {sets}, rev = COALESCE(rev, 0) + 1, mod_info = ? WHERE id = ?"
        params = list(valores_dict.values()) + [marca_modificacion(), id_]
        if rev is not None:
            sql += " AND COALESCE(rev, 0) = ?"
            params.append(rev)
        cur = self.con.execute(sql, params)
        self.con.commit()
        if cur.rowcount:
            return "ok"
        return "eliminado" if self.obtener_rev(id_) is None else "conflicto"

    def actualizar_campo(self, ids, nombre, valor):
        marca = marca_modificacion()
        self.con.executemany(f"UPDATE actividades SET {self.col(nombre)} = ?, rev = COALESCE(rev, 0) + 1, "
                             f"mod_info = ? WHERE id = ?", [(valor, marca, i) for i in ids])
        self.con.commit()

    def eliminar(self, id_):
        self.con.execute("DELETE FROM actividades WHERE id = ?", (id_,))
        self.con.commit()

    def exportar(self, ruta, filas=None):
        wb = Workbook(write_only=True)
        ws = wb.create_sheet(HOJA_BBDD)
        ws.append(self.encabezados)
        for f in (filas if filas is not None else self.consultar(orden="id")):
            ws.append([db_a_excel(v) for v in f[1:]])
        wb.save(ruta)

    # ---------- Planificación ----------
    def tareas(self, id_=None):
        sql = f"SELECT id, {', '.join(TAREA_CAMPOS)} FROM tareas" + (" WHERE id = ?" if id_ else "") + " ORDER BY id"
        filas = [dict(zip(["id"] + TAREA_CAMPOS, r)) for r in self.con.execute(sql, (id_,) if id_ else ())]
        return (filas[0] if filas else None) if id_ else filas

    def _codigos_usados(self):
        usados = {r[0] for r in self.con.execute("SELECT ot_base FROM tareas") if r[0]}
        if not self.vacia:
            c = self.col(OT_SEC)
            usados |= {str(r[0]) for r in self.con.execute(f"SELECT DISTINCT {c} FROM actividades WHERE {c} IS NOT NULL")}
        return usados

    def _generar_ot_base(self, t):
        """INT-MES-NNNN-TIPO-MÁQ (correlativo global, nunca se repite) o P-N-TIPO-SECTOR para las periódicas."""
        usados = self._codigos_usados()
        limpio = lambda s: re.sub(r"[\s\-]+", "_", str(s or "").strip())
        tipo, maq, sector = limpio(t.get("tipo")), limpio(t.get("maquina")), limpio(t.get("sector"))
        if t.get("unidad"):   # periódica
            patron = re.compile(r"P-(\d+)-.*-" + re.escape(sector) + "$")
            n = max([int(m[1]) for u in usados for m in [patron.match(u)] if m] + [0]) + 1
            while f"P-{n}-{tipo}-{sector}" in usados:
                n += 1
            return f"P-{n}-{tipo}-{sector}"
        r = self.con.execute("SELECT valor FROM config WHERE clave = 'contador_int'").fetchone()
        n = int(r[0]) + 1 if r else 1000
        mes = date.today().month
        while f"INT-{mes}-{n:04d}-{tipo}-{maq}" in usados:
            n += 1
        self.con.execute("INSERT OR REPLACE INTO config VALUES ('contador_int', ?)", (str(n),))
        return f"INT-{mes}-{n:04d}-{tipo}-{maq}"

    def guardar_tarea(self, t, id_=None, rev=None):
        """Alta (genera la OT) o modificación con control de cambios. Devuelve (resultado, id, ot_base)."""
        self.con.execute("BEGIN IMMEDIATE")
        try:
            if id_ is None:
                id_nuevo, ot = self._insertar_tarea(t)
                self.con.commit()
                return "ok", id_nuevo, ot
            campos = [c for c in TAREA_CAMPOS if c in t and c not in ("ot_base", "creada", "creada_por", "rev", "mod_info")]
            sql = (f"UPDATE tareas SET {', '.join(c + ' = ?' for c in campos)}, rev = COALESCE(rev, 0) + 1, "
                   f"mod_info = ? WHERE id = ?")
            params = [t[c] for c in campos] + [marca_modificacion(), id_]
            if rev is not None:
                sql += " AND COALESCE(rev, 0) = ?"
                params.append(rev)
            cur = self.con.execute(sql, params)
            self.con.commit()
            if cur.rowcount:
                return "ok", id_, None
            return ("eliminado" if self.tareas(id_) is None else "conflicto"), id_, None
        except Exception:
            self.con.rollback()
            raise

    def _insertar_tarea(self, t):
        """INSERT sin manejar la transacción (la abre quien llama)."""
        t = dict(t)
        t["ot_base"] = t.get("ot_base") or self._generar_ot_base(t)
        t.setdefault("creada", date.today().isoformat())
        t.setdefault("creada_por", EQUIPO)
        t["mod_info"] = marca_modificacion()
        campos = [c for c in TAREA_CAMPOS if c in t and c != "rev"]
        cur = self.con.execute(f"INSERT INTO tareas ({', '.join(campos)}) VALUES ({', '.join('?' * len(campos))})",
                               [t[c] for c in campos])
        return cur.lastrowid, t["ot_base"]

    def eliminar_tarea(self, id_):
        orden = self.ordenes_por_tarea().get(id_)
        if orden:   # es una orden de preventivo: se borran también sus ítems (los ítems del plan quedan)
            self.con.execute("DELETE FROM orden_items WHERE orden_id = ?", (orden,))
            self.con.execute("DELETE FROM ordenes WHERE id = ?", (orden,))
        self.con.execute("DELETE FROM orden_items WHERE tarea_id = ?", (id_,))
        for tabla, col in (("tareas", "id"), ("vinculos", "tarea_id"), ("marcas", "tarea_id")):
            self.con.execute(f"DELETE FROM {tabla} WHERE {col} = ?", (id_,))
        self.con.commit()

    def importar_items(self, nuevos, actualizar=()):
        """Alta de ítems nuevos y actualización de existentes (solo los campos con dato) en una sola transacción."""
        self.con.execute("BEGIN IMMEDIATE")
        try:
            for t in nuevos:
                self._insertar_tarea(dict(t, creada_por=f"Planilla ({EQUIPO})"))
            for actual, t in actualizar:
                campos = [c for c in ("tipo", "conjunto", "subconjunto", "asignado", "prioridad", "comentarios", "fecha",
                                      "periodo", "unidad") if t.get(c) is not None]
                self.con.execute(f"UPDATE tareas SET {', '.join(c + ' = ?' for c in campos)}, rev = COALESCE(rev, 0) + 1, "
                                 f"mod_info = ? WHERE id = ?", [t[c] for c in campos] + [marca_modificacion(), actual["id"]])
            self.con.commit()
        except Exception:
            self.con.rollback()
            raise

    # ---------- Órdenes de preventivo ----------
    def ordenes_por_tarea(self):
        return {t: o for t, o in self.con.execute("SELECT tarea_id, id FROM ordenes")}

    def avance_ordenes(self):
        """{orden_id: (ítems, hechos, no_hechos)}"""
        sql = ("SELECT orden_id, COUNT(*), SUM(resultado = 'Hecho'), SUM(resultado = 'No hecho') FROM orden_items "
               "GROUP BY orden_id")
        return {o: (n, h or 0, nh or 0) for o, n, h, nh in self.con.execute(sql)}

    def items_orden_mapa(self):
        """{(tarea_id, fecha_occ): ítem de orden}. Si una ocurrencia está en varias órdenes, vale la más reciente."""
        sql = ("SELECT i.tarea_id, i.fecha_occ, i.resultado, i.fecha_real, i.comentario, o.ot, o.id "
               "FROM orden_items i JOIN ordenes o ON o.id = i.orden_id ORDER BY o.id")
        return {(t, f): dict(resultado=r, fecha_real=fr, comentario=c, ot=ot, orden_id=oid)
                for t, f, r, fr, c, ot, oid in self.con.execute(sql)}

    def orden(self, orden_id):
        r = self.con.execute("SELECT id, tarea_id, ot, sector, maquina, fecha, creada, creada_por FROM ordenes "
                             "WHERE id = ?", (orden_id,)).fetchone()
        return dict(zip(["id", "tarea_id", "ot", "sector", "maquina", "fecha", "creada", "creada_por"], r)) if r else None

    def ordenes_de(self, sector, maquina):
        sql = ("SELECT o.id, o.ot, o.fecha, COUNT(i.tarea_id), SUM(i.resultado = 'Hecho'), SUM(i.resultado = 'No hecho') "
               "FROM ordenes o LEFT JOIN orden_items i ON i.orden_id = o.id WHERE o.sector = ? AND o.maquina = ? "
               "GROUP BY o.id ORDER BY o.fecha DESC, o.id DESC")
        return [dict(zip(["id", "ot", "fecha", "items", "hechos", "no_hechos"], r))
                for r in self.con.execute(sql, (sector, maquina))]

    def items_de_orden(self, orden_id):
        sql = ("SELECT i.tarea_id, i.fecha_occ, i.resultado, i.fecha_real, i.comentario, i.info, t.tarea, t.ot_base, "
               "t.periodo, t.unidad, t.conjunto, t.subconjunto FROM orden_items i LEFT JOIN tareas t ON t.id = i.tarea_id "
               "WHERE i.orden_id = ? ORDER BY t.conjunto, t.subconjunto, t.tarea")
        claves = ["tarea_id", "fecha_occ", "resultado", "fecha_real", "comentario", "info", "tarea", "ot_base",
                  "periodo", "unidad", "conjunto", "subconjunto"]
        return [dict(zip(claves, r)) for r in self.con.execute(sql, (orden_id,))]

    def candidatos_orden(self, sector, maquina, hasta, hoy=None):
        """Ítems periódicos de la máquina con una ocurrencia sin hacer hasta 'hasta' (vencidos y por vencer)
        que no estén en otra orden. Por ítem se toma la última ocurrencia vencida o, si no hay, la próxima."""
        hoy = hoy or date.today()
        acts, marcas, items = self.actividades_de_tareas(), self.marcas(), self.items_orden_mapa()
        res = []
        for t in self.tareas():
            if t["sector"] != sector or t["maquina"] != maquina or not es_periodica(t):
                continue
            occs = ocurrencias_tarea(t, acts.get(t["id"], []), marcas, items, hasta, hoy) or []
            # sin hacer y sin orden abierta (si estuvo en una orden y quedó 'No hecho', vuelve a ser candidato)
            libres = [o for o in occs if o["fecha"] and o["fecha"] <= hasta and (
                      (o["item"] is None and o["estado"] in ("Vencida", "Pendiente", "No hecha"))
                      or (o["item"] is not None and o["item"]["resultado"] == "No hecho"))]
            if not libres:
                continue
            vencidas = [o for o in libres if o["fecha"] <= hoy]
            o = vencidas[-1] if vencidas else libres[0]
            res.append(dict(tarea=t, fecha_occ=o["fecha"], estado="Vencida" if o["fecha"] < hoy else "Por vencer"))
        res.sort(key=lambda c: (c["fecha_occ"], c["tarea"]["conjunto"] or "", c["tarea"]["tarea"] or ""))
        return res

    def crear_orden(self, sector, maquina, fecha, asignado, items, comentarios=None):
        """Crea la OT general (como tarea puntual, así se carga y se sigue como cualquier otra) y sus ítems.
        items: [(tarea_id, fecha_occ_iso)]. Devuelve (orden_id, ot)."""
        self.con.execute("BEGIN IMMEDIATE")
        try:
            r = self.con.execute("SELECT valor FROM config WHERE clave = 'contador_pm'").fetchone()
            n = int(r[0]) + 1 if r else 1
            maq = re.sub(r"[\s\-]+", "_", str(maquina or "").strip())
            usados = self._codigos_usados()
            while f"PM-{maq}-{n:04d}" in usados:
                n += 1
            ot = f"PM-{maq}-{n:04d}"
            self.con.execute("INSERT OR REPLACE INTO config VALUES ('contador_pm', ?)", (str(n),))
            tarea_id, _ = self._insertar_tarea(dict(
                ot_base=ot, tipo="PM", sector=sector, maquina=maquina, asignado=asignado, fecha=fecha,
                tarea=f"Mantenimiento preventivo {maquina} ({len(items)} ítems)", comentarios=comentarios,
                periodo=None, unidad=None))
            cur = self.con.execute("INSERT INTO ordenes (tarea_id, ot, sector, maquina, fecha, creada, creada_por) "
                                   "VALUES (?, ?, ?, ?, ?, ?, ?)",
                                   (tarea_id, ot, sector, maquina, fecha, date.today().isoformat(), EQUIPO))
            orden_id = cur.lastrowid
            self.con.executemany("INSERT INTO orden_items (orden_id, tarea_id, fecha_occ, resultado) VALUES (?, ?, ?, '')",
                                 [(orden_id, t, f) for t, f in items])
            self.con.commit()
            return orden_id, ot
        except Exception:
            self.con.rollback()
            raise

    def marcar_items(self, orden_id, claves, resultado, fecha_real=None, comentario=None):
        """claves: [(tarea_id, fecha_occ)]. resultado: 'Hecho', 'No hecho' o '' (pendiente)."""
        info = marca_modificacion()
        for tarea_id, fecha_occ in claves:
            sets, params = ["resultado = ?", "fecha_real = ?", "info = ?"], [resultado,
                                                                            fecha_real if resultado == "Hecho" else None, info]
            if comentario is not None:
                sets.append("comentario = ?")
                params.append(comentario)
            self.con.execute(f"UPDATE orden_items SET {', '.join(sets)} WHERE orden_id = ? AND tarea_id = ? AND fecha_occ = ?",
                             params + [orden_id, tarea_id, fecha_occ])
        self.con.commit()

    def quitar_items(self, orden_id, claves):
        self.con.executemany("DELETE FROM orden_items WHERE orden_id = ? AND tarea_id = ? AND fecha_occ = ?",
                             [(orden_id, t, f) for t, f in claves])
        self.con.commit()

    def vincular(self, actividad_id, tarea_id, fecha_occ):
        self.con.execute("INSERT OR REPLACE INTO vinculos VALUES (?, ?, ?)", (actividad_id, tarea_id, fecha_occ or ""))
        self.con.commit()

    def marcar(self, tarea_id, fecha_occ, estado):
        if estado:
            self.con.execute("INSERT OR REPLACE INTO marcas VALUES (?, ?, ?, ?)",
                             (tarea_id, fecha_occ or "", estado, marca_modificacion()))
        else:
            self.con.execute("DELETE FROM marcas WHERE tarea_id = ? AND fecha_occ = ?", (tarea_id, fecha_occ or ""))
        self.con.commit()

    def marcas(self):
        return {(t, f): (e, i) for t, f, e, i in self.con.execute("SELECT tarea_id, fecha_occ, estado, info FROM marcas")}

    def actividades_de_tareas(self):
        """{tarea_id: [actividades cuyo OT SEC es la OT Base de la tarea]} — así se cruza también el histórico."""
        if self.vacia:
            return {}
        c = self.col
        sql = (f"SELECT t.id, a.id, a.{c(OT)}, a.{c(FECHA)}, a.{c(OPERARIO)}, a.{c(ESTADO)}, a.{c(HI)}, a.{c(HF)}, "
               f"a.{c(ACTIVIDAD)}, a.{c(COMENT)}, a.{c(REPUESTOS)}, v.fecha_occ "
               f"FROM actividades a JOIN tareas t ON a.{c(OT_SEC)} = t.ot_base "
               f"LEFT JOIN vinculos v ON v.actividad_id = a.id AND v.tarea_id = t.id")
        claves = ["id", "ot", "fecha", "operario", "estado", "hi", "hf", "actividad", "comentarios", "repuestos", "fecha_occ"]
        res = {}
        for r in self.con.execute(sql):
            res.setdefault(r[0], []).append(dict(zip(claves, r[1:])))
        return res

    def importar_calendario(self, ruta):
        """Importa las hojas 'Calendario' (periódicas) y 'ActividadesAgenda' (puntuales) de CalendarioMTO.xlsm.
        Solo agrega las OT Base que todavía no existen. Devuelve (nuevas, ya_existentes, a_revisar)."""
        wb = load_workbook(ruta, read_only=True, data_only=True)
        existentes = {r[0] for r in self.con.execute("SELECT ot_base FROM tareas")}
        nuevas, repetidas, revisar = [], 0, []
        txt = lambda v: str(v).strip() if v is not None and str(v).strip() else None
        try:
            if "Calendario" in wb.sheetnames:
                for r in wb["Calendario"].iter_rows(min_row=5, max_col=12, values_only=True):
                    r = list(r) + [None] * (12 - len(r))
                    if not txt(r[7]):
                        continue
                    fecha = valor_a_db(r[9]) if isinstance(r[9], (datetime, date)) else None
                    periodo = int(r[10]) if isinstance(r[10], (int, float)) and r[10] > 0 else None
                    unidad = "dias"
                    if not fecha and periodo and periodo <= 12 and PALABRAS_MESES.search(str(r[7])):
                        unidad = "meses"   # p. ej. preventivos de COMP: 1=Mensual, 2=Bimestral, 4=Cuatrimestral
                    t = dict(ot_base=txt(r[1]), tipo=txt(r[2]), sector=txt(r[3]), maquina=txt(r[4]),
                             conjunto=txt(r[5]), subconjunto=txt(r[6]), tarea=txt(r[7]), comentarios=txt(r[11]),
                             fecha=fecha, periodo=periodo, unidad=unidad, creada_por="Importado CalendarioMTO")
                    if not periodo or not fecha:
                        revisar.append(t["ot_base"] or t["tarea"])
                    nuevas.append((t, None))
            if "ActividadesAgenda" in wb.sheetnames:
                for r in wb["ActividadesAgenda"].iter_rows(min_row=4, max_col=15, values_only=True):
                    r = list(r) + [None] * (15 - len(r))
                    if not txt(r[0]) and not txt(r[6]):
                        continue
                    t = dict(ot_base=txt(r[0]), tipo=txt(r[1]), sector=txt(r[2]), maquina=txt(r[3]),
                             conjunto=txt(r[4]), subconjunto=txt(r[5]), tarea=txt(r[6]), asignado=txt(r[10]),
                             comentarios=txt(r[11]),
                             fecha=valor_a_db(r[9]) if isinstance(r[9], (datetime, date)) else None,
                             periodo=None, unidad=None,
                             creada=valor_a_db(r[14]) if isinstance(r[14], (datetime, date)) else None,
                             creada_por="Importado CalendarioMTO")
                    nuevas.append((t, "Hecha" if str(r[7] or "").strip().lower() == "hecho" else None))
        finally:
            wb.close()
        agregadas = 0
        for t, marca in nuevas:
            if t["ot_base"] and t["ot_base"] in existentes:
                repetidas += 1
                continue
            t = {k: v for k, v in t.items() if v is not None or k in ("periodo", "unidad")}
            _, id_, ot = self.guardar_tarea(t)
            existentes.add(ot)
            agregadas += 1
            if marca:
                self.marcar(id_, t.get("fecha") or "", marca)   # solo puntuales traen marca "Hecho"
        return agregadas, repetidas, revisar

    # ---------- Contraseña supervisor ----------
    def hay_clave(self):
        return self.con.execute("SELECT 1 FROM config WHERE clave = 'sup_hash'").fetchone() is not None

    def fijar_clave(self, clave):
        sal = secrets.token_hex(16)
        h = hashlib.sha256((sal + clave).encode()).hexdigest()
        self.con.execute("INSERT OR REPLACE INTO config VALUES ('sup_hash', ?)", (f"{sal}${h}",))
        self.con.commit()

    def verificar_clave(self, clave):
        r = self.con.execute("SELECT valor FROM config WHERE clave = 'sup_hash'").fetchone()
        if not r:
            return False
        sal, h = r[0].split("$")
        return secrets.compare_digest(hashlib.sha256((sal + clave).encode()).hexdigest(), h)


COLORES_ESTADO = {"Vencida": "#c00000", "Cumplida": "#007a33", "Cumplida tarde": "#8a6d00", "En proceso": "#b36b00",
                  "Cargada": "#b36b00", "En orden": "#0060b0", "No hecha": "#7a0050", "Cancelada": "gray",
                  "Sin programar": "#4060a0"}


def texto_retraso(f):
    if f["retraso"] is None:
        return ""
    if f["sin_hacer"]:
        return f"{f['retraso']} d sin hacer"
    return f"{f['retraso']:+d} d" if f["retraso"] else "0 d"


def resumen_plan(filas, tol):
    """Texto de resumen: cantidades por estado, cumplimiento, % a tiempo y retraso promedio."""
    hoy = date.today()
    cuenta = {e: sum(1 for f in filas if f["estado"] == e) for e in ESTADOS_PLAN}
    exigibles = [f for f in filas if f["fecha"] and f["fecha"] <= hoy and f["estado"] not in ("Cancelada", "Sin programar")]
    cumplidas = sum(1 for f in exigibles if f["estado"] == "Cumplida")
    pct = f"{100 * cumplidas / len(exigibles):.0f}%" if exigibles else "-"
    hechas = [f for f in filas if f["realizada"] and f["retraso"] is not None]
    a_tiempo = sum(1 for f in hechas if f["retraso"] <= tol)
    txt = (f"{len(filas)} tareas   |   " + "   ".join(f"{e}: {n}" for e, n in cuenta.items() if n) +
           f"   |   Cumplimiento a la fecha: {pct} ({cumplidas} de {len(exigibles)})")
    if hechas:
        prom = sum(max(f["retraso"], 0) for f in hechas) / len(hechas)
        txt += (f"   |   A tiempo (≤{tol} d): {100 * a_tiempo / len(hechas):.0f}% ({a_tiempo} de {len(hechas)})"
                f"   |   Retraso promedio: {prom:.1f} d")
    return txt


def calcular_plan(db, desde, hasta, atrasadas=True, hoy=None):
    """Filas de planificación entre 'desde' y 'hasta': una por ocurrencia de cada tarea, con su estado
    calculado a partir de las actividades cargadas (OT SEC = OT Base) y de las órdenes de preventivo."""
    hoy = hoy or date.today()
    acts, marcas, items = db.actividades_de_tareas(), db.marcas(), db.items_orden_mapa()
    ordenes, avance = db.ordenes_por_tarea(), db.avance_ordenes()
    filas = []
    for t in db.tareas():
        occs = ocurrencias_tarea(t, acts.get(t["id"], []), marcas, items, hasta, hoy)
        if occs is None:
            filas.append(_fila_plan(t, dict(fecha=None, clave="", estado="Sin programar", acts=acts.get(t["id"], []),
                                             marca=None, item=None, realizada=None), hoy, ordenes, avance))
            continue
        # En las periódicas solo se arrastra la última ocurrencia anterior al período (las previas ya fueron reemplazadas)
        previas = [o["fecha"] for o in occs if o["fecha"] and o["fecha"] < desde]
        ultima_previa = previas[-1] if (previas and es_periodica(t)) else None
        for o in occs:
            f, est = o["fecha"], o["estado"]
            if f is None:      # puntual sin fecha: visible mientras no esté resuelta
                ver = est not in ("Cumplida", "Cancelada", "No hecha")
            elif desde <= f <= hasta:
                ver = True
            else:              # de períodos anteriores: solo si quedó pendiente
                ver = (atrasadas and f < desde and est in ("Vencida", "En proceso", "Cargada", "En orden")
                       and (not es_periodica(t) or f == ultima_previa))
            if ver:
                filas.append(_fila_plan(t, o, hoy, ordenes, avance))
    filas.sort(key=lambda f: (f["fecha"] or date.max, f["sector"] or "", f["ot_base"] or ""))
    return filas


def _fila_plan(t, o, hoy, ordenes, avance=None):
    fecha, acts, marca, item = o["fecha"], o["acts"], o["marca"], o["item"]
    ordenadas = sorted(acts, key=lambda a: (a["fecha"] or "", a["id"]))
    detalle = ""
    if marca:
        detalle = f"Marcada '{marca[0]}' ({marca[1]})"
    elif item:
        detalle = f"Orden {item['ot']}" + (f": {item['resultado']}" if item["resultado"] else " (pendiente)") + \
                  (f" / {item['comentario']}" if item["comentario"] else "")
    elif ordenadas:
        detalle = " / ".join(x for x in (ordenadas[-1]["estado"], ordenadas[-1]["comentarios"]) if x)
    realizada = o["realizada"]
    retraso, sin_hacer = None, False
    if fecha and realizada:
        retraso = (realizada - fecha).days
    elif fecha and o["estado"] in ("Vencida", "En orden", "En proceso", "Cargada", "No hecha") and fecha < hoy:
        retraso, sin_hacer = (hoy - fecha).days, True
    orden_id = ordenes.get(t["id"])
    if orden_id and avance and orden_id in avance:   # OT general de preventivo: avance de su lista de ítems
        n, h, nh = avance[orden_id]
        pend = n - h - nh
        detalle = (f"Ítems: {h} hechos, {nh} no hechos, {pend} pendientes" + (f" | {detalle}" if detalle else ""))
    ot_man = ", ".join(str(a["ot"]) for a in ordenadas)
    if item and not ot_man:
        ot_man = item["ot"]
    return dict(id=t["id"], fecha=fecha, fecha_occ=o["clave"], ot_base=t["ot_base"], tipo=t["tipo"],
                sector=t["sector"], maquina=t["maquina"], tarea=t["tarea"], asignado=t["asignado"],
                periodicidad=texto_periodicidad(t), estado=o["estado"], ot_man=ot_man, detalle=detalle,
                realizada=realizada, retraso=retraso, sin_hacer=sin_hacer, orden_id=orden_id,
                tarea_dict=t)


# ======================= Carga masiva: planilla modelo =======================
PLANILLA_HOJA = "Ítems"
UNIDADES_TEXTO = {"días": "dias", "dias": "dias", "día": "dias", "dia": "dias", "semanas": "semanas",
                  "semana": "semanas", "meses": "meses", "mes": "meses"}


def generar_planilla(db, ruta):
    """Planilla modelo para cargar ítems del plan preventivo, con las listas actuales de la base."""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.comments import Comment
    sectores = db.lista("sectores")
    maquinas = {s: db.lista("maquinas", s) for s in sectores}
    operarios, prioridades, tipos = db.lista("operarios"), db.lista("prioridad"), db.lista("tipos_tarea")
    UNIDADES_P = ["días", "semanas", "meses"]
    FILAS = 1000

    F = lambda **k: Font(**{"name": "Arial", "size": 10, **k})
    borde = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
    wb = Workbook()

    # ---------- Hoja Ítems ----------
    ws = wb.active
    ws.title = "Ítems"
    cols = [("Sector *", 10), ("Máquina *", 14), ("Ítem / Tarea *", 50), ("Cada *", 8), ("Unidad *", 10),
            ("Fecha de inicio", 14), ("Tipo", 8), ("Conjunto", 16), ("Subconjunto", 16), ("Asignado", 14),
            ("Prioridad", 10), ("Comentarios", 40)]
    for i, (t, w) in enumerate(cols, 1):
        c = ws.cell(1, i, t)
        c.font = F(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = borde
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[1].height = 30
    ejemplo = ["IMP", "IMP3", "Lubricar rodamientos del cilindro central", 3, "meses", datetime(2026, 11, 2), "MEC",
               "Cilindro central", "Rodamientos", operarios[0] if operarios else None, "MEDIA",
               "EJEMPLO (esta fila no se importa)"]
    for i, v in enumerate(ejemplo, 1):
        c = ws.cell(2, i, v)
        c.font = F(italic=True, color="7F7F7F")
        c.fill = PatternFill("solid", fgColor="F2F2F2")
        c.border = borde
    for r in range(3, FILAS + 2):
        for i in range(1, len(cols) + 1):
            c = ws.cell(r, i)
            c.font = F()
            c.fill = PatternFill("solid", fgColor="FFF2CC")   # celdas a completar
            c.border = borde
            if i in (3, 12):
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r, 6).number_format = "dd-mm-yyyy"
    ws.cell(2, 6).number_format = "dd-mm-yyyy"
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A1:L{FILAS + 1}"
    ws["A1"].comment = Comment("Sector de la lista. Es obligatorio porque hay máquinas con el mismo nombre "
                               "en distintos sectores (ej. GENERAL).", "Mantenimiento")
    ws["B1"].comment = Comment("La lista muestra las máquinas del sector elegido. Si la máquina no está, "
                               "se puede escribir (Excel avisa, pero permite continuar).", "Mantenimiento")
    ws["F1"].comment = Comment("Primera fecha programada (dd-mm-aaaa). Si se deja vacía, el ítem queda "
                               "'Sin programar' hasta cargarla en el programa.", "Mantenimiento")

    # ---------- Hoja Listas (oculta) ----------
    wl = wb.create_sheet("Listas")
    listas = [("Sectores", sectores), ("Unidades", UNIDADES_P), ("Operarios", operarios), ("Prioridades", prioridades),
              ("Tipos", tipos)]
    col = 1
    rangos = {}
    for titulo, vals in listas:
        wl.cell(1, col, titulo).font = F(bold=True)
        for r, v in enumerate(vals, 2):
            wl.cell(r, col, v).font = F()
        L = wl.cell(1, col).column_letter
        rangos[titulo] = f"Listas!${L}$2:${L}${max(len(vals), 1) + 1}"
        col += 1
    col += 1
    for s in sectores:   # un rango con nombre por sector: M_<sector>, para la lista dependiente de máquinas
        wl.cell(1, col, s).font = F(bold=True)
        vals = maquinas[s] or [""]
        for r, v in enumerate(vals, 2):
            wl.cell(r, col, v).font = F()
        L = wl.cell(1, col).column_letter
        nombre = "M_" + re.sub(r"\W", "_", s)
        wb.defined_names[nombre] = DefinedName(nombre, attr_text=f"Listas!${L}$2:${L}${len(vals) + 1}")
        col += 1
    wl.sheet_state = "hidden"

    # ---------- Validaciones ----------
    rng = lambda L: f"{L}3:{L}{FILAS + 1}"
    def lista(formula, L, estricta=True, titulo="", msg=""):
        dv = DataValidation(type="list", formula1=formula, allow_blank=True,
                            errorStyle="stop" if estricta else "warning", showErrorMessage=True,
                            errorTitle=titulo, error=msg)
        ws.add_data_validation(dv)
        dv.add(rng(L))
    lista(f"{rangos['Sectores']}", "A", True, "Sector", "Elija un sector de la lista.")
    lista('INDIRECT("M_"&SUBSTITUTE($A3," ","_"))', "B", False, "Máquina",
          "La máquina no está en la lista del sector. ¿Continuar igual?")
    lista(f"{rangos['Unidades']}", "E", True, "Unidad", "Elija días, semanas o meses.")
    lista(f"{rangos['Tipos']}", "G", False, "Tipo", "El tipo no está en la lista. ¿Continuar igual?")
    lista(f"{rangos['Operarios']}", "J", True, "Asignado", "Elija un operario de la lista.")
    lista(f"{rangos['Prioridades']}", "K", True, "Prioridad", "Elija una prioridad de la lista.")
    dv = DataValidation(type="whole", operator="greaterThan", formula1="0", allow_blank=True, showErrorMessage=True,
                        errorTitle="Cada", error="Ingrese un número entero mayor a 0.")
    ws.add_data_validation(dv); dv.add(rng("D"))
    dv = DataValidation(type="date", operator="greaterThan", formula1="43831", allow_blank=True, showErrorMessage=True,
                        errorTitle="Fecha", error="Ingrese una fecha válida (dd-mm-aaaa).")
    ws.add_data_validation(dv); dv.add(rng("F"))

    # ---------- Hoja Instrucciones ----------
    wi = wb.create_sheet("Instrucciones", 0)
    wi.column_dimensions["A"].width = 110
    texto = [
        ("PLANILLA DE CARGA MASIVA - PLAN DE MANTENIMIENTO PREVENTIVO", True),
        ("", False),
        ("Cómo completarla", True),
        ("• En la hoja 'Ítems', una fila por cada ítem del plan (ej.: cada tarea del preventivo de IMP3).", False),
        ("• Completar solo las celdas amarillas. La fila 2 (gris) es un ejemplo y no se importa.", False),
        ("• Columnas con * son obligatorias: Sector, Máquina, Ítem / Tarea, Cada y Unidad.", False),
        ("• 'Cada' + 'Unidad' define la periodicidad: 3 + meses = cada 3 meses; 2 + semanas = cada 2 semanas.", False),
        ("• 'Fecha de inicio' es la primera fecha programada. Las siguientes se calculan desde ella (calendario fijo).", False),
        ("  Si queda vacía, el ítem se importa como 'Sin programar' y se completa después en el programa.", False),
        ("• Sector, Unidad, Asignado y Prioridad se eligen de las listas desplegables.", False),
        ("• La lista de Máquina depende del Sector elegido. Si una máquina no figura, se puede escribir igual.", False),
        ("", False),
        ("Al importar en el programa", True),
        ("• Antes de cargar, se muestra un resumen: ítems nuevos, ítems con errores (fila y motivo) y ya existentes.", False),
        ("• Un ítem 'ya existe' si coincide Sector + Máquina + Ítem / Tarea. Se puede elegir actualizarlo o saltearlo,", False),
        ("  así la misma planilla sirve para corregir periodicidades o fechas.", False),
        ("", False),
        (f"Listas generadas desde la base del programa el {datetime.now():%d-%m-%Y}. "
         "Si se agregan sectores, máquinas u operarios, generar una planilla nueva.", False),
    ]
    for r, (t, negrita) in enumerate(texto, 1):
        c = wi.cell(r, 1, t)
        c.font = F(bold=negrita, size=12 if r == 1 else 10, color="1F4E78" if negrita else "000000")
    wb.active = 1
    wb.save(ruta)


def _norm(v):
    return re.sub(r"\s+", " ", str(v or "").strip()).upper()


def leer_planilla_items(db, ruta):
    """Lee y valida la hoja 'Ítems' de la planilla modelo.
    Devuelve dict(nuevos=[tarea], existentes=[(tarea_actual, tarea_nueva)], errores=[(fila, motivo)])."""
    wb = load_workbook(ruta, read_only=True, data_only=True)
    try:
        if PLANILLA_HOJA not in wb.sheetnames:
            raise ValueError(f"El archivo no tiene la hoja '{PLANILLA_HOJA}'. Use la planilla modelo del programa.")
        sectores = {_norm(x): x for x in db.lista("sectores")}
        operarios = {_norm(x): x for x in db.lista("operarios")}
        prioridades = {_norm(x): x for x in db.lista("prioridad")}
        actuales = {(_norm(t["sector"]), _norm(t["maquina"]), _norm(t["tarea"])): t for t in db.tareas() if es_periodica(t)}
        res = dict(nuevos=[], existentes=[], errores=[])
        vistos = set()
        txt = lambda v: str(v).strip() if v is not None and str(v).strip() else None
        for nro, fila in enumerate(wb[PLANILLA_HOJA].iter_rows(min_row=2, max_col=12, values_only=True), 2):
            fila = list(fila) + [None] * (12 - len(fila))
            if all(txt(v) is None for v in fila):
                continue
            sector, maq, tarea, cada, unidad, fecha, tipo, conj, subconj, asig, prio, com = fila
            if (txt(com) or "").upper().startswith("EJEMPLO"):
                continue
            errs = []
            if _norm(sector) not in sectores:
                errs.append(f"sector '{sector or ''}' no existe")
            if not txt(maq):
                errs.append("falta la máquina")
            if not txt(tarea):
                errs.append("falta el ítem / tarea")
            try:
                periodo = int(float(str(cada).replace(",", ".")))
                assert periodo > 0
            except (ValueError, TypeError, AssertionError):
                periodo = None
                errs.append(f"'Cada' inválido ({cada if cada is not None else 'vacío'})")
            u = UNIDADES_TEXTO.get(str(unidad or "").strip().lower())
            if not u:
                errs.append(f"unidad inválida ({unidad or 'vacía'})")
            f_iso = None
            if isinstance(fecha, (datetime, date)):
                f_iso = valor_a_db(fecha)
            elif txt(fecha):
                f_iso = fecha_a_db(txt(fecha))
                if not f_iso:
                    errs.append(f"fecha inválida ({fecha})")
            if txt(asig) and _norm(asig) not in operarios:
                errs.append(f"operario '{asig}' no existe")
            if txt(prio) and _norm(prio) not in prioridades:
                errs.append(f"prioridad '{prio}' no existe")
            if errs:
                res["errores"].append((nro, "; ".join(errs)))
                continue
            clave = (_norm(sector), _norm(maq), _norm(tarea))
            if clave in vistos:
                res["errores"].append((nro, "ítem repetido dentro de la planilla"))
                continue
            vistos.add(clave)
            t = dict(sector=sectores[_norm(sector)], maquina=txt(maq), tarea=txt(tarea), periodo=periodo, unidad=u,
                     fecha=f_iso, tipo=txt(tipo), conjunto=txt(conj), subconjunto=txt(subconj),
                     asignado=operarios.get(_norm(asig)) if txt(asig) else None,
                     prioridad=prioridades.get(_norm(prio)) if txt(prio) else None, comentarios=txt(com))
            if clave in actuales:
                res["existentes"].append((actuales[clave], t))
            else:
                res["nuevos"].append(t)
        return res
    finally:
        wb.close()


# ======================= Interfaz =======================
def iniciar_interfaz():
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox, simpledialog

    TECLAS_IGNORADAS = {"Tab", "ISO_Left_Tab", "Return", "KP_Enter", "Escape", "Up", "Down", "Left", "Right",
                        "Shift_L", "Shift_R", "Control_L", "Control_R", "Alt_L", "Alt_R", "Home", "End",
                        "Caps_Lock", "Prior", "Next"}

    def autocompletar(cmb, estricto=True, avisar=None):
        """Desplegable con autocompletado: mientras se escribe completa en línea y filtra la lista
        (busca en cualquier parte del texto). Si es estricto, al salir solo acepta valores de la lista."""
        cmb._auto, cmb._todos, cmb._antes = True, None, cmb.get()
        if str(cmb.cget("state")) == "readonly":
            cmb.config(state="normal")

        def todos():
            if cmb._todos is None:
                cmb._todos = [str(v) for v in cmb.tk.splitlist(cmb.cget("values"))]
            return cmb._todos

        def al_entrar(e):
            cmb._todos = None
            cmb._antes = cmb.get()

        def tecla(e):
            if e.keysym in TECLAS_IGNORADAS or str(cmb.cget("state")) == "disabled":
                return
            vals = todos()
            borrar = e.keysym in ("BackSpace", "Delete")
            escrito = cmb.get() if borrar else cmb.get()[:cmb.index("insert")]
            bajo = escrito.lower()
            cmb["values"] = [v for v in vals if bajo in v.lower()] if bajo else vals
            if borrar or not escrito:
                return
            empieza = [v for v in vals if v.lower().startswith(bajo)]
            if empieza:   # completa en línea: lo agregado queda seleccionado y se reemplaza al seguir escribiendo
                cmb.delete(0, "end")
                cmb.insert(0, empieza[0])
                cmb.select_range(len(escrito), "end")
                cmb.icursor(len(escrito))

        def confirmar(e=None):
            if str(cmb.cget("state")) == "disabled":
                return
            vals = todos()
            txt = cmb.get().strip()
            cmb["values"] = vals
            final = txt
            if txt:
                exacto = [v for v in vals if v.lower() == txt.lower()]
                if exacto:
                    final = exacto[0]
                elif estricto:
                    contiene = [v for v in vals if txt.lower() in v.lower()]
                    final = contiene[0] if len(contiene) == 1 else ""
                    if not final:
                        (avisar or (lambda m: cmb.bell()))(f"'{txt}' no está en la lista: elija una opción válida.")
            if final != cmb.get():
                cmb.set(final)
            cmb.selection_clear()
            cmb.icursor("end")
            cmb._todos = None
            if final != cmb._antes:
                cmb._antes = final
                cmb.event_generate("<<ComboboxSelected>>")

        def al_salir(e):
            def revisar():   # si el foco pasó a la lista desplegable del propio combo, todavía no se confirma
                try:
                    if not str(cmb.tk.call("focus")).startswith(str(cmb)):
                        confirmar()
                except tk.TclError:   # la ventana se cerró
                    pass
            cmb.after(30, revisar)

        cmb.bind("<FocusIn>", al_entrar, add="+")
        cmb.bind("<KeyRelease>", tecla, add="+")
        cmb.bind("<Return>", confirmar, add="+")
        cmb.bind("<KP_Enter>", confirmar, add="+")
        cmb.bind("<FocusOut>", al_salir, add="+")

    class EditorListas(tk.Toplevel):
        """Edición de las listas de la hoja Datos (solo supervisor). Los cambios se aplican al presionar GUARDAR."""

        def __init__(self, app):
            super().__init__(app)
            self.app = app
            self.title("Editar listas")
            self.geometry("720x480")
            self.transient(app)
            self.grab_set()
            self.listas = app.db.listas_todas()
            self.renombres = []   # (lista, viejo, nuevo) para informar
            self.modificado = False

            izq = ttk.Frame(self, padding=8)
            izq.pack(side="left", fill="y")
            ttk.Label(izq, text="Lista").pack(anchor="w")
            self.lb_listas = tk.Listbox(izq, width=22, height=12, exportselection=False)
            for _, nombre in NOMBRES_LISTAS:
                self.lb_listas.insert("end", nombre)
            self.lb_listas.pack(fill="y", expand=True)
            self.lb_listas.bind("<<ListboxSelect>>", lambda e: self._cambiar_lista())

            centro = ttk.Frame(self, padding=8)
            centro.pack(side="left", fill="both", expand=True)
            self.frm_sector = ttk.Frame(centro)
            ttk.Label(self.frm_sector, text="Sector:").pack(side="left")
            self.cmb_sec = ttk.Combobox(self.frm_sector, state="readonly", width=15)
            self.cmb_sec.pack(side="left", padx=4)
            self.cmb_sec.bind("<<ComboboxSelected>>", lambda e: self._mostrar_valores())
            self.lbl_vals = ttk.Label(centro, text="Valores (en el orden en que aparecen en el desplegable)")
            self.lbl_vals.pack(anchor="w")
            self.lb_valores = tk.Listbox(centro, exportselection=False)
            self.lb_valores.pack(fill="both", expand=True)
            self.lb_valores.bind("<<ListboxSelect>>", lambda e: self._al_elegir_valor())

            der = ttk.Frame(self, padding=8)
            der.pack(side="left", fill="y")
            ttk.Label(der, text="Valor").pack(anchor="w")
            self.txt = ttk.Entry(der, width=22)
            self.txt.pack(pady=(0, 8))
            self.txt.bind("<Return>", lambda e: self.agregar())
            for texto, cmd in [("Agregar", self.agregar), ("Renombrar", self.renombrar),
                               ("Eliminar", self.eliminar), ("▲ Subir", lambda: self.mover(-1)),
                               ("▼ Bajar", lambda: self.mover(1))]:
                ttk.Button(der, text=texto, width=18, command=cmd).pack(pady=2)
            ttk.Separator(der).pack(fill="x", pady=10)
            ttk.Button(der, text="GUARDAR", width=18, command=self.guardar).pack(pady=2)
            ttk.Button(der, text="Cancelar", width=18, command=self.cerrar).pack(pady=2)
            self.lbl_nota = ttk.Label(der, text="", wraplength=160, foreground="gray")
            self.lbl_nota.pack(pady=8)
            self.protocol("WM_DELETE_WINDOW", self.cerrar)

            self.lb_listas.selection_set(0)
            self._cambiar_lista()

        # ----- navegación -----
        def _lista_actual(self):
            return NOMBRES_LISTAS[self.lb_listas.curselection()[0]][0]

        def _clave(self):
            lista = self._lista_actual()
            if lista == "maquinas":
                return ("maquinas", self.cmb_sec.get()) if self.cmb_sec.get() else None
            return (lista, None)

        def _valores(self):
            clave = self._clave()
            return self.listas.setdefault(clave, []) if clave else []

        def _cambiar_lista(self):
            if not self.lb_listas.curselection():
                return
            if self._lista_actual() == "maquinas":
                sectores = self.listas.get(("sectores", None), [])
                self.cmb_sec["values"] = sectores
                if self.cmb_sec.get() not in sectores:
                    self.cmb_sec.set(sectores[0] if sectores else "")
                self.frm_sector.pack(anchor="w", pady=(0, 6), before=self.lbl_vals)
            else:
                self.frm_sector.pack_forget()
            notas = {"operarios": "Las 2 primeras letras forman el código de OT.",
                     "sectores": "Al renombrar un sector se mantienen sus máquinas.",
                     "estado": "EP_FM exige n° de SRH en REPUESTOS.",
                     "extras": "EXTRA = rojo, DEV = azul en la lista."}
            self.lbl_nota.config(text=notas.get(self._lista_actual(), ""))
            self._mostrar_valores()

        def _mostrar_valores(self, seleccionar=None):
            self.lb_valores.delete(0, "end")
            for v in self._valores():
                self.lb_valores.insert("end", v)
            if seleccionar is not None and 0 <= seleccionar < self.lb_valores.size():
                self.lb_valores.selection_set(seleccionar)
                self.lb_valores.see(seleccionar)
            self.txt.delete(0, "end")

        def _al_elegir_valor(self):
            i = self._indice()
            if i is not None:
                self.txt.delete(0, "end")
                self.txt.insert(0, self._valores()[i])

        def _indice(self):
            s = self.lb_valores.curselection()
            return s[0] if s else None

        # ----- validaciones -----
        def _validar_nuevo(self, valor, excepto=None):
            if not valor:
                messagebox.showwarning("Valor", "Escriba un valor.", parent=self)
                return False
            if not self._clave():
                messagebox.showwarning("Máquinas", "Primero elija un sector.", parent=self)
                return False
            if any(v.upper() == valor.upper() for i, v in enumerate(self._valores()) if i != excepto):
                messagebox.showwarning("Valor", f"'{valor}' ya está en la lista.", parent=self)
                return False
            if self._lista_actual() == "operarios":
                if len(valor) < 2:
                    messagebox.showwarning("Operario", "El nombre debe tener al menos 2 letras.", parent=self)
                    return False
                iguales = [v for i, v in enumerate(self._valores()) if i != excepto and v[:2].upper() == valor[:2].upper()]
                if iguales and not messagebox.askyesno(
                        "Código de OT compartido",
                        f"'{valor}' comparte las 2 primeras letras ({valor[:2].upper()}) con: {', '.join(iguales)}.\n"
                        "Sus OT tendrán el mismo código y BUSCAR mostrará las actividades de ambos.\n\n¿Continuar?",
                        parent=self):
                    return False
            return True

        def _confirmar_uso(self, valor, accion):
            uso = VALORES_CON_USO.get((self._lista_actual(), valor))
            return not uso or messagebox.askyesno(
                "Valor usado por el programa",
                f"'{valor}' se usa en {uso}.\nSi lo {accion}, esa función deja de aplicarse.\n\n¿Continuar?",
                icon="warning", parent=self)

        # ----- acciones -----
        def agregar(self):
            valor = self.txt.get().strip()
            if not self._validar_nuevo(valor):
                return
            vals = self._valores()
            vals.append(valor)
            if self._lista_actual() == "sectores":
                self.listas.setdefault(("maquinas", valor), [])
            self.modificado = True
            self._mostrar_valores(len(vals) - 1)

        def renombrar(self):
            i = self._indice()
            if i is None:
                messagebox.showinfo("Renombrar", "Seleccione el valor a renombrar.", parent=self)
                return
            vals = self._valores()
            viejo, nuevo = vals[i], self.txt.get().strip()
            if nuevo == viejo or not self._validar_nuevo(nuevo, excepto=i) or not self._confirmar_uso(viejo, "renombra"):
                return
            vals[i] = nuevo
            if self._lista_actual() == "sectores" and ("maquinas", viejo) in self.listas:
                self.listas[("maquinas", nuevo)] = self.listas.pop(("maquinas", viejo))
            self.renombres.append((self._lista_actual(), viejo, nuevo))
            self.modificado = True
            self._mostrar_valores(i)

        def eliminar(self):
            i = self._indice()
            if i is None:
                messagebox.showinfo("Eliminar", "Seleccione el valor a eliminar.", parent=self)
                return
            vals = self._valores()
            valor = vals[i]
            if not self._confirmar_uso(valor, "elimina"):
                return
            if self._lista_actual() == "sectores":
                maqs = self.listas.get(("maquinas", valor), [])
                if not messagebox.askyesno("Eliminar sector",
                                           f"¿Eliminar el sector '{valor}'"
                                           + (f" y sus {len(maqs)} máquinas?" if maqs else "?"), parent=self):
                    return
                self.listas.pop(("maquinas", valor), None)
            del vals[i]
            self.modificado = True
            self._mostrar_valores(min(i, len(vals) - 1))

        def mover(self, delta):
            i = self._indice()
            vals = self._valores()
            if i is None or not 0 <= i + delta < len(vals):
                return
            vals[i], vals[i + delta] = vals[i + delta], vals[i]
            self.modificado = True
            self._mostrar_valores(i + delta)

        def guardar(self):
            if not self.modificado:
                self.destroy()
                return
            limpio = {k: v for k, v in self.listas.items() if v}
            try:
                self.app.db.guardar_listas(limpio)
            except sqlite3.Error as ex:
                messagebox.showerror("Error", f"No se pudieron guardar las listas:\n{ex}", parent=self)
                return
            self.app._cargar_combos()
            msg = "Listas guardadas."
            if self.renombres:
                msg += ("\n\nLas actividades ya cargadas conservan el nombre anterior:\n"
                        + "\n".join(f"  {v} → {n}" for _, v, n in self.renombres[:8]))
            messagebox.showinfo("Editar listas", msg, parent=self)
            self.destroy()

        def cerrar(self):
            if self.modificado and not messagebox.askyesno("Editar listas", "Hay cambios sin guardar. ¿Descartarlos?",
                                                           parent=self):
                return
            self.destroy()

    class FormTarea(tk.Toplevel):
        """Alta / modificación de una tarea de planificación (solo supervisor)."""

        def __init__(self, app, tarea=None, defecto=None):
            super().__init__(app)
            self.app, self.db, self.t = app, app.db, tarea
            self.title("Editar tarea" if tarea else "Nueva tarea")
            self.transient(app)
            self.grab_set()
            self.resizable(False, False)
            f = ttk.Frame(self, padding=12)
            f.pack(fill="both", expand=True)

            def etiqueta(texto, fila, col):
                ttk.Label(f, text=texto).grid(row=fila, column=col, sticky="w", padx=4, pady=(6, 0))

            etiqueta("OT Base", 0, 0)
            ttk.Label(f, text=(tarea or {}).get("ot_base") or "(se genera al guardar)",
                      font=("Segoe UI", 10, "bold")).grid(row=1, column=0, columnspan=2, sticky="w", padx=4)
            etiqueta("Tipo", 2, 0)
            self.tipo = ttk.Combobox(f, width=14, values=self.db.lista("tipos_tarea"))
            self.tipo.grid(row=3, column=0, sticky="w", padx=4)
            etiqueta("Prioridad", 2, 1)
            self.prio = ttk.Combobox(f, width=14, state="readonly", values=[""] + self.db.lista("prioridad"))
            self.prio.grid(row=3, column=1, sticky="w", padx=4)
            etiqueta("Sector *", 4, 0)
            self.sector = ttk.Combobox(f, width=14, state="readonly", values=self.db.lista("sectores"))
            self.sector.grid(row=5, column=0, sticky="w", padx=4)
            self.sector.bind("<<ComboboxSelected>>",
                             lambda e: (self.maq.config(values=self.db.lista("maquinas", self.sector.get())), self.maq.set("")))
            etiqueta("Máquina", 4, 1)
            self.maq = ttk.Combobox(f, width=18)
            self.maq.grid(row=5, column=1, sticky="w", padx=4)
            etiqueta("Conjunto", 6, 0)
            self.conj = ttk.Entry(f, width=17)
            self.conj.grid(row=7, column=0, sticky="w", padx=4)
            etiqueta("Subconjunto", 6, 1)
            self.subconj = ttk.Entry(f, width=21)
            self.subconj.grid(row=7, column=1, sticky="w", padx=4)
            etiqueta("Tarea *", 8, 0)
            self.tarea = tk.Text(f, width=52, height=4, wrap="word", font=("Segoe UI", 9))
            self.tarea.grid(row=9, column=0, columnspan=2, sticky="we", padx=4)
            etiqueta("Asignado a", 10, 0)
            self.asig = ttk.Combobox(f, width=14, state="readonly", values=[""] + self.db.lista("operarios"))
            self.asig.grid(row=11, column=0, sticky="w", padx=4)
            etiqueta("Comentarios", 12, 0)
            self.coment = tk.Text(f, width=52, height=2, wrap="word", font=("Segoe UI", 9))
            self.coment.grid(row=13, column=0, columnspan=2, sticky="we", padx=4)

            per = ttk.LabelFrame(f, text="Programación", padding=8)
            per.grid(row=14, column=0, columnspan=2, sticky="we", padx=4, pady=(10, 4))
            self.var_per = tk.BooleanVar(value=bool(tarea and es_periodica(tarea)) or bool((defecto or {}).get("periodica")))
            ttk.Checkbutton(per, text="Periódica", variable=self.var_per,
                            command=self._act_periodica).grid(row=0, column=0, sticky="w")
            self.lbl_fecha = ttk.Label(per, text="Fecha a realizar")
            self.lbl_fecha.grid(row=1, column=0, sticky="w", pady=(6, 0))
            self.fecha = ttk.Entry(per, width=12)
            self.fecha.grid(row=2, column=0, sticky="w")
            ttk.Label(per, text="Cada").grid(row=1, column=1, sticky="w", padx=(12, 0), pady=(6, 0))
            self.periodo = ttk.Spinbox(per, from_=1, to=999, width=5)
            self.periodo.grid(row=2, column=1, sticky="w", padx=(12, 0))
            self.unidad = ttk.Combobox(per, width=9, state="readonly", values=list(UNIDADES.values()))
            self.unidad.grid(row=2, column=2, sticky="w", padx=4)
            ttk.Label(per, text="Hasta (opcional)").grid(row=1, column=3, sticky="w", padx=(12, 0), pady=(6, 0))
            self.fin = ttk.Entry(per, width=12)
            self.fin.grid(row=2, column=3, sticky="w", padx=(12, 0))
            ttk.Label(per, text="Fechas en formato dd-mm-aaaa.", foreground="gray").grid(
                row=3, column=0, columnspan=4, sticky="w", pady=(6, 0))
            self.wid_per = [self.periodo, self.unidad, self.fin]

            bot = ttk.Frame(f)
            bot.grid(row=15, column=0, columnspan=2, sticky="e", pady=(8, 0))
            ttk.Button(bot, text="Cancelar", command=self.destroy).pack(side="right")
            ttk.Button(bot, text="GUARDAR", command=self.guardar).pack(side="right", padx=6)

            if tarea:
                self.tipo.set(tarea["tipo"] or "")
                self.prio.set(tarea["prioridad"] or "")
                self.sector.set(tarea["sector"] or "")
                self.maq.config(values=self.db.lista("maquinas", tarea["sector"] or ""))
                self.maq.set(tarea["maquina"] or "")
                self.conj.insert(0, tarea["conjunto"] or "")
                self.subconj.insert(0, tarea["subconjunto"] or "")
                self.tarea.insert("1.0", tarea["tarea"] or "")
                self.asig.set(tarea["asignado"] or "")
                self.coment.insert("1.0", tarea["comentarios"] or "")
                self.fecha.insert(0, db_a_texto(tarea["fecha"]))
                self.periodo.set(tarea["periodo"] or 1)
                self.unidad.set(UNIDADES.get(tarea["unidad"] or "dias"))
                self.fin.insert(0, db_a_texto(tarea["fecha_fin"]))
            else:
                self.periodo.set(1)
                self.unidad.set(UNIDADES["meses"])
                if defecto:
                    self.sector.set(defecto.get("sector") or "")
                    self.maq.config(values=self.db.lista("maquinas", self.sector.get()))
                    self.maq.set(defecto.get("maquina") or "")
            self._act_periodica()
            for w, estricto in ((self.sector, True), (self.prio, True), (self.asig, True), (self.maq, False),
                                (self.tipo, False)):
                autocompletar(w, estricto)

        def _act_periodica(self):
            per = self.var_per.get()
            self.lbl_fecha.config(text="Fecha de inicio" if per else "Fecha a realizar")
            for w in self.wid_per:
                w.config(state=("readonly" if w is self.unidad else "normal") if per else "disabled")

        def guardar(self):
            tarea = self.tarea.get("1.0", "end").strip()
            if not self.sector.get() or not tarea:
                messagebox.showwarning("Falta dato", "Sector y Tarea son obligatorios.", parent=self)
                return
            fecha_txt, fin_txt = self.fecha.get().strip(), self.fin.get().strip()
            fecha = fecha_a_db(fecha_txt) if fecha_txt else None
            fin = fecha_a_db(fin_txt) if fin_txt else None
            if (fecha_txt and not fecha) or (fin_txt and not fin):
                messagebox.showwarning("Fecha", "Fecha inválida. Use dd-mm-aaaa.", parent=self)
                return
            per = self.var_per.get()
            periodo = unidad = None
            if per:
                try:
                    periodo = int(self.periodo.get())
                    assert periodo > 0
                except (ValueError, AssertionError):
                    messagebox.showwarning("Frecuencia", "Ingrese una frecuencia válida (número mayor a 0).", parent=self)
                    return
                unidad = {v: k for k, v in UNIDADES.items()}[self.unidad.get()]
                if not fecha and not messagebox.askyesno(
                        "Sin fecha de inicio", "La tarea quedará 'Sin programar' hasta que le cargue una fecha de "
                                               "inicio. ¿Guardar igual?", parent=self):
                    return
            t = dict(tipo=self.tipo.get().strip() or None, prioridad=self.prio.get() or None,
                     sector=self.sector.get(), maquina=self.maq.get().strip() or None,
                     conjunto=self.conj.get().strip() or None, subconjunto=self.subconj.get().strip() or None,
                     tarea=tarea, asignado=self.asig.get() or None,
                     comentarios=self.coment.get("1.0", "end").strip() or None,
                     fecha=fecha, periodo=periodo, unidad=unidad, fecha_fin=fin if per else None)
            try:
                if self.t is None:
                    _, _, ot = self.db.guardar_tarea(t)
                    messagebox.showinfo("Nueva tarea", f"Tarea creada con OT {ot}", parent=self)
                else:
                    res, _, _ = self.db.guardar_tarea(t, self.t["id"], self.t["rev"] or 0)
                    if res == "eliminado":
                        messagebox.showwarning("Editar tarea", "La tarea fue eliminada por otro usuario.", parent=self)
                    elif res == "conflicto":
                        if not messagebox.askyesno("Editar tarea",
                                                   "Otro usuario modificó esta tarea mientras usted la editaba.\n"
                                                   "¿Guardar sus cambios igual?", icon="warning", parent=self):
                            return
                        self.db.guardar_tarea(t, self.t["id"])
            except sqlite3.Error as ex:
                messagebox.showerror("Error", f"No se pudo guardar:\n{ex}", parent=self)
                return
            self.app._refrescar_planes()
            self.destroy()

    class Historico(tk.Toplevel):
        """Todas las actividades cargadas para una tarea (todas sus fechas). Doble clic abre la actividad."""

        def __init__(self, app, t):
            super().__init__(app)
            self.app = app
            self.title(f"Histórico {t['ot_base']}")
            self.geometry("1100x420")
            self.transient(app)
            ttk.Label(self, text=f"{t['ot_base']}   |   {t['sector'] or ''} {t['maquina'] or ''}   |   "
                                 f"{texto_periodicidad(t)}", font=("Segoe UI", 10, "bold"),
                      padding=(10, 8, 10, 0)).pack(anchor="w")
            ttk.Label(self, text=t["tarea"] or "", wraplength=1050, padding=(10, 2, 10, 6)).pack(anchor="w")
            cols = [("fecha", "Fecha", 90), ("ot", "OT MAN", 85), ("operario", "Operario", 90), ("estado", "Estado", 65),
                    ("hi", "HI", 50), ("hf", "HF", 50), ("actividad", "Actividad", 280),
                    ("comentarios", "Comentarios", 250), ("repuestos", "Repuestos / SRH", 150)]
            frm = ttk.Frame(self, padding=(10, 0, 10, 10))
            frm.pack(fill="both", expand=True)
            tree = ttk.Treeview(frm, columns=[c[0] for c in cols], show="headings")
            for c, txt, w in cols:
                tree.heading(c, text=txt)
                tree.column(c, width=w, anchor="w", stretch=c in ("actividad", "comentarios"))
            sy = ttk.Scrollbar(frm, orient="vertical", command=tree.yview)
            tree.configure(yscrollcommand=sy.set)
            tree.pack(side="left", fill="both", expand=True)
            sy.pack(side="right", fill="y")
            acts = sorted(app.db.actividades_de_tareas().get(t["id"], []), key=lambda a: (a["fecha"] or "", a["id"]))
            for a in acts:
                tree.insert("", "end", iid=str(a["id"]), values=[db_a_texto(a[c[0]]) for c in cols])
            if not acts:
                tree.insert("", "end", values=["", "", "", "", "", "", "Todavía no hay actividades cargadas."])
            tree.bind("<Double-1>", lambda e: self._abrir(tree))

        def _abrir(self, tree):
            sel = tree.selection()
            if sel and sel[0].isdigit():
                self.app.nb.select(0)
                self.app._cargar_id(int(sel[0]))
                self.destroy()

    def _ventana(win, app, titulo, tamano=None):
        win.title(titulo)
        if tamano:
            win.geometry(tamano)
        win.transient(app)

    class OrdenChecklist(tk.Toplevel):
        """Lista de verificación de una orden de preventivo: cada ítem se marca Hecho / No hecho."""

        def __init__(self, app, orden_id):
            super().__init__(app)
            self.app, self.db, self.orden_id = app, app.db, orden_id
            self.o = self.db.orden(orden_id)
            if not self.o:
                self.destroy()
                messagebox.showwarning("Orden", "La orden ya no existe.")
                return
            _ventana(self, app, f"Orden {self.o['ot']}", "1150x560")
            cab = ttk.Frame(self, padding=(10, 8, 10, 4))
            cab.pack(fill="x")
            ttk.Label(cab, text=f"{self.o['ot']}   |   {self.o['sector']} {self.o['maquina']}   |   Programada: "
                                f"{db_a_texto(self.o['fecha'])}   |   Creada: {db_a_texto(self.o['creada'])} "
                                f"({self.o['creada_por']})", font=("Segoe UI", 10, "bold")).pack(side="left")
            self.lbl_avance = ttk.Label(cab, text="", font=("Segoe UI", 10, "bold"), foreground="#007a33")
            self.lbl_avance.pack(side="right")

            frm = ttk.Frame(self, padding=(10, 0, 10, 0))
            frm.pack(fill="both", expand=True)
            cols = [("conjunto", "Conjunto", 110), ("subconjunto", "Subconjunto", 110), ("tarea", "Ítem", 330),
                    ("periodicidad", "Periodicidad", 105), ("programada", "Programada", 92), ("resultado", "Resultado", 80),
                    ("fecha_real", "Realizada", 92), ("comentario", "Comentario", 220), ("info", "Registrado", 140)]
            self.tree = ttk.Treeview(frm, columns=[c[0] for c in cols], show="headings", selectmode="extended")
            for c, t, w in cols:
                self.tree.heading(c, text=t)
                self.tree.column(c, width=w, anchor="w", stretch=c in ("tarea", "comentario"))
            sy = ttk.Scrollbar(frm, orient="vertical", command=self.tree.yview)
            self.tree.configure(yscrollcommand=sy.set)
            self.tree.pack(side="left", fill="both", expand=True)
            sy.pack(side="right", fill="y")
            self.tree.tag_configure("Hecho", foreground="#007a33")
            self.tree.tag_configure("No hecho", foreground="#7a0050")

            ctl = ttk.Frame(self, padding=(10, 6, 10, 2))
            ctl.pack(fill="x")
            ttk.Label(ctl, text="Fecha realizada").pack(side="left")
            self.fecha = ttk.Entry(ctl, width=11)
            self.fecha.insert(0, date.today().strftime("%d-%m-%Y"))
            self.fecha.pack(side="left", padx=(4, 12))
            ttk.Label(ctl, text="Comentario").pack(side="left")
            self.coment = ttk.Entry(ctl, width=40)
            self.coment.pack(side="left", padx=4)
            ttk.Label(self, text="Seleccione uno o varios ítems (Ctrl / Shift + clic) y marque el resultado. "
                                 "El comentario se aplica solo si se escribe algo.", foreground="gray",
                      padding=(10, 0)).pack(anchor="w")
            bot = ttk.Frame(self, padding=(10, 4, 10, 10))
            bot.pack(fill="x")
            ttk.Button(bot, text="✔ HECHO", command=lambda: self.marcar("Hecho")).pack(side="left")
            ttk.Button(bot, text="✘ NO HECHO", command=lambda: self.marcar("No hecho")).pack(side="left", padx=4)
            ttk.Button(bot, text="Volver a pendiente", command=lambda: self.marcar("")).pack(side="left")
            ttk.Button(bot, text="Todos los pendientes → Hecho", command=self.todos_hechos).pack(side="left", padx=12)
            b = ttk.Button(bot, text="Quitar de la orden", command=self.quitar)
            b.pack(side="left")
            if not app.supervisor:
                b.state(["disabled"])
            ttk.Button(bot, text="CARGAR ACTIVIDAD DE ESTA ORDEN", command=self.cargar_actividad).pack(side="right")
            self.refrescar()

        def refrescar(self):
            self.items = self.db.items_de_orden(self.orden_id)
            sel = self.tree.selection()
            self.tree.delete(*self.tree.get_children())
            for it in self.items:
                per = texto_periodicidad(dict(periodo=it["periodo"], unidad=it["unidad"]))
                self.tree.insert("", "end", iid=f"{it['tarea_id']}|{it['fecha_occ']}", tags=(it["resultado"] or "",),
                                 values=[it["conjunto"] or "", it["subconjunto"] or "", it["tarea"] or "(ítem eliminado)",
                                         per, db_a_texto(it["fecha_occ"]), it["resultado"] or "Pendiente",
                                         db_a_texto(it["fecha_real"]), it["comentario"] or "", it["info"] or ""])
            for i in sel:
                if self.tree.exists(i):
                    self.tree.selection_add(i)
            n = len(self.items)
            h = sum(1 for i in self.items if i["resultado"] == "Hecho")
            nh = sum(1 for i in self.items if i["resultado"] == "No hecho")
            self.lbl_avance.config(text=f"Hechos {h} de {n}   ·   No hechos {nh}   ·   Pendientes {n - h - nh}")

        def _claves(self, iids):
            return [(int(i.split("|")[0]), i.split("|", 1)[1]) for i in iids]

        def marcar(self, resultado, iids=None):
            iids = iids if iids is not None else self.tree.selection()
            if not iids:
                messagebox.showinfo("Checklist", "Seleccione uno o más ítems.", parent=self)
                return
            fecha = None
            if resultado == "Hecho":
                fecha = fecha_a_db(self.fecha.get())
                if not fecha:
                    messagebox.showwarning("Fecha", "Fecha realizada inválida (dd-mm-aaaa).", parent=self)
                    return
            com = self.coment.get().strip() or None
            self.db.marcar_items(self.orden_id, self._claves(iids), resultado, fecha, com)
            self.coment.delete(0, "end")
            self.refrescar()
            self.app._refrescar_planes()

        def todos_hechos(self):
            pend = [f"{i['tarea_id']}|{i['fecha_occ']}" for i in self.items if not i["resultado"]]
            if not pend:
                return
            if messagebox.askyesno("Checklist", f"¿Marcar los {len(pend)} ítems pendientes como HECHOS con fecha "
                                                f"{self.fecha.get()}?", parent=self):
                self.marcar("Hecho", pend)

        def quitar(self):
            iids = self.tree.selection()
            if iids and messagebox.askyesno("Quitar", f"¿Quitar {len(iids)} ítem(s) de esta orden? Vuelven a quedar "
                                                      "pendientes para la próxima orden.", parent=self):
                self.db.quitar_items(self.orden_id, self._claves(iids))
                self.refrescar()
                self.app._refrescar_planes()

        def cargar_actividad(self):
            t = self.db.tareas(self.o["tarea_id"])
            if t:
                self.app._cargar_desde_tarea(t, t["fecha"] or "", iso_a_date(t["fecha"]))
                self.destroy()

    class GenerarOrden(tk.Toplevel):
        """Junta los ítems vencidos y por vencer de la máquina en una OT general de preventivo."""

        def __init__(self, app, sector, maquina):
            super().__init__(app)
            self.app, self.db, self.sector, self.maquina = app, app.db, sector, maquina
            _ventana(self, app, f"Generar orden de preventivo - {maquina}", "980x560")
            self.grab_set()
            f = ttk.Frame(self, padding=(10, 8))
            f.pack(fill="x")
            ttk.Label(f, text=f"{sector} / {maquina}", font=("Segoe UI", 10, "bold")).grid(row=0, column=0, columnspan=6,
                                                                                         sticky="w")
            ttk.Label(f, text="Fecha de la orden").grid(row=1, column=0, sticky="w", pady=(6, 0))
            self.fecha = ttk.Entry(f, width=11)
            self.fecha.insert(0, date.today().strftime("%d-%m-%Y"))
            self.fecha.grid(row=2, column=0, sticky="w")
            ttk.Label(f, text="Incluir los que vencen en los próximos (días)").grid(row=1, column=1, sticky="w",
                                                                                   padx=12, pady=(6, 0))
            self.ant = ttk.Spinbox(f, from_=0, to=365, width=6)
            self.ant.set(app.cfg.getint("plan", "anticipacion_orden_dias", fallback=15))
            self.ant.grid(row=2, column=1, sticky="w", padx=12)
            ttk.Label(f, text="Asignado a").grid(row=1, column=2, sticky="w", pady=(6, 0))
            self.asig = ttk.Combobox(f, width=14, state="readonly", values=[""] + self.db.lista("operarios"))
            self.asig.grid(row=2, column=2, sticky="w")
            ttk.Label(f, text="Comentarios").grid(row=1, column=3, sticky="w", padx=12, pady=(6, 0))
            self.coment = ttk.Entry(f, width=40)
            self.coment.grid(row=2, column=3, sticky="w", padx=12)
            ttk.Button(f, text="Buscar ítems", command=self.buscar).grid(row=2, column=4, sticky="w")
            autocompletar(self.asig, True)

            frm = ttk.Frame(self, padding=(10, 4, 10, 0))
            frm.pack(fill="both", expand=True)
            cols = [("inc", "✔", 30), ("conjunto", "Conjunto", 110), ("tarea", "Ítem", 360),
                    ("periodicidad", "Periodicidad", 105), ("programada", "Programada", 85), ("estado", "Estado", 85)]
            self.tree = ttk.Treeview(frm, columns=[c[0] for c in cols], show="headings", selectmode="extended")
            for c, t, w in cols:
                self.tree.heading(c, text=t)
                self.tree.column(c, width=w, anchor="center" if c == "inc" else "w", stretch=c == "tarea")
            sy = ttk.Scrollbar(frm, orient="vertical", command=self.tree.yview)
            self.tree.configure(yscrollcommand=sy.set)
            self.tree.pack(side="left", fill="both", expand=True)
            sy.pack(side="right", fill="y")
            self.tree.tag_configure("Vencida", foreground="#c00000")
            self.tree.tag_configure("fuera", foreground="gray")
            self.tree.bind("<Double-1>", lambda e: self.alternar())
            self.tree.bind("<space>", lambda e: self.alternar())
            bot = ttk.Frame(self, padding=(10, 6, 10, 10))
            bot.pack(fill="x")
            ttk.Button(bot, text="Incluir / quitar seleccionados", command=self.alternar).pack(side="left")
            ttk.Button(bot, text="Incluir todos", command=lambda: self._todos(True)).pack(side="left", padx=4)
            ttk.Button(bot, text="Quitar todos", command=lambda: self._todos(False)).pack(side="left")
            self.lbl = ttk.Label(bot, text="")
            self.lbl.pack(side="left", padx=12)
            ttk.Button(bot, text="Cancelar", command=self.destroy).pack(side="right")
            ttk.Button(bot, text="GENERAR ORDEN", command=self.generar).pack(side="right", padx=6)
            self.cands, self.incluidos = [], set()
            self.buscar()

        def buscar(self):
            fecha = iso_a_date(fecha_a_db(self.fecha.get()))
            try:
                ant = int(self.ant.get())
            except ValueError:
                ant = 0
            if not fecha:
                messagebox.showwarning("Fecha", "Fecha inválida (dd-mm-aaaa).", parent=self)
                return
            self.cands = self.db.candidatos_orden(self.sector, self.maquina, fecha + timedelta(days=ant), fecha)
            self.incluidos = {self._iid(c) for c in self.cands}
            self._mostrar()

        def _iid(self, c):
            return f"{c['tarea']['id']}|{c['fecha_occ'].isoformat()}"

        def _mostrar(self):
            self.tree.delete(*self.tree.get_children())
            for c in self.cands:
                iid, t = self._iid(c), c["tarea"]
                dentro = iid in self.incluidos
                self.tree.insert("", "end", iid=iid, tags=(c["estado"] if dentro else "fuera",), values=[
                    "✔" if dentro else "", t["conjunto"] or "", t["tarea"] or "", texto_periodicidad(t),
                    c["fecha_occ"].strftime("%d-%m-%Y"), c["estado"]])
            self.lbl.config(text=f"{len(self.incluidos)} de {len(self.cands)} ítems incluidos" if self.cands else
                            "No hay ítems vencidos ni por vencer para esta máquina.")

        def alternar(self):
            for i in self.tree.selection():
                self.incluidos.symmetric_difference_update({i})
            sel = self.tree.selection()
            self._mostrar()
            self.tree.selection_set([i for i in sel if self.tree.exists(i)])

        def _todos(self, incluir):
            self.incluidos = {self._iid(c) for c in self.cands} if incluir else set()
            self._mostrar()

        def generar(self):
            if not self.incluidos:
                messagebox.showwarning("Orden", "No hay ítems incluidos.", parent=self)
                return
            fecha = fecha_a_db(self.fecha.get())
            items = [(int(i.split("|")[0]), i.split("|")[1]) for i in self.incluidos]
            try:
                oid, ot = self.db.crear_orden(self.sector, self.maquina, fecha, self.asig.get() or None, items,
                                              self.coment.get().strip() or None)
            except sqlite3.Error as ex:
                messagebox.showerror("Orden", f"No se pudo crear la orden:\n{ex}", parent=self)
                return
            self.app._refrescar_planes()
            self.destroy()
            if messagebox.askyesno("Orden de preventivo", f"Orden {ot} creada con {len(items)} ítems.\n\n"
                                                          "¿Abrir su checklist?"):
                OrdenChecklist(self.app, oid)

    class EdicionLote(tk.Toplevel):
        """Cambia los mismos campos en varios ítems a la vez. Lo que se deja vacío no se modifica."""

        def __init__(self, app, ids):
            super().__init__(app)
            self.app, self.db, self.ids = app, app.db, ids
            _ventana(self, app, f"Editar {len(ids)} ítems")
            self.grab_set()
            self.resizable(False, False)
            f = ttk.Frame(self, padding=12)
            f.pack()
            ttk.Label(f, text=f"Se modificarán {len(ids)} ítems. Los campos vacíos NO se cambian.",
                      font=("Segoe UI", 9, "bold")).grid(row=0, column=0, columnspan=4, sticky="w", pady=(0, 8))
            self.w = {}

            def campo(texto, fila, col, w):
                ttk.Label(f, text=texto).grid(row=fila, column=col, sticky="w", padx=4, pady=(4, 0))
                w.grid(row=fila + 1, column=col, sticky="w", padx=4)
                return w
            self.w["asignado"] = campo("Asignado", 1, 0, ttk.Combobox(f, width=16, state="readonly",
                                                                      values=["", "(quitar asignado)"] + self.db.lista("operarios")))
            self.w["prioridad"] = campo("Prioridad", 1, 1, ttk.Combobox(f, width=12, state="readonly",
                                                                        values=[""] + self.db.lista("prioridad")))
            self.w["tipo"] = campo("Tipo", 1, 2, ttk.Combobox(f, width=10, values=[""] + self.db.lista("tipos_tarea")))
            self.w["fecha"] = campo("Fecha de inicio (dd-mm-aaaa)", 3, 0, ttk.Entry(f, width=12))
            self.w["periodo"] = campo("Cada", 3, 1, ttk.Spinbox(f, from_=1, to=999, width=6))
            self.w["periodo"].set("")
            self.w["unidad"] = campo("Unidad", 3, 2, ttk.Combobox(f, width=10, state="readonly",
                                                                  values=[""] + list(UNIDADES.values())))
            self.w["fecha_fin"] = campo("Hasta (dd-mm-aaaa)", 3, 3, ttk.Entry(f, width=12))
            bot = ttk.Frame(f)
            bot.grid(row=5, column=0, columnspan=4, sticky="e", pady=(12, 0))
            ttk.Button(bot, text="Cancelar", command=self.destroy).pack(side="right")
            ttk.Button(bot, text="APLICAR", command=self.aplicar).pack(side="right", padx=6)
            autocompletar(self.w["asignado"], True)
            autocompletar(self.w["prioridad"], True)

        def aplicar(self):
            c = {}
            asig = self.w["asignado"].get()
            if asig:
                c["asignado"] = None if asig == "(quitar asignado)" else asig
            for k in ("prioridad", "tipo"):
                if self.w[k].get().strip():
                    c[k] = self.w[k].get().strip()
            for k in ("fecha", "fecha_fin"):
                txt = self.w[k].get().strip()
                if txt:
                    if not fecha_a_db(txt):
                        messagebox.showwarning("Fecha", f"Fecha inválida: {txt}", parent=self)
                        return
                    c[k] = fecha_a_db(txt)
            per, uni = self.w["periodo"].get().strip(), self.w["unidad"].get()
            if per:
                try:
                    c["periodo"] = int(per)
                    assert c["periodo"] > 0
                except (ValueError, AssertionError):
                    messagebox.showwarning("Cada", "Ingrese un número mayor a 0.", parent=self)
                    return
                if not uni:
                    messagebox.showwarning("Unidad", "Indique la unidad (días, semanas o meses).", parent=self)
                    return
            if uni:
                c["unidad"] = {v: k for k, v in UNIDADES.items()}[uni]
            if not c:
                messagebox.showinfo("Editar en lote", "No completó ningún campo.", parent=self)
                return
            if not messagebox.askyesno("Editar en lote", f"¿Aplicar los cambios a {len(self.ids)} ítems?", parent=self):
                return
            for i in self.ids:
                self.db.guardar_tarea(c, i)
            self.app._refrescar_planes()
            self.destroy()

    class ImportarPreview(tk.Toplevel):
        """Resumen de la planilla antes de importar: nuevos, existentes y errores por fila."""

        def __init__(self, app, res):
            super().__init__(app)
            self.app, self.res = app, res
            _ventana(self, app, "Importar ítems desde planilla", "760x460")
            self.grab_set()
            f = ttk.Frame(self, padding=12)
            f.pack(fill="both", expand=True)
            n, e, x = len(res["nuevos"]), len(res["existentes"]), len(res["errores"])
            ttk.Label(f, text=f"Ítems nuevos: {n}      Ya existentes: {e}      Filas con errores (no se importan): {x}",
                      font=("Segoe UI", 10, "bold")).pack(anchor="w")
            self.act = tk.BooleanVar(value=False)
            cb = ttk.Checkbutton(f, text=f"Actualizar los {e} ítems existentes con los datos de la planilla "
                                         "(periodicidad, fecha, asignado, etc.)", variable=self.act)
            cb.pack(anchor="w", pady=6)
            if not e:
                cb.state(["disabled"])
            ttk.Label(f, text="Errores:" if x else "Sin errores.").pack(anchor="w", pady=(6, 2))
            txt = tk.Text(f, height=14, wrap="word", font=("Consolas", 9))
            txt.pack(fill="both", expand=True)
            for fila, msg in res["errores"]:
                txt.insert("end", f"Fila {fila}: {msg}\n")
            txt.config(state="disabled")
            bot = ttk.Frame(f)
            bot.pack(fill="x", pady=(8, 0))
            ttk.Button(bot, text="Cancelar", command=self.destroy).pack(side="right")
            b = ttk.Button(bot, text="IMPORTAR", command=self.importar)
            b.pack(side="right", padx=6)
            if not (n or e):
                b.state(["disabled"])

        def importar(self):
            act = self.res["existentes"] if self.act.get() else []
            try:
                self.app.db.importar_items(self.res["nuevos"], act)
            except sqlite3.Error as ex:
                messagebox.showerror("Importar", f"No se pudo importar:\n{ex}", parent=self)
                return
            messagebox.showinfo("Importar", f"Se agregaron {len(self.res['nuevos'])} ítems"
                                + (f" y se actualizaron {len(act)}." if act else "."), parent=self)
            self.app._refrescar_planes()
            self.app.nb.select(2)
            self.destroy()

    class ResumenOperarios(tk.Toplevel):
        """Tareas programadas por operario y por mes: cumplidas / programadas y vencidas."""

        def __init__(self, app):
            super().__init__(app)
            self.app = app
            _ventana(self, app, "Resumen por operario", "1150x480")
            top = ttk.Frame(self, padding=(10, 8))
            top.pack(fill="x")
            ttk.Label(top, text="Año").pack(side="left")
            self.anio = ttk.Spinbox(top, from_=2020, to=2100, width=6, command=self.refrescar)
            self.anio.set(date.today().year)
            self.anio.pack(side="left", padx=4)
            self.anio.bind("<Return>", lambda e: self.refrescar())
            ttk.Button(top, text="Exportar a Excel", command=self.exportar).pack(side="right")
            ttk.Label(top, text="Celda: cumplidas / programadas  ·  v = vencidas sin hacer.  % = cumplidas sobre las "
                                "ya exigibles (fecha ≤ hoy).", foreground="gray").pack(side="left", padx=15)
            cols = ["Operario"] + [m[:3] for m in MESES] + ["Total", "% cumpl."]
            frm = ttk.Frame(self, padding=(10, 0, 10, 10))
            frm.pack(fill="both", expand=True)
            self.tree = ttk.Treeview(frm, columns=cols, show="headings")
            for c in cols:
                self.tree.heading(c, text=c)
                self.tree.column(c, width=130 if c == "Operario" else 70, anchor="w" if c == "Operario" else "center")
            sy = ttk.Scrollbar(frm, orient="vertical", command=self.tree.yview)
            self.tree.configure(yscrollcommand=sy.set)
            self.tree.pack(side="left", fill="both", expand=True)
            sy.pack(side="right", fill="y")
            self.filas = []
            self.refrescar()

        def refrescar(self):
            try:
                y = int(self.anio.get())
            except ValueError:
                return
            hoy = date.today()
            plan = calcular_plan(self.app.db, date(y, 1, 1), date(y, 12, 31), atrasadas=False)
            datos = {}
            for f in plan:
                if not f["fecha"] or f["estado"] in ("Sin programar", "Cancelada"):
                    continue
                op = f["asignado"] or "(sin asignar)"
                d = datos.setdefault(op, {m: [0, 0, 0] for m in range(1, 13)})
                celda = d[f["fecha"].month]
                celda[0] += 1
                celda[1] += f["estado"] == "Cumplida"
                celda[2] += f["estado"] == "Vencida"
            self.tree.delete(*self.tree.get_children())
            self.filas = []
            for op in sorted(datos, key=lambda o: (o == "(sin asignar)", o)):
                d = datos[op]
                txt = lambda t, c, v: (f"{c}/{t}" + (f" · {v}v" if v else "")) if t else ""
                tot = [sum(d[m][i] for m in d) for i in range(3)]
                exig = sum(1 for f in plan if (f["asignado"] or "(sin asignar)") == op and f["fecha"]
                           and f["fecha"] <= hoy and f["estado"] not in ("Sin programar", "Cancelada"))
                pct = f"{100 * tot[1] / exig:.0f}%" if exig else "-"
                fila = [op] + [txt(*d[m]) for m in range(1, 13)] + [txt(*tot), pct]
                self.filas.append(fila)
                self.tree.insert("", "end", values=fila)

        def exportar(self):
            ruta = filedialog.asksaveasfilename(parent=self, defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")],
                                                initialfile=f"resumen_operarios_{self.anio.get()}.xlsx")
            if not ruta:
                return
            wb = Workbook(write_only=True)
            ws = wb.create_sheet("Resumen")
            ws.append(["Operario"] + MESES + ["Total", "% cumplimiento"])
            for f in self.filas:
                ws.append(f)
            try:
                wb.save(ruta)
            except PermissionError:
                messagebox.showerror("Error", "No se pudo guardar: el archivo está abierto en Excel.", parent=self)
                return
            messagebox.showinfo("Exportar", f"Exportado a:\n{ruta}", parent=self)

    class App(tk.Tk):
        def __init__(self):
            super().__init__()
            self.title("SECTOR DE MANTENIMIENTO")
            # Se adapta a pantallas chicas (ej. 1366x768 con barra de tareas)
            alto = min(760, self.winfo_screenheight() - 70)
            self.geometry(f"{min(1320, self.winfo_screenwidth() - 20)}x{alto}+5+5")
            self.minsize(1000, 560)
            self.cfg = leer_config()
            ruta = self._resolver_ruta_db(ruta_abs(self.cfg.get("base", "ruta")))
            if not ruta:
                self.destroy()
                raise SystemExit
            self.msg_respaldo = self._respaldo_automatico(ruta)
            try:
                self.db = BaseDatos(ruta)
            except sqlite3.Error as ex:
                messagebox.showerror("Base de datos", f"No se pudo abrir la base:\n{ruta}\n\n{ex}")
                self.destroy()
                raise SystemExit
            self.supervisor = False
            self.id_cargado = None       # registro cargado en el formulario
            self.rev_cargado = None      # versión del registro al cargarlo (para detectar cambios de otros)
            self.filas_lista = []        # filas mostradas en la lista
            self.consulta = None         # función que vuelve a generar la lista (BUSCAR o Filtrar)
            self.cambios_pendientes = False
            self.ultimo_refresco = None
            self.vinculo_pendiente = None  # (tarea_id, fecha_occ, ot_base) al cargar desde Planificación
            self.plan_filas = []
            self.campos = {}             # encabezado -> widget
            self.solo_sup = []           # widgets habilitados solo en modo supervisor
            self._armar()
            self._aplicar_permisos()
            self.dv = self.db.data_version()
            self._actualizar_estado()
            if self.db.vacia:
                self.after(200, self.importar)
            else:
                self._cargar_listas()
            seg = self.cfg.getint("app", "refresco_seg", fallback=30)
            if seg > 0:
                self.after(seg * 1000, lambda: self._vigilar(seg))

        # ---------------- Base compartida ----------------
        def _resolver_ruta_db(self, ruta):
            """Verifica que la carpeta de la base sea accesible (por ejemplo, la carpeta de red)."""
            while not os.path.isdir(os.path.dirname(ruta) or "."):
                if not messagebox.askyesno(
                        "Base de datos no accesible",
                        f"No se encuentra la carpeta de la base:\n{os.path.dirname(ruta)}\n\n"
                        "Puede ser un problema de red o una ruta mal configurada en config.ini.\n\n"
                        "¿Desea elegir otra ubicación? (No = cerrar el programa)", icon="warning"):
                    return None
                nueva = self._elegir_ruta_db()
                if nueva:
                    ruta = nueva
            return ruta

        def _elegir_ruta_db(self):
            ruta = filedialog.asksaveasfilename(
                title="Elegir la base existente o una nueva", defaultextension=".db",
                filetypes=[("Base de datos", "*.db")], initialfile="mantenimiento.db", confirmoverwrite=False)
            if ruta:
                self.cfg.set("base", "ruta", os.path.normpath(ruta))
                guardar_config(self.cfg)
            return ruta

        def cambiar_ubicacion(self):
            if self._elegir_ruta_db():
                messagebox.showinfo("Ubicación de la base",
                                    "Ubicación guardada en config.ini.\nCierre y vuelva a abrir el programa.")

        def _respaldo_automatico(self, ruta):
            try:
                hecho = respaldar(ruta, self.cfg.get("respaldo", "carpeta"),
                                  self.cfg.getint("respaldo", "dias", fallback=30))
                return f"Respaldo del día creado: {os.path.basename(hecho)}" if hecho else ""
            except Exception as ex:
                return f"⚠ No se pudo crear el respaldo diario: {ex}"

        def respaldar_ahora(self):
            try:
                hecho = respaldar(self.db.ruta, self.cfg.get("respaldo", "carpeta"), manual=True)
                messagebox.showinfo("Respaldo", f"Respaldo creado:\n{hecho}")
            except Exception as ex:
                messagebox.showerror("Respaldo", f"No se pudo crear el respaldo:\n{ex}")

        def _vigilar(self, seg):
            """Cada 'seg' segundos revisa si otra PC guardó cambios y refresca la lista."""
            try:
                dv = self.db.data_version()
                if dv != self.dv:
                    self.dv = dv
                    if not self.db.vacia:
                        self._cargar_combos()   # por si otro supervisor editó las listas
                    pestana = self.nb.index(self.nb.select())
                    if pestana == 1:
                        self.plan_refrescar()
                    elif pestana == 2:
                        self.maq_refrescar()
                    if self.consulta and len(self.filas_lista) <= LIMITE_REFRESCO_AUTO:
                        self.refrescar()
                    elif self.consulta:
                        self.cambios_pendientes = True
                        self._actualizar_estado()
                    self._revisar_cargado()
                self.conexion_ok = True
            except sqlite3.Error:
                self.conexion_ok = False
                self._actualizar_estado()
            self.after(seg * 1000, lambda: self._vigilar(seg))

        def _revisar_cargado(self):
            """Avisa si la actividad que se está editando fue modificada o eliminada por otro usuario."""
            if self.id_cargado is None:
                return
            r = self.db.obtener_rev(self.id_cargado)
            if r is None:
                self._ayuda("⚠ La actividad que está editando fue ELIMINADA por otro usuario.")
            elif (r[0] or 0) != self.rev_cargado:
                self._ayuda(f"⚠ Otro usuario modificó esta actividad ({r[1]}). "
                            "Si guarda, se le preguntará qué versión conservar.")

        def _actualizar_estado(self):
            partes = [f"{len(self.filas_lista)} actividades en la lista"]
            if getattr(self, "busq_total", None) is not None and self.consulta is not None:
                partes[0] = (f"Búsqueda: {self.busq_total} encontradas" +
                             (f" (se muestran las {len(self.filas_lista)} más recientes, agregue palabras para "
                              f"acotar)" if self.busq_total > len(self.filas_lista) else ""))
            if self.cambios_pendientes:
                partes.append("⚠ Hay cambios de otros usuarios: presione REFRESCAR")
            if getattr(self, "conexion_ok", True) is False:
                partes.append("⚠ SIN CONEXIÓN CON LA BASE")
            if self.ultimo_refresco:
                partes.append(f"Actualizado {self.ultimo_refresco:%H:%M:%S}")
            partes.append(f"Base: {self.db.ruta}")
            if self.msg_respaldo:
                partes.append(self.msg_respaldo)
            self.lbl_cuenta.config(text="   |   ".join(partes))

        # ---------------- Armado ----------------
        def _armar(self):
            menu = tk.Menu(self)
            archivo = tk.Menu(menu, tearoff=0)
            archivo.add_command(label="Importar desde Excel...", command=self.importar)
            archivo.add_command(label="Exportar lista actual a Excel...", command=lambda: self.exportar(False))
            archivo.add_command(label="Exportar base completa a Excel...", command=lambda: self.exportar(True))
            archivo.add_separator()
            archivo.add_command(label="Refrescar lista (F5)", command=self.refrescar)
            archivo.add_command(label="Respaldar ahora", command=self.respaldar_ahora)
            archivo.add_command(label="Ubicación de la base...", command=self.cambiar_ubicacion)
            archivo.add_separator()
            archivo.add_command(label="Editar listas (supervisor)...", command=self.editar_listas)
            archivo.add_command(label="Importar planificación desde CalendarioMTO (supervisor)...",
                                command=self.plan_importar)
            archivo.add_command(label="Generar planilla modelo (ítems de preventivo)...", command=self.generar_planilla)
            archivo.add_command(label="Importar ítems desde planilla (supervisor)...", command=self.importar_planilla)
            archivo.add_command(label="Configuración de planificación (supervisor)...", command=self.config_plan)
            archivo.add_separator()
            archivo.add_command(label="Salir", command=self.destroy)
            menu.add_cascade(label="Archivo", menu=archivo)
            self.config(menu=menu)
            # Para el operario solo quedan habilitados Refrescar y Salir; el resto es de supervisor
            self.menu_archivo = archivo
            self.menu_solo_sup = [i for i in range(archivo.index("end") + 1)
                                  if archivo.type(i) == "command"
                                  and archivo.entrycget(i, "label") not in ("Refrescar lista (F5)", "Salir")]

            self.nb = ttk.Notebook(self)
            self.nb.pack(fill="both", expand=True, pady=(4, 0))
            # Modo y botón de supervisor en la misma línea de las pestañas (no ocupan una fila aparte)
            barra = ttk.Frame(self)
            self.btn_sup = ttk.Button(barra, text="PASAR A SUP", command=self.cambiar_modo)
            self.btn_sup.pack(side="right")
            self.lbl_modo = ttk.Label(barra, text="Modo: OPERARIO", foreground="gray", font=("Segoe UI", 9, "bold"))
            self.lbl_modo.pack(side="right", padx=8)
            barra.place(in_=self.nb, relx=1.0, x=-10, y=0, anchor="ne")
            barra.lift()
            self.tab_act = ttk.Frame(self.nb)
            self.tab_plan = ttk.Frame(self.nb)
            self.nb.add(self.tab_act, text="  Actividades  ")
            self.nb.add(self.tab_plan, text="  Planificación  ")
            self.nb.bind("<<NotebookTabChanged>>", lambda e: self._al_cambiar_pestana())
            self._armar_cabecera()
            self._armar_lista()
            self._armar_formulario()
            self._armar_planificacion()
            self._armar_plan_maquina()
            self._activar_autocompletado()

        def _armar_cabecera(self):
            cab = ttk.Frame(self.tab_act, padding=(10, 8, 10, 0))
            cab.pack(fill="x")

            ttk.Label(cab, text="OPERARIO").grid(row=0, column=0, sticky="w")
            self.cmb_oper = ttk.Combobox(cab, width=22, state="readonly")
            self.cmb_oper.grid(row=1, column=0, padx=(0, 10))
            ttk.Label(cab, text="AÑO").grid(row=0, column=1)
            self.var_anio = tk.StringVar(value=str(date.today().year)[-2:])
            ttk.Entry(cab, width=5, textvariable=self.var_anio, justify="center").grid(row=1, column=1, padx=3)
            ttk.Label(cab, text="CICLO").grid(row=0, column=2)
            self.var_ciclo = tk.StringVar(value=ciclo_actual())
            ttk.Entry(cab, width=5, textvariable=self.var_ciclo, justify="center").grid(row=1, column=2, padx=3)
            ttk.Label(cab, text="CÓDIGO").grid(row=0, column=3)
            self.var_codigo = tk.StringVar()
            ttk.Entry(cab, width=12, textvariable=self.var_codigo, state="readonly",
                      justify="center").grid(row=1, column=3, padx=(3, 15))
            for v in (self.var_anio, self.var_ciclo):
                v.trace_add("write", lambda *a: self._actualizar_codigo())
            self.cmb_oper.bind("<<ComboboxSelected>>", lambda e: self._actualizar_codigo())

            botones = ttk.Frame(cab)
            botones.grid(row=0, column=4, rowspan=2, padx=5)
            for i, (txt, cmd) in enumerate([("BUSCAR", self.buscar), ("AGREGAR", self.agregar),
                                            ("EDITAR", self.editar), ("ELIMINAR", self.eliminar),
                                            ("LIMPIAR", self.limpiar), ("REFRESCAR", self.refrescar)]):
                ttk.Button(botones, text=txt, width=11, command=cmd).grid(row=i // 3, column=i % 3, padx=2, pady=2)
            self.bind("<F5>", lambda e: self.refrescar())

            # Filtros de supervisor (reemplazan L. CÓD / L. SEC / L. MÁQ / L. PROY)
            filt = ttk.LabelFrame(cab, text="Filtros (supervisor)", padding=5)
            filt.grid(row=0, column=5, rowspan=2, padx=10, sticky="n")
            ttk.Label(filt, text="Código").grid(row=0, column=0, sticky="w")
            self.f_codigo = ttk.Entry(filt, width=10)
            self.f_codigo.grid(row=1, column=0, padx=2)
            ttk.Label(filt, text="Sector").grid(row=0, column=1, sticky="w")
            self.f_sector = ttk.Combobox(filt, width=9, state="readonly")
            self.f_sector.grid(row=1, column=1, padx=2)
            self.f_sector.bind("<<ComboboxSelected>>", lambda e: self._maquinas_de(self.f_sector, self.f_maquina))
            ttk.Label(filt, text="Máquina").grid(row=0, column=2, sticky="w")
            self.f_maquina = ttk.Combobox(filt, width=10)
            self.f_maquina.grid(row=1, column=2, padx=2)
            ttk.Label(filt, text="Proyecto").grid(row=0, column=3, sticky="w")
            self.f_proyecto = ttk.Entry(filt, width=9)
            self.f_proyecto.grid(row=1, column=3, padx=2)
            b_filtrar = ttk.Button(filt, text="Filtrar", command=self.filtrar)
            b_filtrar.grid(row=1, column=4, padx=4)
            for w in (self.f_codigo, self.f_sector, self.f_maquina, self.f_proyecto, b_filtrar):
                self.solo_sup.append(w)
            self.f_codigo.bind("<FocusIn>", lambda e: self._ayuda("Filtra por OT que contenga el texto. 'TT' trae todo."))


        def _armar_lista(self):
            barra = ttk.Frame(self.tab_act, padding=(10, 6, 10, 2))
            barra.pack(fill="x")
            ttk.Label(barra, text="VISTA DE ACTIVIDADES").pack(side="left")
            ttk.Label(barra, text="Buscar en todas las actividades:").pack(side="left", padx=(25, 4))
            self.a_texto = ttk.Entry(barra, width=40)
            self.a_texto.pack(side="left")
            self.a_texto.bind("<KeyRelease>", lambda e: self._act_buscar_tecla(e))
            self.a_texto.bind("<Escape>", lambda e: self._act_buscar_limpiar())
            self.a_texto.bind("<FocusIn>", lambda e: self._ayuda(
                "Escriba palabras (ej: 'rodamiento imp3'): se listan las actividades que contengan TODAS, en "
                "cualquier campo (OT, operario, máquina, actividad, comentarios, repuestos...). Esc borra."), add="+")
            ttk.Button(barra, text="✕", width=3, command=self._act_buscar_limpiar).pack(side="left", padx=2)
            self.busq_total = None
            frm = ttk.Frame(self.tab_act, padding=(10, 0, 10, 0))
            frm.pack(fill="both", expand=True)
            self.tree = ttk.Treeview(frm, show="headings", selectmode="extended", height=4)
            sy = ttk.Scrollbar(frm, orient="vertical", command=self.tree.yview)
            sx = ttk.Scrollbar(frm, orient="horizontal", command=self.tree.xview)
            self.tree.configure(yscrollcommand=sy.set, xscrollcommand=sx.set)
            self.tree.grid(row=0, column=0, sticky="nsew")
            sy.grid(row=0, column=1, sticky="ns")
            sx.grid(row=1, column=0, sticky="ew")
            frm.rowconfigure(0, weight=1)
            frm.columnconfigure(0, weight=1)
            self.tree.tag_configure("extra", foreground="red")
            self.tree.tag_configure("dev", foreground="blue")
            self.tree.bind("<Double-1>", lambda e: self.cargar_seleccion())
            self.lbl_cuenta = ttk.Label(self.tab_act, text="", padding=(10, 0))
            self.lbl_cuenta.pack(anchor="w")

        def _armar_formulario(self):
            frm = ttk.Frame(self.tab_act, padding=(10, 4, 10, 8))
            frm.pack(fill="x")

            def entrada(padre, clave, fila, col, ancho=14, sup=False):
                ttk.Label(padre, text=etiquetas.get(clave, clave)).grid(row=fila, column=col, sticky="w", padx=4)
                w = ttk.Entry(padre, width=ancho)
                w.grid(row=fila + 1, column=col, sticky="we", padx=4, pady=(0, 4))
                self._registrar(clave, w, sup)
                return w

            def combo(padre, clave, fila, col, ancho=14, sup=False, editable=False):
                ttk.Label(padre, text=etiquetas.get(clave, clave)).grid(row=fila, column=col, sticky="w", padx=4)
                w = ttk.Combobox(padre, width=ancho, state="normal" if editable else "readonly")
                w.grid(row=fila + 1, column=col, sticky="we", padx=4, pady=(0, 4))
                self._registrar(clave, w, sup)
                return w

            def texto(padre, clave, fila, col, alto, ancho=48, colspan=1):
                ttk.Label(padre, text=etiquetas.get(clave, clave)).grid(row=fila, column=col, columnspan=colspan,
                                                                     sticky="w", padx=4)
                w = tk.Text(padre, height=alto, width=ancho, wrap="word", font=("Segoe UI", 9))
                w.grid(row=fila + 1, column=col, columnspan=colspan, sticky="we", padx=4, pady=(0, 4))
                self._registrar(clave, w, False)
                return w

            etiquetas = {ACTIVIDAD: "ACTIVIDAD A REALIZAR O REALIZADA (Detallar lo máximo posible)",
                         REPUESTOS: "REPUESTOS UTILIZADOS (ESCRIBIR SRH EN SU DEFECTO)",
                         PLANIF: "¿FUE PLANIFICADO?", MAQ_PAR: "¿MAQ PARADA?", TIPO_END: "EXTRAS/DEV",
                         TIPO_PCPG: "TIPO (PCPG)", CAT: "ID PROYECTO", HI: "HORA INICIO", HF: "HORA FIN",
                         SUBCONJ: "SUBCONJUNTO", MAQUINA: "MÁQUINA"}

            # Columna 1 y 2
            c1 = ttk.Frame(frm)
            c1.grid(row=0, column=0, sticky="n")
            entrada(c1, OT_SEC, 0, 0, sup=True)
            entrada(c1, OT_REPLAN, 2, 0, sup=True)
            entrada(c1, QREPLAN, 4, 0, sup=True)
            entrada(c1, FECHA, 6, 0)
            self.cmb_sector = combo(c1, SECTOR, 8, 0)
            self.cmb_maquina = combo(c1, MAQUINA, 0, 1, editable=True)
            combo(c1, CONJ, 2, 1, sup=True, editable=True)
            combo(c1, SUBCONJ, 4, 1, sup=True, editable=True)
            combo(c1, CODACT, 6, 1, sup=True, editable=True)
            combo(c1, PRIORIDAD, 8, 1)
            self.cmb_sector.bind("<<ComboboxSelected>>",
                                 lambda e: self._maquinas_de(self.cmb_sector, self.cmb_maquina, limpiar=True))

            # Columna central: textos
            c2 = ttk.Frame(frm)
            c2.grid(row=0, column=1, sticky="n", padx=6)
            texto(c2, ACTIVIDAD, 0, 0, 3, colspan=2)
            entrada(c2, MEJORAS, 2, 0, ancho=48).grid(columnspan=2)
            entrada(c2, REPUESTOS, 4, 0, ancho=34)
            combo(c2, MAQ_PAR, 4, 1, ancho=12)
            texto(c2, COMENT, 6, 0, 3, ancho=34)
            c2b = ttk.Frame(c2)
            c2b.grid(row=7, column=1, sticky="n")
            combo(c2b, TIPO_END, 0, 0, ancho=12)

            # Columna derecha
            c3 = ttk.Frame(frm)
            c3.grid(row=0, column=2, sticky="n")
            combo(c3, PLANIF, 0, 0, ancho=18)
            combo(c3, ESTADO, 2, 0, ancho=18)
            entrada(c3, CAT, 4, 0, ancho=18, sup=True)
            combo(c3, TIPO_PCPG, 6, 0, ancho=18, sup=True)
            entrada(c3, HI, 8, 0, ancho=18)
            entrada(c3, HF, 10, 0, ancho=18)

            # Ayuda + control de actividad
            c4 = ttk.Frame(frm)
            c4.grid(row=0, column=3, sticky="nw", padx=(15, 0))
            self.lbl_ayuda = ttk.Label(c4, text="Espacio para mensajes de ayuda", wraplength=240,
                                       justify="left", foreground="#333")
            self.lbl_ayuda.grid(row=0, column=0, columnspan=3, sticky="nw", pady=(0, 10))
            c4.rowconfigure(0, minsize=100)
            ttk.Label(c4, text="CONTROL DE ACTIVIDAD").grid(row=1, column=0, columnspan=3, sticky="w")
            self.cmb_control = ttk.Combobox(c4, width=12, state="readonly")
            self.cmb_control.grid(row=2, column=0, columnspan=2, sticky="we")
            b_ind = ttk.Button(c4, text="OK IND", command=lambda: self.control(False))
            b_ind.grid(row=3, column=0, sticky="w", pady=3)
            b_todo = ttk.Button(c4, text="OK TODO", command=lambda: self.control(True))
            b_todo.grid(row=3, column=1, sticky="w", padx=3, pady=3)
            self.solo_sup += [self.cmb_control, b_ind, b_todo]
            ttk.Label(c4, text="Todas las actividades deben ser evaluadas por un supervisor.",
                      wraplength=240, foreground="gray").grid(row=4, column=0, columnspan=3, sticky="w", pady=4)

            self.cmb_oper.bind("<FocusIn>", lambda e: self._ayuda("Escriba o seleccione su nombre."), add="+")

        def _activar_autocompletado(self):
            """Se aplica al final del armado, después de todos los bind de los desplegables."""
            for w in (self.cmb_oper, self.cmb_sector, self.cmb_control, self.f_sector):
                autocompletar(w, True, self._ayuda)
            for clave in (PRIORIDAD, ESTADO, PLANIF, TIPO_END, MAQ_PAR, TIPO_PCPG):
                autocompletar(self.campos[clave], True, self._ayuda)
            for w in (self.cmb_maquina, self.f_maquina):
                autocompletar(w, False)
            for w in (self.p_sector, self.p_maquina, self.p_asig, self.p_estado, self.m_sector, self.m_maquina):
                autocompletar(w, True)

        def _registrar(self, clave, w, sup):
            self.campos[clave] = w
            if sup:
                self.solo_sup.append(w)
            if clave in AYUDA:
                w.bind("<FocusIn>", lambda e, c=clave: self._ayuda(AYUDA[c]), add="+")
            if clave == FECHA:
                w.bind("<FocusOut>", lambda e: self._normalizar(FECHA, fecha_a_db, "Ingrese una fecha válida (ej: 10-06-2025)."))
            if clave in (HI, HF):
                w.bind("<FocusOut>", lambda e, c=clave: self._normalizar(c, hora_a_db, "Ingrese una hora válida (ej: 08:00 o 14:30)."))

        # ---------------- Utilidades de interfaz ----------------
        def _ayuda(self, texto):
            self.lbl_ayuda.config(text=texto)

        def _get(self, clave):
            w = self.campos[clave]
            return w.get("1.0", "end").strip() if isinstance(w, tk.Text) else w.get().strip()

        def _set(self, clave, valor):
            w = self.campos[clave]
            valor = "" if valor is None else str(valor)
            if isinstance(w, tk.Text):
                w.delete("1.0", "end")
                w.insert("1.0", valor)
            elif isinstance(w, ttk.Combobox):
                estado = w.cget("state")
                w.config(state="normal")
                w.set(valor)
                w.config(state=estado)
            else:
                estado = w.cget("state")
                w.config(state="normal")
                w.delete(0, "end")
                w.insert(0, valor)
                w.config(state=estado)

        def _normalizar(self, clave, func, msg):
            t = self._get(clave)
            if not t:
                return
            v = func(t)
            if v is None:
                messagebox.showwarning("Dato inválido", msg)
                self.after(10, self.campos[clave].focus_set)
            else:
                self._set(clave, db_a_texto(v))

        def _actualizar_codigo(self):
            self.var_codigo.set(armar_codigo(self.var_anio.get(), self.var_ciclo.get(), self.cmb_oper.get()))

        def _maquinas_de(self, cmb_sec, cmb_maq, limpiar=False):
            cmb_maq["values"] = self.db.lista("maquinas", cmb_sec.get())
            if limpiar:
                cmb_maq.set("")

        def _cargar_combos(self):
            """Carga las listas desplegables desde la base (se llama también cuando otra PC las edita)."""
            self.cmb_oper["values"] = self.db.lista("operarios")
            sectores = self.db.lista("sectores")
            self.cmb_sector["values"] = sectores
            self.f_sector["values"] = [""] + sectores
            for clave, lista in [(ESTADO, "estado"), (PRIORIDAD, "prioridad"), (PLANIF, "planificado"),
                                 (TIPO_END, "extras"), (MAQ_PAR, "maq_parada"), (TIPO_PCPG, "tipo_pcpg")]:
                self.campos[clave]["values"] = [""] + self.db.lista(lista)
            self.cmb_control["values"] = [""] + self.db.lista("supervisores")
            self.p_sector["values"] = [""] + sectores
            self.p_asig["values"] = [""] + self.db.lista("operarios")
            self.m_sector["values"] = sectores
            self._plan_maquinas_filtro()
            if self.cmb_sector.get():
                self.cmb_maquina["values"] = self.db.lista("maquinas", self.cmb_sector.get())

        def editar_listas(self):
            if not self.supervisor:
                messagebox.showwarning("Editar listas", "Solo un supervisor puede editar las listas.")
                return
            if self.db.vacia:
                return
            EditorListas(self)

        def _cargar_listas(self):
            self._cargar_combos()
            # Columnas de la lista: las 31 de bbdd
            cols = self.db.encabezados
            self.tree["columns"] = cols
            anchos = {ACTIVIDAD: 300, COMENT: 200, REPUESTOS: 160, MEJORAS: 160, MAQUINA: 90, OPERARIO: 85,
                      FECHA: 90, HI: 50, HF: 50, TIPO_PCPG: 110, CONTROL: 110}
            for c in cols:
                self.tree.heading(c, text=c)
                self.tree.column(c, width=anchos.get(c, 70), stretch=False, anchor="w")

        def _aplicar_permisos(self):
            for w in self.solo_sup:
                if isinstance(w, ttk.Combobox):
                    editable = getattr(w, "_auto", False) or w in (self.f_maquina,) or w in (
                        self.campos.get(CONJ), self.campos.get(SUBCONJ), self.campos.get(CODACT))
                    w.config(state=("normal" if editable else "readonly") if self.supervisor else "disabled")
                else:
                    w.config(state="normal" if self.supervisor else "disabled")
            self.btn_sup.config(text="VOLVER A OP" if self.supervisor else "PASAR A SUP")
            self.lbl_modo.config(text="Modo: SUPERVISOR" if self.supervisor else "Modo: OPERARIO",
                                 foreground="darkred" if self.supervisor else "gray")
            for i in self.menu_solo_sup:
                self.menu_archivo.entryconfig(i, state="normal" if self.supervisor else "disabled")
            for tab in (self.tab_plan, self.tab_maq):   # Planificación y Plan por máquina: solo supervisor
                self.nb.tab(tab, state="normal" if self.supervisor else "hidden")
            if not self.supervisor:
                self.nb.select(self.tab_act)
                for w in self.winfo_children():   # cierra las ventanas de supervisor que hubieran quedado abiertas
                    if isinstance(w, tk.Toplevel):
                        w.destroy()

        def _mostrar(self, filas, conservar=False):
            sel = self.tree.selection() if conservar else ()
            vista = self.tree.yview()[0] if conservar else 0
            self.filas_lista = filas
            self.tree.delete(*self.tree.get_children())
            i_tipo = self.db.encabezados.index(TIPO_END) + 1
            for f in filas:
                tipo = str(f[i_tipo] or "").strip().lower()
                tag = ("extra",) if tipo == "extra" else ("dev",) if tipo == "dev" else ()
                self.tree.insert("", "end", iid=str(f[0]), values=[db_a_texto(v) for v in f[1:]], tags=tag)
            if conservar:
                existentes = [i for i in sel if self.tree.exists(i)]
                if existentes:
                    self.tree.selection_set(existentes)
                self.tree.yview_moveto(vista)
            self._actualizar_estado()

        def refrescar(self):
            """Vuelve a ejecutar la última búsqueda/filtro: muestra altas, cambios y bajas de todas las PC."""
            if self.consulta is None:
                return
            try:
                filas = self.consulta()
            except sqlite3.Error as ex:
                self.conexion_ok = False
                self._actualizar_estado()
                messagebox.showerror("Base de datos", f"No se pudo leer la base:\n{ex}")
                return
            self.conexion_ok = True
            self.cambios_pendientes = False
            self.ultimo_refresco = datetime.now()
            self.dv = self.db.data_version()
            self._mostrar(filas, conservar=True)

        def _limpiar_campos(self):
            for clave in self.campos:
                self._set(clave, "")
            self.cmb_maquina["values"] = []
            self.id_cargado = None
            self.rev_cargado = None
            self.vinculo_pendiente = None
            self.cmb_oper.focus_set()

        def _datos_formulario(self):
            d = {k: self._get(k) for k in self.campos}
            d[OPERARIO] = self.cmb_oper.get().strip()
            return d

        def _a_guardar(self, d):
            """Convierte los datos del formulario al formato de la base."""
            out = {}
            for k, v in d.items():
                if k == FECHA:
                    iso = fecha_a_db(v)
                    if iso and iso >= "1900":
                        out[k] = iso
                        out[ANIO] = int(iso[:4])
                    elif iso:  # fechas especiales (11/11/1111 = actividad semanal): se guardan como texto
                        out[k] = f"{iso[8:10]}/{iso[5:7]}/{iso[:4]}"
                    else:
                        out[k] = v or None
                elif k in (HI, HF):
                    out[k] = hora_a_db(v) or (v or None)
                else:
                    out[k] = v or None
            return out

        # ---------------- Acciones ----------------
        def buscar(self):
            if not self.cmb_oper.get():
                messagebox.showwarning("Atención", "Nombre de operario en blanco.")
                return
            codigo = self.var_codigo.get()
            if len(codigo) < 6:
                messagebox.showwarning("Atención", "Faltan datos del usuario (año, ciclo u operario).")
                return
            prefijo = codigo[:6]
            self.a_texto.delete(0, "end")
            self.busq_total = None
            self.consulta = lambda: self.db.por_codigo(prefijo)
            self.refrescar()
            if not self.filas_lista:
                messagebox.showinfo("Buscar", "No se encontraron coincidencias.")

        def agregar(self):
            if self.db.vacia:
                return
            d = self._datos_formulario()
            err = validar(d, self.supervisor)
            if err:
                messagebox.showwarning("Falta dato", err)
                return
            codigo = self.var_codigo.get()
            if len(codigo) < 6:
                messagebox.showwarning("Atención", "No se pudo armar el código (año, ciclo y operario).")
                return
            datos = self._a_guardar(d)
            datos[SEM] = self.var_ciclo.get().strip()
            datos.setdefault(ANIO, 2000 + int(self.var_anio.get()[-2:]))
            try:
                ot = self.db.insertar_nueva(codigo, datos)
            except sqlite3.OperationalError as ex:
                messagebox.showerror("No se pudo guardar",
                                     f"La base está ocupada o no accesible. Intente de nuevo.\n\n{ex}")
                return
            vp = self.vinculo_pendiente
            if vp and datos.get(OT_SEC) == vp[2]:   # vino de Planificación: queda vinculada a esa fecha
                self.db.vincular(self.db.ultimo_id, vp[0], vp[1])
            if self.consulta is None:  # sin lista activa: mostrar las actividades del operario en el ciclo
                prefijo = codigo[:6]
                self.consulta = lambda: self.db.por_codigo(prefijo)
            self.refrescar()
            nuevo = self.db.consultar(f"{self.db.col(OT)} = ?", (ot,))
            if nuevo and not self.tree.exists(str(nuevo[0][0])):  # no entra en el filtro actual: se agrega igual
                self._mostrar(self.filas_lista + nuevo, conservar=True)
            if nuevo:
                self.tree.see(str(nuevo[0][0]))
            self._limpiar_campos()
            self._ayuda(f"Actividad {ot} guardada.")

        def cargar_seleccion(self):
            sel = self.tree.selection()
            if sel:
                self._cargar_id(int(sel[0]))

        def _cargar_id(self, id_):
            fila = self.db.obtener(id_)
            if not fila:
                messagebox.showwarning("Atención", "La actividad ya no existe (fue eliminada por otro usuario).")
                self.refrescar()
                return
            enc = self.db.encabezados
            self._limpiar_campos()
            self.id_cargado = fila[0]
            self.rev_cargado = (self.db.obtener_rev(id_) or (0,))[0] or 0
            sector = fila[enc.index(SECTOR) + 1]
            self._set(SECTOR, sector)
            self.cmb_maquina["values"] = self.db.lista("maquinas", sector or "")
            for clave in self.campos:
                if clave in enc and clave != SECTOR:
                    self._set(clave, db_a_texto(fila[enc.index(clave) + 1]))
            self.cmb_control.set(db_a_texto(fila[enc.index(CONTROL) + 1]))
            self._ayuda(f"Editando OT {fila[enc.index(OT) + 1]}. Modifique y presione EDITAR.")

        def editar(self):
            if self.id_cargado is None:
                messagebox.showwarning("Atención", "Haga doble clic en una fila de la lista para cargarla.")
                return
            d = self._datos_formulario()
            err = validar(d, self.supervisor)
            if err:   # se valida ANTES de modificar (el VBA lo hacía después)
                messagebox.showwarning("Falta dato", err)
                return
            datos = self._a_guardar(d)
            if not self.supervisor:  # el operario no modifica campos de supervisor
                for k in (OT_SEC, OT_REPLAN, QREPLAN, CONJ, SUBCONJ, CODACT, CAT, TIPO_PCPG):
                    datos.pop(k, None)
            try:
                res = self.db.actualizar(self.id_cargado, datos, self.rev_cargado)
            except sqlite3.OperationalError as ex:
                messagebox.showerror("No se pudo guardar",
                                     f"La base está ocupada o no accesible. Intente de nuevo.\n\n{ex}")
                return
            if res == "eliminado":
                messagebox.showwarning("Atención", "No se guardó: la actividad fue eliminada por otro usuario.")
                self._limpiar_campos()
                self.refrescar()
                return
            if res == "conflicto":
                info = (self.db.obtener_rev(self.id_cargado) or (None, "?"))[1]
                r = messagebox.askyesnocancel(
                    "Actividad modificada por otro usuario",
                    f"Otro usuario modificó esta actividad mientras usted la editaba ({info}).\n\n"
                    "Sí: guardar SUS cambios igual (reemplaza los del otro usuario).\n"
                    "No: descartar sus cambios y cargar la versión actual.\n"
                    "Cancelar: volver al formulario sin guardar.", icon="warning")
                if r is None:
                    return
                if r is False:
                    self._cargar_id(self.id_cargado)
                    self.refrescar()
                    return
                self.db.actualizar(self.id_cargado, datos)
            self.refrescar()
            self._limpiar_campos()
            self._ayuda("Cambios guardados.")

        def eliminar(self):
            sel = self.tree.selection()
            if not sel:
                messagebox.showwarning("Atención", "Seleccione una fila de la lista para eliminar.")
                return
            if not self.supervisor:
                messagebox.showwarning("Atención", "Debe ser supervisor para eliminar de la base de datos.")
                return
            id_ = int(sel[0])
            ot = self.tree.set(sel[0], OT)
            if not messagebox.askyesno("Confirmar eliminación",
                                       f"¿Está seguro que desea eliminar la OT '{ot}'?\nEsta acción no se puede deshacer.",
                                       icon="warning", default="no"):
                return
            self.db.eliminar(id_)
            self._mostrar([f for f in self.filas_lista if f[0] != id_], conservar=True)
            self._limpiar_campos()

        def _act_buscar_tecla(self, e=None):
            """Busca 300 ms después de la última tecla (no en cada tecla)."""
            if e is not None and e.keysym in ("Up", "Down", "Left", "Right", "Shift_L", "Shift_R", "Control_L",
                                              "Control_R", "Home", "End", "Tab", "Escape"):
                return
            if getattr(self, "_a_busqueda", None):
                self.after_cancel(self._a_busqueda)
            self._a_busqueda = self.after(300, self._act_buscar)

        def _act_buscar(self):
            self._a_busqueda = None
            texto = self.a_texto.get().strip()
            if len(texto.replace(" ", "")) < 2:
                if not texto:
                    self._act_buscar_limpiar()
                return
            limite = 500

            def consulta():
                self.busq_total, filas = self.db.buscar_texto(texto, limite)
                return filas
            self.consulta = consulta
            self.refrescar()
            self.tree.yview_moveto(0)

        def _act_buscar_limpiar(self):
            self.a_texto.delete(0, "end")
            if self.busq_total is not None:
                self.busq_total = None
                self.consulta = None
                self._mostrar([])

        def limpiar(self):
            self.a_texto.delete(0, "end")
            self.busq_total = None
            self.consulta = None
            self.cambios_pendientes = False
            self._mostrar([])
            self._limpiar_campos()
            self._ayuda("Espacio para mensajes de ayuda")

        def filtrar(self):
            codigo = self.f_codigo.get().strip()
            if not any([codigo, self.f_sector.get(), self.f_maquina.get().strip(), self.f_proyecto.get().strip()]):
                messagebox.showinfo("Filtros", "Complete al menos un filtro ('TT' en Código trae todo).")
                return
            self.config(cursor="watch")
            self.update()
            args = (codigo, self.f_sector.get(), self.f_maquina.get().strip(), self.f_proyecto.get().strip())
            self.a_texto.delete(0, "end")
            self.busq_total = None
            try:
                self.consulta = lambda: self.db.filtrar(*args)
                self.refrescar()
                self.tree.yview_moveto(0)
            finally:
                self.config(cursor="")

        def control(self, todo):
            sup = self.cmb_control.get()
            if not sup:
                messagebox.showwarning("Atención", "Debe seleccionar un supervisor.")
                return
            if todo:
                ids = [f[0] for f in self.filas_lista]
                if not ids:
                    messagebox.showwarning("Atención", "No hay filas en la lista.")
                    return
                if not messagebox.askyesno("Confirmar", f"¿Marcar las {len(ids)} actividades de la lista como "
                                                        f"controladas por {sup}?"):
                    return
            else:
                ids = [int(i) for i in self.tree.selection()]
                if not ids:
                    messagebox.showwarning("Atención", "Debe seleccionar una fila.")
                    return
            self.db.actualizar_campo(ids, CONTROL, sup)
            self.refrescar()
            self._ayuda(f"Control de actividad actualizado en {len(ids)} fila(s).")

        def cambiar_modo(self):
            if self.supervisor:
                self.supervisor = False
            else:
                if not self.db.hay_clave():
                    c1 = simpledialog.askstring("Contraseña de supervisor",
                                                "No hay contraseña definida.\nIngrese una nueva:", show="*", parent=self)
                    if not c1:
                        return
                    c2 = simpledialog.askstring("Contraseña de supervisor", "Repítala:", show="*", parent=self)
                    if c1 != c2:
                        messagebox.showerror("Error", "Las contraseñas no coinciden.")
                        return
                    self.db.fijar_clave(c1)
                else:
                    c = simpledialog.askstring("Acceso restringido", "Ingrese la contraseña de supervisor:",
                                               show="*", parent=self)
                    if c is None:
                        return
                    if not self.db.verificar_clave(c):
                        messagebox.showerror("Error", "Contraseña incorrecta. Permanece en modo OPERARIO.")
                        return
                self.supervisor = True
            self._aplicar_permisos()

        # ---------------- Planificación ----------------
        def _armar_planificacion(self):
            p = self.tab_plan
            hoy = date.today()
            self.p_modo, self.p_sem_ini = "mes", hoy - timedelta(days=hoy.weekday())

            # Fila 1: período, vista y título
            per = ttk.Frame(p, padding=(10, 8, 10, 2))
            per.pack(fill="x")
            ttk.Button(per, text="◀", width=3, command=lambda: self._plan_mover(-1)).pack(side="left")
            self.p_mes = ttk.Combobox(per, width=11, state="readonly", values=MESES)
            self.p_mes.current(hoy.month - 1)
            self.p_mes.pack(side="left", padx=2)
            self.p_anio = ttk.Spinbox(per, from_=2020, to=2100, width=6, command=self._plan_modo_mes)
            self.p_anio.set(hoy.year)
            self.p_anio.pack(side="left", padx=2)
            ttk.Button(per, text="▶", width=3, command=lambda: self._plan_mover(1)).pack(side="left")
            ttk.Button(per, text="Esta semana", command=self._plan_semana).pack(side="left", padx=(6, 0))
            ttk.Button(per, text="Este mes", command=self._plan_hoy).pack(side="left", padx=2)
            self.p_mes.bind("<<ComboboxSelected>>", lambda e: self._plan_modo_mes())
            self.p_anio.bind("<Return>", lambda e: self._plan_modo_mes())
            ttk.Label(per, text="Vista:").pack(side="left", padx=(20, 4))
            self.p_vista = tk.StringVar(value="lista")
            for txt, val in (("Lista", "lista"), ("Calendario", "calendario")):
                ttk.Radiobutton(per, text=txt, value=val, variable=self.p_vista,
                                command=self._plan_cambiar_vista).pack(side="left", padx=2)
            ttk.Button(per, text="Resumen por operario", command=lambda: ResumenOperarios(self)).pack(side="left", padx=(15, 0))
            self.p_titulo = ttk.Label(per, text="", font=("Segoe UI", 11, "bold"))
            self.p_titulo.pack(side="right", padx=10)

            # Fila 2: filtros
            fil = ttk.Frame(p, padding=(10, 2, 10, 4))
            fil.pack(fill="x")
            ttk.Label(fil, text="Sector").pack(side="left", padx=(0, 2))
            self.p_sector = ttk.Combobox(fil, width=9, state="readonly")
            self.p_sector.pack(side="left")
            ttk.Label(fil, text="Máquina").pack(side="left", padx=(10, 2))
            self.p_maquina = ttk.Combobox(fil, width=12, state="readonly")
            self.p_maquina.pack(side="left")
            ttk.Label(fil, text="Asignado").pack(side="left", padx=(10, 2))
            self.p_asig = ttk.Combobox(fil, width=12, state="readonly")
            self.p_asig.pack(side="left")
            ttk.Label(fil, text="Estado").pack(side="left", padx=(10, 2))
            self.p_estado = ttk.Combobox(fil, width=12, state="readonly", values=[""] + ESTADOS_PLAN)
            self.p_estado.pack(side="left")
            ttk.Label(fil, text="Buscar").pack(side="left", padx=(10, 2))
            self.p_texto = ttk.Entry(fil, width=18)
            self.p_texto.pack(side="left")
            self.p_texto.bind("<KeyRelease>", lambda e: self._plan_buscar_tecla())
            self.p_texto.bind("<Escape>", lambda e: (self.p_texto.delete(0, "end"), self._plan_mostrar()))
            self.p_atras = tk.BooleanVar(value=True)
            ttk.Checkbutton(fil, text="Incluir atrasadas", variable=self.p_atras,
                            command=self.plan_refrescar).pack(side="left", padx=10)
            ttk.Button(fil, text="Limpiar filtros", command=self._plan_limpiar_filtros).pack(side="left")
            self.p_sector.bind("<<ComboboxSelected>>", lambda e: (self._plan_maquinas_filtro(), self._plan_mostrar()))
            for c in (self.p_maquina, self.p_asig, self.p_estado):
                c.bind("<<ComboboxSelected>>", lambda e: self._plan_mostrar())

            # Fila 3: acciones
            acc = ttk.Frame(p, padding=(10, 0, 10, 4))
            acc.pack(fill="x")
            ttk.Button(acc, text="CARGAR COMO ACTIVIDAD", command=self.plan_cargar_actividad).pack(side="left")
            ttk.Button(acc, text="Ver histórico / checklist", command=self.plan_abrir).pack(side="left", padx=4)
            ttk.Separator(acc, orient="vertical").pack(side="left", fill="y", padx=8)
            for txt, cmd in [("Nueva tarea", lambda: FormTarea(self)), ("Editar tarea", self.plan_editar),
                             ("Marca manual ▾", self.plan_marca), ("Eliminar tarea", self.plan_eliminar)]:
                b = ttk.Button(acc, text=txt, command=cmd)
                b.pack(side="left", padx=2)
                self.solo_sup.append(b)
            ttk.Button(acc, text="Exportar a Excel", command=self.plan_exportar).pack(side="right")
            ttk.Button(acc, text="REFRESCAR", command=self.plan_refrescar).pack(side="right", padx=4)

            # Vista lista
            self.p_frm_lista = frm = ttk.Frame(p, padding=(10, 0, 10, 0))
            frm.pack(fill="both", expand=True)
            cols = [("fecha", "Programada", 92), ("sem", "Sem", 38), ("ot_base", "OT Base", 145), ("tipo", "Tipo", 42),
                    ("sector", "Sector", 60), ("maquina", "Máquina", 80), ("tarea", "Tarea", 300),
                    ("asignado", "Asignado", 85), ("periodicidad", "Periodicidad", 100), ("estado", "Estado", 85),
                    ("realizada", "Realizada", 90), ("retraso", "Retraso", 105), ("ot_man", "OT MAN", 120),
                    ("detalle", "Último estado / comentario", 240)]
            self.p_tree = ttk.Treeview(frm, columns=[c[0] for c in cols], show="headings", selectmode="browse",
                                       height=6)
            for c, t, w in cols:
                self.p_tree.heading(c, text=t)
                self.p_tree.column(c, width=w, stretch=c in ("tarea", "detalle"), anchor="w")
            sy = ttk.Scrollbar(frm, orient="vertical", command=self.p_tree.yview)
            sx = ttk.Scrollbar(frm, orient="horizontal", command=self.p_tree.xview)
            self.p_tree.configure(yscrollcommand=sy.set, xscrollcommand=sx.set)
            self.p_tree.grid(row=0, column=0, sticky="nsew")
            sy.grid(row=0, column=1, sticky="ns")
            sx.grid(row=1, column=0, sticky="ew")
            frm.rowconfigure(0, weight=1)
            frm.columnconfigure(0, weight=1)
            for est, color in COLORES_ESTADO.items():
                self.p_tree.tag_configure(est, foreground=color)
            self.p_tree.bind("<Double-1>", lambda e: self.plan_abrir())

            # Vista calendario
            self.p_frm_cal = ttk.Frame(p, padding=(10, 0, 10, 0))
            self.p_resumen = ttk.Label(p, text="", padding=(10, 4))
            self.p_resumen.pack(anchor="w")

        def _pestana_plan(self):
            return self.nb.index(self.nb.select()) == 1

        def _al_cambiar_pestana(self):
            i = self.nb.index(self.nb.select())
            if i == 1:
                self.plan_refrescar()
            elif i == 2:
                self.maq_refrescar()

        def _plan_rango(self):
            if self.p_modo == "semana":
                return self.p_sem_ini, self.p_sem_ini + timedelta(days=6)
            y, m = int(self.p_anio.get()), self.p_mes.current() + 1
            return date(y, m, 1), date(y, m, calendar.monthrange(y, m)[1])

        def _plan_modo_mes(self):
            self.p_modo = "mes"
            self.plan_refrescar()

        def _plan_semana(self):
            hoy = date.today()
            self.p_modo, self.p_sem_ini = "semana", hoy - timedelta(days=hoy.weekday())
            self.plan_refrescar()

        def _plan_hoy(self):
            hoy = date.today()
            self.p_mes.current(hoy.month - 1)
            self.p_anio.set(hoy.year)
            self._plan_modo_mes()

        def _plan_mover(self, delta):
            if self.p_modo == "semana":
                self.p_sem_ini += timedelta(days=7 * delta)
            else:
                m = self.p_mes.current() + delta
                y = int(self.p_anio.get()) + (m // 12)
                self.p_mes.current(m % 12)
                self.p_anio.set(y)
            self.plan_refrescar()

        def _plan_maquinas_filtro(self):
            """Máquinas del filtro: las de la lista del sector más las que tengan tareas cargadas."""
            sec = self.p_sector.get()
            maqs = set(self.db.lista("maquinas", sec)) if sec else set()
            maqs |= {t["maquina"] for t in self.db.tareas() if t["maquina"] and (not sec or t["sector"] == sec)}
            self.p_maquina["values"] = [""] + sorted(maqs)
            if self.p_maquina.get() not in maqs:
                self.p_maquina.set("")

        def _plan_limpiar_filtros(self):
            for c in (self.p_sector, self.p_maquina, self.p_asig, self.p_estado):
                c.set("")
            self.p_texto.delete(0, "end")
            self._plan_maquinas_filtro()
            self._plan_mostrar()

        def _plan_cambiar_vista(self):
            if self.p_vista.get() == "calendario":
                self.p_frm_lista.pack_forget()
                self.p_frm_cal.pack(fill="both", expand=True, before=self.p_resumen)
            else:
                self.p_frm_cal.pack_forget()
                self.p_frm_lista.pack(fill="both", expand=True, before=self.p_resumen)
            self._plan_mostrar()

        def plan_refrescar(self):
            try:
                desde, hasta = self._plan_rango()
            except ValueError:
                return
            try:
                self.plan_base = calcular_plan(self.db, desde, hasta, self.p_atras.get())
            except sqlite3.Error as ex:
                self.p_resumen.config(text=f"⚠ No se pudo leer la base: {ex}")
                return
            self._plan_mostrar()

        def _plan_buscar_tecla(self):
            """Búsqueda dinámica: filtra en memoria 250 ms después de la última tecla (no relee la base)."""
            if getattr(self, "_p_busqueda", None):
                self.after_cancel(self._p_busqueda)
            self._p_busqueda = self.after(250, self._plan_mostrar)

        def _tolerancia(self):
            return self.cfg.getint("plan", "tolerancia_dias", fallback=3)

        def _plan_mostrar(self):
            self._p_busqueda = None
            desde, hasta = self._plan_rango()
            filas = getattr(self, "plan_base", [])
            sec, maq, asig, est = self.p_sector.get(), self.p_maquina.get(), self.p_asig.get(), self.p_estado.get()
            if sec:
                filas = [f for f in filas if f["sector"] == sec]
            if maq:
                filas = [f for f in filas if f["maquina"] == maq]
            if asig:
                filas = [f for f in filas if f["asignado"] == asig]
            if est:
                filas = [f for f in filas if f["estado"] == est]
            palabras = self.p_texto.get().lower().split()
            if palabras:   # todas las palabras deben aparecer, en cualquier campo y en cualquier orden
                campos = ("ot_base", "tipo", "sector", "maquina", "tarea", "asignado", "estado", "ot_man", "detalle")
                filas = [f for f in filas
                         if all(p in " ".join(str(f[k] or "") for k in campos).lower() for p in palabras)]
            self.plan_filas = filas
            tol = self._tolerancia()
            sel = self.p_tree.selection()
            vista = self.p_tree.yview()[0]
            self.p_tree.delete(*self.p_tree.get_children())
            for f in filas:
                d = f["fecha"]
                tag = f["estado"]
                if f["estado"] == "Cumplida" and f["retraso"] is not None and f["retraso"] > tol:
                    tag = "Cumplida tarde"
                self.p_tree.insert("", "end", iid=f"{f['id']}|{f['fecha_occ']}", tags=(tag,), values=[
                    d.strftime("%d-%m-%Y") if d else "Sin fecha", semana_excel(d) if d else "", f["ot_base"] or "",
                    f["tipo"] or "", f["sector"] or "", f["maquina"] or "", f["tarea"] or "", f["asignado"] or "",
                    f["periodicidad"], f["estado"], f["realizada"].strftime("%d-%m-%Y") if f["realizada"] else "",
                    texto_retraso(f), f["ot_man"], f["detalle"]])
            sel = [i for i in sel if self.p_tree.exists(i)]
            if sel:
                self.p_tree.selection_set(sel)
            self.p_tree.yview_moveto(vista)
            if self.p_vista.get() == "calendario":
                self._plan_dibujar_calendario(filas, desde, hasta)
            if self.p_modo == "semana":
                self.p_titulo.config(text=f"Semana {desde:%d/%m} al {hasta:%d/%m/%Y}")
            else:
                self.p_titulo.config(text=f"{MESES[desde.month - 1]} {desde.year}")
            self.p_resumen.config(text=resumen_plan(filas, tol))

        def _plan_dibujar_calendario(self, filas, desde, hasta):
            for w in self.p_frm_cal.winfo_children():
                w.destroy()
            hoy = date.today()
            inicio = desde - timedelta(days=desde.weekday())
            fin = hasta + timedelta(days=6 - hasta.weekday())
            por_dia = {}
            fuera = 0
            for f in filas:
                if f["fecha"] and inicio <= f["fecha"] <= fin:
                    por_dia.setdefault(f["fecha"], []).append(f)
                else:
                    fuera += 1
            if fuera:
                ttk.Label(self.p_frm_cal, text=f"{fuera} tareas fuera del calendario (atrasadas de períodos anteriores o "
                                               "sin fecha): se ven en la vista Lista.", foreground="#c00000").grid(
                    row=0, column=0, columnspan=7, sticky="w", pady=(0, 4))
            for c, nombre in enumerate(["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]):
                ttk.Label(self.p_frm_cal, text=nombre, anchor="center", font=("Segoe UI", 9, "bold")).grid(
                    row=1, column=c, sticky="we")
                self.p_frm_cal.columnconfigure(c, weight=1, uniform="dias")
            d, i = inicio, 0
            while d <= fin:
                r, c = 2 + i // 7, i % 7
                del_periodo = desde <= d <= hasta
                fondo = "#fff8dc" if d == hoy else ("white" if del_periodo else "#ececec")
                celda = tk.Frame(self.p_frm_cal, bg=fondo, highlightbackground="#b0b0b0", highlightthickness=1)
                celda.grid(row=r, column=c, sticky="nsew", padx=1, pady=1)
                tareas = por_dia.get(d, [])
                tk.Label(celda, text=f"{d.day}" + (f"   ({len(tareas)})" if tareas else ""), bg=fondo, anchor="w",
                         font=("Segoe UI", 9, "bold" if d == hoy else "normal"),
                         fg="black" if del_periodo else "gray").pack(fill="x", padx=3)
                lb = tk.Listbox(celda, height=4, bd=0, highlightthickness=0, activestyle="none", font=("Segoe UI", 8),
                                bg=fondo, exportselection=False)
                lb.pack(fill="both", expand=True, padx=2, pady=(0, 2))
                for f in tareas:
                    lb.insert("end", f"{f['maquina'] or f['sector'] or ''}: {f['tarea'] or ''}")
                    lb.itemconfig("end", fg=COLORES_ESTADO.get(f["estado"], "black"))
                lb.bind("<Double-1>", lambda e, lst=tareas, w=lb: lst and w.curselection() and
                        self._plan_abrir_fila(lst[w.curselection()[0]]))
                d += timedelta(days=1)
                i += 1
            for r in range(2, 2 + (i + 6) // 7):
                self.p_frm_cal.rowconfigure(r, weight=1, uniform="semanas")

        def _plan_sel(self):
            sel = self.p_tree.selection()
            if not sel:
                messagebox.showinfo("Planificación", "Seleccione una tarea de la lista.")
                return None
            return next((f for f in self.plan_filas if f"{f['id']}|{f['fecha_occ']}" == sel[0]), None)

        def plan_cargar_actividad(self):
            f = self._plan_sel()
            if not f:
                return
            if f["estado"] == "Cancelada":
                messagebox.showwarning("Planificación", "La tarea está cancelada.")
                return
            if f["estado"] == "Cumplida" and not messagebox.askyesno(
                    "Planificación", "La tarea ya figura cumplida. ¿Cargar otra actividad igual?"):
                return
            self._cargar_desde_tarea(f["tarea_dict"], f["fecha_occ"], f["fecha"])

        def _cargar_desde_tarea(self, t, fecha_occ, fecha=None):
            """Abre la pestaña Actividades con los datos de la tarea, lista para completar y AGREGAR."""
            self.nb.select(0)
            self._limpiar_campos()
            if not self.cmb_oper.get() and t["asignado"] in self.db.lista("operarios"):
                self.cmb_oper.set(t["asignado"])
                self._actualizar_codigo()
            self._set(OT_SEC, t["ot_base"])
            self._set(SECTOR, t["sector"])
            self.cmb_maquina["values"] = self.db.lista("maquinas", t["sector"] or "")
            for clave, campo in ((MAQUINA, "maquina"), (CONJ, "conjunto"), (SUBCONJ, "subconjunto"),
                                 (ACTIVIDAD, "tarea"), (PRIORIDAD, "prioridad")):
                self._set(clave, t[campo])
            self._set(FECHA, date.today().strftime("%d-%m-%Y"))
            self.vinculo_pendiente = (t["id"], fecha_occ, t["ot_base"])
            self._ayuda(f"Tarea planificada {t['ot_base']}"
                        + (f" del {fecha:%d-%m-%Y}" if fecha else "")
                        + ". Complete horas, estado y comentarios y presione AGREGAR: la actividad queda "
                          "vinculada a la planificación.")
            if not self.cmb_oper.get():
                self.cmb_oper.focus_set()

        def plan_abrir(self):
            f = self._plan_sel()
            if f:
                self._plan_abrir_fila(f)

        def _plan_abrir_fila(self, f):
            """Una orden de preventivo abre su checklist; cualquier otra tarea, su histórico."""
            if f.get("orden_id"):
                OrdenChecklist(self, f["orden_id"])
            else:
                Historico(self, f["tarea_dict"])

        plan_historico = plan_abrir

        def plan_editar(self):
            f = self._plan_sel()
            if f:
                FormTarea(self, self.db.tareas(f["id"]))

        def plan_marca(self):
            f = self._plan_sel()
            if not f:
                return
            menu = tk.Menu(self, tearoff=0)
            for txt, marca in (("Marcar cumplida", "Hecha"), ("Marcar cancelada", "Cancelada"),
                               ("Quitar marca manual", None)):
                menu.add_command(label=txt, command=lambda m=marca: self._plan_aplicar_marca(f, m))
            menu.tk_popup(self.winfo_pointerx(), self.winfo_pointery())

        def _plan_aplicar_marca(self, f, marca):
            self.db.marcar(f["id"], f["fecha_occ"], marca)
            self.plan_refrescar()

        def plan_eliminar(self):
            f = self._plan_sel()
            if not f:
                return
            n = len(self.db.actividades_de_tareas().get(f["id"], []))
            if not messagebox.askyesno(
                    "Eliminar tarea", f"¿Eliminar la tarea {f['ot_base']} de la planificación"
                    + (" (todas sus fechas)" if es_periodica(f["tarea_dict"]) else "") + "?\n\n"
                    + (f"Sus {n} actividades cargadas NO se borran (siguen en Actividades con OT SEC = {f['ot_base']}).\n\n"
                       if n else "") + "Para dejar de programar una periódica sin borrar su historial, "
                    "edítela y ponga una fecha 'Hasta'.", icon="warning", default="no"):
                return
            self.db.eliminar_tarea(f["id"])
            self.plan_refrescar()

        def plan_importar(self):
            if not self.supervisor:
                messagebox.showwarning("Importar", "Solo un supervisor puede importar la planificación.")
                return
            ruta = filedialog.askopenfilename(title="Seleccionar CalendarioMTO", filetypes=[("Excel", "*.xlsm *.xlsx")])
            if not ruta:
                return
            self.config(cursor="watch")
            self.update()
            try:
                nuevas, repetidas, revisar = self.db.importar_calendario(ruta)
            except Exception as ex:
                messagebox.showerror("Importar", f"No se pudo importar:\n{ex}")
                return
            finally:
                self.config(cursor="")
            msg = f"Tareas agregadas: {nuevas}\nYa existentes (no se tocaron): {repetidas}"
            if revisar:
                msg += (f"\n\n{len(revisar)} tareas periódicas no tienen fecha de inicio o frecuencia. "
                        "Aparecen como 'Sin programar': edítelas para completarlas.")
            messagebox.showinfo("Importar planificación", msg)
            self.nb.select(1)
            self.plan_refrescar()

        def plan_exportar(self):
            if not self.plan_filas:
                messagebox.showinfo("Exportar", "No hay tareas en la vista.")
                return
            ruta = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")],
                                                initialfile=f"planificacion_{self.p_titulo.cget('text').replace(' ', '_').replace('/', '-')}.xlsx")
            if not ruta:
                return
            wb = Workbook(write_only=True)
            ws = wb.create_sheet("Planificación")
            ws.append(["Programada", "Semana", "OT Base", "Tipo", "Sector", "Máquina", "Tarea", "Asignado", "Periodicidad",
                       "Estado", "Realizada", "Retraso (días)", "OT MAN", "Último estado / comentario"])
            for f in self.plan_filas:
                ws.append([f["fecha"], semana_excel(f["fecha"]) if f["fecha"] else None, f["ot_base"], f["tipo"],
                           f["sector"], f["maquina"], f["tarea"], f["asignado"], f["periodicidad"], f["estado"],
                           f["realizada"], f["retraso"], f["ot_man"], f["detalle"]])
            try:
                wb.save(ruta)
            except PermissionError:
                messagebox.showerror("Error", "No se pudo guardar: el archivo está abierto en Excel.")
                return
            messagebox.showinfo("Exportar", f"Exportado a:\n{ruta}")

        # ---------------- Plan por máquina ----------------
        def _armar_plan_maquina(self):
            self.tab_maq = ttk.Frame(self.nb)
            self.nb.add(self.tab_maq, text="  Plan por máquina  ")
            p = self.tab_maq
            fil = ttk.Frame(p, padding=(10, 8, 10, 4))
            fil.pack(fill="x")
            ttk.Label(fil, text="Sector").pack(side="left", padx=(0, 2))
            self.m_sector = ttk.Combobox(fil, width=10, state="readonly")
            self.m_sector.pack(side="left")
            ttk.Label(fil, text="Máquina").pack(side="left", padx=(10, 2))
            self.m_maquina = ttk.Combobox(fil, width=14, state="readonly")
            self.m_maquina.pack(side="left")
            self.m_solo_per = tk.BooleanVar(value=True)
            ttk.Checkbutton(fil, text="Solo periódicas", variable=self.m_solo_per,
                            command=self.maq_refrescar).pack(side="left", padx=10)
            ttk.Label(fil, text="Buscar").pack(side="left", padx=(10, 2))
            self.m_texto = ttk.Entry(fil, width=18)
            self.m_texto.pack(side="left")
            self.m_texto.bind("<KeyRelease>", lambda e: self._maq_buscar_tecla())
            self.m_texto.bind("<Escape>", lambda e: (self.m_texto.delete(0, "end"), self._maq_mostrar()))
            self.m_sector.bind("<<ComboboxSelected>>", lambda e: (self._maq_maquinas(), self.maq_refrescar()))
            self.m_maquina.bind("<<ComboboxSelected>>", lambda e: self.maq_refrescar())
            self.m_info = ttk.Label(fil, text="", foreground="#333")
            self.m_info.pack(side="right")

            acc = ttk.Frame(p, padding=(10, 0, 10, 4))
            acc.pack(fill="x")
            ttk.Button(acc, text="Ver histórico", command=self.maq_historico).pack(side="left")
            ttk.Separator(acc, orient="vertical").pack(side="left", fill="y", padx=8)
            for txt, cmd in [("Nuevo ítem", self.maq_nuevo), ("Editar", self.maq_editar),
                             ("Editar en lote", self.maq_lote), ("Eliminar", self.maq_eliminar)]:
                b = ttk.Button(acc, text=txt, command=cmd)
                b.pack(side="left", padx=2)
                self.solo_sup.append(b)
            ttk.Separator(acc, orient="vertical").pack(side="left", fill="y", padx=8)
            b = ttk.Button(acc, text="GENERAR ORDEN DE PREVENTIVO", command=self.maq_generar_orden)
            b.pack(side="left")
            self.solo_sup.append(b)
            for txt, cmd in [("Generar planilla modelo", self.generar_planilla),
                             ("Importar planilla", self.importar_planilla)]:
                ttk.Button(acc, text=txt, command=cmd).pack(side="right", padx=2)

            panel = ttk.PanedWindow(p, orient="vertical")
            panel.pack(fill="both", expand=True, padx=10, pady=(0, 8))
            arriba = ttk.Frame(panel)
            abajo = ttk.Frame(panel)
            panel.add(arriba, weight=3)
            panel.add(abajo, weight=1)
            cols = [("ot_base", "OT Base", 120), ("maquina", "Máquina", 80), ("conjunto", "Conjunto", 110),
                    ("subconjunto", "Subconjunto", 110), ("tarea", "Ítem / Tarea", 320), ("periodicidad", "Periodicidad", 105),
                    ("inicio", "Inicio", 90), ("asignado", "Asignado", 90), ("ultima", "Última realizada", 110),
                    ("proxima", "Próxima", 90), ("estado", "Estado actual", 115)]
            self.m_tree = ttk.Treeview(arriba, columns=[c[0] for c in cols], show="headings", selectmode="extended",
                                       height=6)
            for c, t, w in cols:
                self.m_tree.heading(c, text=t)
                self.m_tree.column(c, width=w, stretch=c == "tarea", anchor="w")
            sy = ttk.Scrollbar(arriba, orient="vertical", command=self.m_tree.yview)
            self.m_tree.configure(yscrollcommand=sy.set)
            self.m_tree.pack(side="left", fill="both", expand=True)
            sy.pack(side="right", fill="y")
            for est, color in dict(COLORES_ESTADO, **{"Al día": "#007a33"}).items():
                self.m_tree.tag_configure(est, foreground=color)
            self.m_tree.bind("<Double-1>", lambda e: self.maq_editar() if self.supervisor else self.maq_historico())

            cab = ttk.Frame(abajo)
            cab.pack(fill="x", pady=(6, 2))
            ttk.Label(cab, text="Órdenes de preventivo de la máquina", font=("Segoe UI", 9, "bold")).pack(side="left")
            ttk.Button(cab, text="Abrir checklist", command=self.maq_abrir_orden).pack(side="left", padx=8)
            cols = [("ot", "OT", 130), ("fecha", "Fecha", 92), ("items", "Ítems", 55), ("hechos", "Hechos", 60),
                    ("no_hechos", "No hechos", 75), ("pend", "Pendientes", 80), ("estado", "Estado OT", 100),
                    ("ot_man", "OT MAN", 200)]
            self.m_ord = ttk.Treeview(abajo, columns=[c[0] for c in cols], show="headings", height=4, selectmode="browse")
            for c, t, w in cols:
                self.m_ord.heading(c, text=t)
                self.m_ord.column(c, width=w, anchor="w", stretch=c == "ot_man")
            self.m_ord.pack(fill="both", expand=True)
            self.m_ord.bind("<Double-1>", lambda e: self.maq_abrir_orden())
            self.m_filas = []

        def _refrescar_planes(self):
            self.plan_refrescar()
            self.maq_refrescar()

        def _maq_maquinas(self):
            sec = self.m_sector.get()
            maqs = set(self.db.lista("maquinas", sec)) if sec else set()
            maqs |= {t["maquina"] for t in self.db.tareas() if t["maquina"] and t["sector"] == sec and t["tipo"] != "PM"}
            self.m_maquina["values"] = [""] + sorted(maqs)
            if self.m_maquina.get() not in maqs:
                self.m_maquina.set("")

        def maq_refrescar(self):
            sec, maq = self.m_sector.get(), self.m_maquina.get()
            self.m_ord.delete(*self.m_ord.get_children())
            if not sec:
                self.m_filas = []
                self._maq_mostrar()
                self.m_info.config(text="Elija un sector (y opcionalmente una máquina).")
                return
            hoy = date.today()
            try:
                acts, marcas, items = self.db.actividades_de_tareas(), self.db.marcas(), self.db.items_orden_mapa()
                ordenes = self.db.ordenes_por_tarea()
                tareas = self.db.tareas()
            except sqlite3.Error as ex:
                self.m_info.config(text=f"⚠ No se pudo leer la base: {ex}")
                return
            filas = []
            for t in tareas:
                if t["sector"] != sec or (maq and t["maquina"] != maq) or t["id"] in ordenes:
                    continue
                if self.m_solo_per.get() and not es_periodica(t):
                    continue
                occs = ocurrencias_tarea(t, acts.get(t["id"], []), marcas, items, hoy + timedelta(days=400), hoy)
                ultima = proxima = None
                if occs is None:
                    estado = "Sin programar"
                else:
                    hechas = [o["realizada"] or o["fecha"] for o in occs if o["estado"] == "Cumplida"]
                    ultima = max([d for d in hechas if d], default=None)
                    futuras = [o["fecha"] for o in occs if o["fecha"] and o["fecha"] >= hoy
                               and o["estado"] not in ("Cumplida", "Cancelada")]
                    proxima = futuras[0] if futuras else None
                    pasadas = [o for o in occs if o["fecha"] and o["fecha"] < hoy]
                    actual = pasadas[-1] if pasadas else (occs[0] if occs else None)
                    estado = actual["estado"] if actual else "Pendiente"
                    if estado == "Cumplida":
                        estado = "Al día"
                filas.append(dict(t=t, ultima=ultima, proxima=proxima, estado=estado))
            filas.sort(key=lambda f: (f["t"]["maquina"] or "", f["t"]["conjunto"] or "", f["t"]["subconjunto"] or "",
                                      f["t"]["tarea"] or ""))
            self.m_filas = filas
            self._maq_mostrar()
            if maq:   # órdenes de la máquina
                for o in self.db.ordenes_de(sec, maq):
                    od = self.db.orden(o["id"])
                    tt = self.db.tareas(od["tarea_id"]) if od else None
                    est, ots = "", ""
                    if tt:
                        oc = ocurrencias_tarea(tt, acts.get(tt["id"], []), marcas, items, hoy, hoy) or []
                        if oc:
                            est = oc[0]["estado"]
                            ots = ", ".join(str(a["ot"]) for a in oc[0]["acts"])
                    pend = (o["items"] or 0) - (o["hechos"] or 0) - (o["no_hechos"] or 0)
                    self.m_ord.insert("", "end", iid=str(o["id"]), values=[
                        o["ot"], db_a_texto(o["fecha"]), o["items"] or 0, o["hechos"] or 0, o["no_hechos"] or 0, pend,
                        est, ots])

        def _maq_buscar_tecla(self):
            if getattr(self, "_m_busqueda", None):
                self.after_cancel(self._m_busqueda)
            self._m_busqueda = self.after(250, self._maq_mostrar)

        def _maq_mostrar(self):
            self._m_busqueda = None
            filas = self.m_filas
            palabras = self.m_texto.get().lower().split()
            if palabras:
                campos = ("ot_base", "maquina", "conjunto", "subconjunto", "tarea", "asignado")
                filas = [f for f in filas if all(p in " ".join(str(f["t"][k] or "") for k in campos).lower()
                                                 for p in palabras)]
            sel = self.m_tree.selection()
            self.m_tree.delete(*self.m_tree.get_children())
            for f in filas:
                t = f["t"]
                self.m_tree.insert("", "end", iid=str(t["id"]), tags=(f["estado"],), values=[
                    t["ot_base"] or "", t["maquina"] or "", t["conjunto"] or "", t["subconjunto"] or "", t["tarea"] or "",
                    texto_periodicidad(t), db_a_texto(t["fecha"]), t["asignado"] or "",
                    f["ultima"].strftime("%d-%m-%Y") if f["ultima"] else "",
                    f["proxima"].strftime("%d-%m-%Y") if f["proxima"] else "", f["estado"]])
            sel = [i for i in sel if self.m_tree.exists(i)]
            if sel:
                self.m_tree.selection_set(sel)
            if self.m_sector.get():
                cuenta = {}
                for f in filas:
                    cuenta[f["estado"]] = cuenta.get(f["estado"], 0) + 1
                self.m_info.config(text=f"{len(filas)} ítems   " + "   ".join(f"{e}: {n}" for e, n in sorted(cuenta.items())))

        def _maq_sel(self, varias=False):
            sel = self.m_tree.selection()
            if not sel:
                messagebox.showinfo("Plan por máquina", "Seleccione " + ("uno o más ítems." if varias else "un ítem."))
                return None
            if varias:
                return [int(i) for i in sel]
            return self.db.tareas(int(sel[0]))

        def maq_historico(self):
            t = self._maq_sel()
            if t:
                Historico(self, t)

        def maq_nuevo(self):
            if not self.m_sector.get():
                messagebox.showinfo("Nuevo ítem", "Elija primero el sector y la máquina.")
                return
            FormTarea(self, defecto=dict(sector=self.m_sector.get(), maquina=self.m_maquina.get(), periodica=True))

        def maq_editar(self):
            t = self._maq_sel()
            if t:
                FormTarea(self, t)

        def maq_lote(self):
            ids = self._maq_sel(varias=True)
            if ids:
                EdicionLote(self, ids)

        def maq_eliminar(self):
            ids = self._maq_sel(varias=True)
            if not ids:
                return
            if not messagebox.askyesno("Eliminar", f"¿Eliminar {len(ids)} ítem(s) del plan?\n\nLas actividades ya "
                                                   "cargadas no se borran. Para dejar de programar un ítem sin perder su "
                                                   "historial, edítelo y ponga una fecha 'Hasta'.", icon="warning",
                                       default="no"):
                return
            for i in ids:
                self.db.eliminar_tarea(i)
            self._refrescar_planes()

        def maq_generar_orden(self):
            if not (self.m_sector.get() and self.m_maquina.get()):
                messagebox.showinfo("Orden de preventivo", "Elija el sector y la máquina.")
                return
            GenerarOrden(self, self.m_sector.get(), self.m_maquina.get())

        def maq_abrir_orden(self):
            sel = self.m_ord.selection()
            if not sel:
                messagebox.showinfo("Órdenes", "Seleccione una orden.")
                return
            OrdenChecklist(self, int(sel[0]))

        # ---------------- Carga masiva y configuración ----------------
        def generar_planilla(self):
            ruta = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")],
                                                initialfile="Plantilla_Plan_Preventivo.xlsx")
            if not ruta:
                return
            try:
                generar_planilla(self.db, ruta)
            except PermissionError:
                messagebox.showerror("Error", "No se pudo guardar: el archivo está abierto en Excel.")
                return
            messagebox.showinfo("Planilla modelo", f"Planilla creada:\n{ruta}")

        def importar_planilla(self):
            if not self.supervisor:
                messagebox.showwarning("Importar", "Solo un supervisor puede importar ítems.")
                return
            ruta = filedialog.askopenfilename(title="Seleccionar planilla de ítems", filetypes=[("Excel", "*.xlsx *.xlsm")])
            if not ruta:
                return
            try:
                res = leer_planilla_items(self.db, ruta)
            except Exception as ex:
                messagebox.showerror("Importar", f"No se pudo leer la planilla:\n{ex}")
                return
            ImportarPreview(self, res)

        def config_plan(self):
            if not self.supervisor:
                messagebox.showwarning("Configuración", "Solo un supervisor puede cambiar la configuración.")
                return
            tol = simpledialog.askinteger("Tolerancia", "Días de tolerancia para considerar una tarea 'a tiempo':",
                                          initialvalue=self._tolerancia(), minvalue=0, maxvalue=365, parent=self)
            if tol is None:
                return
            ant = simpledialog.askinteger("Anticipación", "Al generar una orden de preventivo, incluir los ítems que "
                                                          "vencen en los próximos ... días:",
                                          initialvalue=self.cfg.getint("plan", "anticipacion_orden_dias", fallback=15),
                                          minvalue=0, maxvalue=365, parent=self)
            if ant is None:
                return
            self.cfg.set("plan", "tolerancia_dias", str(tol))
            self.cfg.set("plan", "anticipacion_orden_dias", str(ant))
            guardar_config(self.cfg)
            self._plan_mostrar()
            messagebox.showinfo("Configuración", "Guardado en config.ini (vale para todas las PC que usen ese archivo).")

        def importar(self):
            if not self.db.vacia:
                if not self.supervisor:
                    messagebox.showwarning("Importar", "Solo un supervisor puede reimportar el Excel.")
                    return
                if not messagebox.askyesno(
                        "Importar", "Esto REEMPLAZA todos los datos y listas de la base por los del Excel.\n\n"
                                    "Asegúrese de que nadie más esté usando el programa en otra PC.\n"
                                    "Antes de importar se crea un respaldo.\n\n¿Continuar?", icon="warning"):
                    return
            con_listas = self.db.vacia or messagebox.askyesno(
                "Listas", f"¿Reemplazar también las listas (operarios, sectores, máquinas, etc.) por las de la "
                          f"hoja '{HOJA_DATOS}'?\n\nNo = conservar las listas editadas en el programa.")
            ruta = filedialog.askopenfilename(title="Seleccionar CargaActividades",
                                              filetypes=[("Excel", "*.xlsm *.xlsx")])
            if not ruta:
                return
            self.config(cursor="watch")
            self.update()
            try:
                if not self.db.vacia:
                    respaldar(self.db.ruta, self.cfg.get("respaldo", "carpeta"), manual=True)
                n = self.db.importar_excel(ruta, listas=con_listas)
            except KeyError as ex:
                messagebox.showerror("Error", f"El archivo no tiene la hoja {ex}.")
                return
            except Exception as ex:
                messagebox.showerror("Error al importar", str(ex))
                return
            finally:
                self.config(cursor="")
            self._cargar_listas()
            self.limpiar()
            messagebox.showinfo("Importación", f"Se importaron {n} registros"
                                               + (f" y las listas de '{HOJA_DATOS}'." if con_listas else "."))

        def exportar(self, completa):
            if self.db.vacia:
                return
            if not completa and not self.filas_lista:
                messagebox.showinfo("Exportar", "La lista está vacía.")
                return
            ruta = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")],
                                                initialfile=f"actividades_{datetime.now():%Y%m%d_%H%M}.xlsx")
            if not ruta:
                return
            try:
                self.db.exportar(ruta, None if completa else self.filas_lista)
            except PermissionError:
                messagebox.showerror("Error", "No se pudo guardar: el archivo está abierto en Excel.")
                return
            messagebox.showinfo("Exportar", f"Exportado a:\n{ruta}")

    App().mainloop()


if __name__ == "__main__":
    iniciar_interfaz()
